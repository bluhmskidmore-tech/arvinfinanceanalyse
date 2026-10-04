"""Hand-derived amounts through real workbook ingestion, persistence, and HTTP reads.

See docs/pnl/product-category-manual-reference.md for the arithmetic and scope.
Expected values are literals, never captured from the implementation under test.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import duckdb
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.api.routes.product_category_pnl import router
from backend.app.governance.settings import get_settings
from backend.app.security.auth_context import ROLE_HEADER_TRUST_ENV
from backend.app.tasks.product_category_pnl import materialize_product_category_pnl_sync

# Per currency: asset average balance, asset income credit, liability expense debit.
# All monetary inputs are synthetic yuan amounts; no production data is loaded.
SOURCE_AMOUNTS = {
    "2026-01-31": {"CNX": (36_500_000, 620_000, 310_000), "CNY": (21_900_000, 372_000, 186_000)},
    "2026-02-28": {"CNX": (73_000_000, 1_120_000, 280_000), "CNY": (36_500_000, 560_000, 168_000)},
}
LIABILITY_SCALE = {"CNX": -36_500_000, "CNY": -21_900_000}

FIELDS = (
    "cnx_scale", "cny_scale", "foreign_scale", "cnx_cash", "cny_cash", "foreign_cash",
    "cny_ftp", "foreign_ftp", "cny_net", "foreign_net", "business_net_income", "weighted_yield",
)
MANUAL_REFERENCE = {
    "january": {
        "asset": (
            "36500000", "21900000", "14600000", "620000", "372000", "248000",
            "29760", "19840", "342240", "228160", "570400", "20",
        ),
        "liability": (
            "-36500000", "-21900000", "-14600000", "-310000", "-186000", "-124000",
            "-29760", "-19840", "-156240", "-104160", "-260400", "10",
        ),
        "grand": (
            "0", "0", "0", "310000", "186000", "124000",
            "0", "0", "186000", "124000", "310000", None,
        ),
    },
    "february": {
        "asset": (
            "73000000", "36500000", "36500000", "1120000", "560000", "560000",
            "44800", "44800", "515200", "515200", "1030400", "20",
        ),
        "liability": (
            "-36500000", "-21900000", "-14600000", "-280000", "-168000", "-112000",
            "-26880", "-17920", "-141120", "-94080", "-235200", "10",
        ),
        "grand": (
            "0", "0", "0", "840000", "392000", "448000",
            "17920", "26880", "374080", "421120", "795200", None,
        ),
    },
    "year_to_february": {
        "asset": (
            "53822033.89830508", "28828813.55932203", "24993220.33898305",
            "1740000", "932000", "808000", "74560", "64640",
            "857440", "743360", "1600800", "20",
        ),
        "liability": (
            "-36500000", "-21900000", "-14600000", "-590000", "-354000", "-236000",
            "-56640", "-37760", "-297360", "-198240", "-495600", "10",
        ),
        "grand": (
            "0", "0", "0", "1150000", "578000", "572000",
            "17920", "26880", "560080", "545120", "1105200", None,
        ),
    },
}


def _write_reference_workbooks(source_dir: Path, report_date: str) -> None:
    month_key = report_date[:7].replace("-", "")
    ledger = Workbook()
    ledger.remove(ledger.active)
    for sheet_name, currency in (("综本", "CNX"), ("人民币", "CNY")):
        _, asset_credit, liability_debit = SOURCE_AMOUNTS[report_date][currency]
        sheet = ledger.create_sheet(sheet_name)
        sheet.cell(1, 1, "合成验收样例")
        sheet.cell(5, 1, f"会计期间： {report_date[:8]}01--{report_date}")
        sheet.append(["科目代码", "科目名称", "币种", "期初余额", "本期借方", "本期贷方", "期末余额"])
        # Distinct ending balances catch accidental use of stock values as income.
        sheet.append(["50204000001", "合成拆放收入", currency, -9_000_000, 0, asset_credit, -8_000_000])
        sheet.append(["52206000001", "合成存放支出", currency, 7_000_000, liability_debit, 0, 6_000_000])
    ledger.save(source_dir / f"总账对账{month_key}.xlsx")
    ledger.close()

    averages = Workbook()
    averages.active.title = "年"
    averages.create_sheet("月")
    for sheet in averages:
        sheet.append(["合成验收样例"])
        sheet.append([f"日期： {report_date}"])
        sheet.append(["币种", "科目", "日均余额", None])
        for currency, (asset_scale, _, _) in SOURCE_AMOUNTS[report_date].items():
            selected_sheet = "年" if report_date == "2026-01-31" else "月"
            # Intentionally different unused fields detect an incorrect period basis.
            factor = 1 if sheet.title == selected_sheet else 7
            # Scale mapping uses exact aggregate accounts; leaf rows must not double it.
            sheet.append([currency, "120", asset_scale * factor, None])
            sheet.append([currency, "234", LIABILITY_SCALE[currency] * factor, None])
            sheet.append([currency, "12000000001", asset_scale * factor, None])
            sheet.append([currency, "23400000001", LIABILITY_SCALE[currency] * factor, None])
    averages.save(source_dir / f"日均{month_key}.xlsx")
    averages.close()


@pytest.fixture
def reference_flow(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, seed_wildcard_scope):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    database = tmp_path / "reference.duckdb"
    governance = tmp_path / "governance"
    for name, value in {
        "MOSS_ENVIRONMENT": "development",
        "MOSS_GOVERNANCE_BACKEND": "jsonl",
        "MOSS_GOVERNANCE_SQL_DSN": "",
        "MOSS_DUCKDB_PATH": str(database),
        "MOSS_PRODUCT_CATEGORY_SOURCE_DIR": str(source_dir),
        "MOSS_GOVERNANCE_PATH": str(governance),
        # The report-year rule must win over a different configured fallback.
        "MOSS_FTP_RATE_PCT": "9.99",
        ROLE_HEADER_TRUST_ENV: "1",
    }.items():
        monkeypatch.setenv(name, value)
    get_settings.cache_clear()
    for report_date in SOURCE_AMOUNTS:
        _write_reference_workbooks(source_dir, report_date)
    receipt = materialize_product_category_pnl_sync(run_id="manual-reference-initial")
    assert receipt["status"] == "completed"
    assert receipt["report_dates"] == ["2026-01-31", "2026-02-28"]
    app = FastAPI()
    app.include_router(router)
    try:
        with TestClient(app) as client:
            yield client, source_dir, database, receipt
    finally:
        get_settings.cache_clear()


def _read_result(client: TestClient, report_date: str, view: str) -> dict:
    response = client.get(
        "/ui/pnl/product-category",
        params={"report_date": report_date, "view": view},
        headers={"X-User-Id": "manual-reference-reader", "X-User-Role": "viewer"},
    )
    assert response.status_code == 200, response.text
    envelope = response.json()
    assert envelope["result_meta"]["basis"] == "formal"
    assert envelope["result_meta"]["formal_use_allowed"] is True
    result = envelope["result"]
    assert result["report_date"] == report_date
    assert result["view"] == view
    assert result["scenario_rate_pct"] is None
    return result


def _assert_manual_amounts(row: dict, expected: tuple, context: str) -> None:
    assert Decimal(str(row["baseline_ftp_rate_pct"])) == Decimal("1.60"), context
    for field, value in zip(FIELDS, expected, strict=True):
        if value is None:
            assert row[field] is None, (context, field)
        else:
            assert Decimal(str(row[field])) == Decimal(value), (context, field, row[field], value)


def test_hand_calculated_reference_reaches_persisted_rows_and_http(reference_flow) -> None:
    client, _, database, _ = reference_flow
    cases = (
        ("2026-01-31", "monthly", "january"),
        ("2026-02-28", "monthly", "february"),
        ("2026-02-28", "ytd", "year_to_february"),
        ("2026-02-28", "qtd", "year_to_february"),
        ("2026-02-28", "year_to_report_month_end", "year_to_february"),
    )
    categories = {
        "interbank_lending_assets": "asset", "interest_earning_assets": "asset", "asset_total": "asset",
        "interbank_deposits": "liability", "liability_total": "liability", "grand_total": "grand",
    }
    for report_date, view, reference in cases:
        expected = MANUAL_REFERENCE[reference]
        result = _read_result(client, report_date, view)
        api_rows = {row["category_id"]: row for row in result["rows"]}
        with duckdb.connect(str(database), read_only=True) as connection:
            cursor = connection.execute(
                "select * from product_category_pnl_formal_read_model where report_date = ? and view = ?",
                [report_date, view],
            )
            columns = [column[0] for column in cursor.description]
            stored_rows = {
                row["category_id"]: row
                for row in (dict(zip(columns, values, strict=True)) for values in cursor.fetchall())
            }
        for category_id, side in categories.items():
            context = f"{report_date} {view} {category_id}"
            _assert_manual_amounts(stored_rows[category_id], expected[side], f"stored {context}")
            _assert_manual_amounts(api_rows[category_id], expected[side], f"API {context}")
        for side in ("asset", "liability", "grand"):
            _assert_manual_amounts(result[f"{side}_total"], expected[side], f"headline {report_date} {view} {side}")


def test_historical_source_revision_updates_cumulative_amount_but_preserves_later_month(reference_flow) -> None:
    client, source_dir, _, initial_receipt = reference_flow
    february_before = _read_result(client, "2026-02-28", "monthly")
    # A synthetic January income correction affects only the CNX-minus-CNY portion.
    ledger_path = source_dir / "总账对账202601.xlsx"
    workbook = load_workbook(ledger_path)
    workbook["综本"]["F7"] = 630_000
    workbook.save(ledger_path)
    workbook.close()
    receipt = materialize_product_category_pnl_sync(run_id="manual-reference-january-revised")
    assert receipt["status"] == "completed"
    assert receipt["source_version"] != initial_receipt["source_version"]
    february_after = _read_result(client, "2026-02-28", "monthly")
    for field in ("rows", "asset_total", "liability_total", "grand_total"):
        assert february_after[field] == february_before[field], field
    for report_date, view, expected_net, expected_foreign in (
        ("2026-01-31", "monthly", "320000", "134000"),
        ("2026-02-28", "ytd", "1115200", "555120"),
        ("2026-02-28", "qtd", "1115200", "555120"),
        ("2026-02-28", "year_to_report_month_end", "1115200", "555120"),
    ):
        result = _read_result(client, report_date, view)
        assert Decimal(str(result["grand_total"]["business_net_income"])) == Decimal(expected_net)
        assert Decimal(str(result["grand_total"]["foreign_net"])) == Decimal(expected_foreign)
