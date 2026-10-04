# 回归：position_scope=asset/liability 时，跨口径表（期限缺口/监管阈值/现金流日历/
# 利率分布/对手方类型/概览卡片/右栏运营区块）必须使用全量口径数据构建，
# 不得把负债/缺口/利差静默降级为 0。单口径表仍按用户 scope 过滤。
from __future__ import annotations

from dataclasses import replace
from datetime import date
from decimal import Decimal
from importlib import import_module, reload
from inspect import isfunction

import pytest

from backend.app.core_finance.balance_analysis import (
    FormalTywBalanceFactRow,
    FormalZqtzBalanceFactRow,
)
from backend.app.core_finance.balance_analysis_workbook import (
    build_balance_analysis_workbook_payload,
)
from backend.app.core_finance.balance_workbook import (
    build_balance_analysis_workbook_payload as build_modular_balance_workbook_payload,
)
from backend.app.schemas.balance_analysis import BalanceAnalysisRiskAlertRow

RD = date(2026, 3, 31)
MAT = date(2027, 3, 31)

CROSS_SCOPE_TABLE_KEYS = (
    "maturity_gap",
    "regulatory_limits",
    "cashflow_calendar",
    "rate_distribution",
    "counterparty_types",
)
CROSS_SCOPE_SECTION_KEYS = (
    "decision_items",
    "event_calendar",
    "risk_alerts",
)

_COMPATIBILITY_CALCULATION_EXPORTS = {
    "_utils": (
        "_ZERO",
        "_TEN_THOUSAND",
        "_MATURITY_BUCKETS",
        "_MISSING_MATURITY_FALLBACK_BUCKET",
        "_RATE_BUCKETS",
        "_LIQUIDITY_LAYER_ORDER",
        "_LIQUIDITY_LEVEL1_BOND_TYPES",
        "_LIQUIDITY_HQLA_HAIRCUTS",
        "_LIQUIDITY_HIGH_RATING",
        "_CAMPISI_POLICY_BOND",
        "_group_rows",
        "_to_finite_decimal",
        "_sum_decimal",
        "_weighted_average",
        "_merged_weighted_average",
        "_remaining_years",
        "_optional_remaining_years",
        "_match_bucket",
        "_matches_maturity_bucket",
        "_safe_ratio",
        "_spread_bp",
        "_rate_value",
        "_normalize_interest_mode",
        "_to_wanyuan",
        "_decimal_value",
        "_severity_from_gap",
        "_month_ladder",
        "_month_key",
        "_month_key_from_index",
        "_card",
        "_section",
        "_table",
    ),
    "_bond_tables": (
        "_build_cards",
        "_build_bond_business_type_table",
        "_bond_business_type_label",
        "_build_maturity_gap_table",
        "_build_issuance_business_type_table",
        "_build_issuer_concentration_table",
        "_classify_liquidity_layer",
        "_build_liquidity_layers_table",
        "_build_portfolio_comparison_table",
        "_build_cashflow_calendar_table",
        "_build_vintage_analysis_table",
        "_build_customer_attribute_analysis_table",
    ),
    "_ifrs9_tables": (
        "_build_ifrs9_classification_table",
        "_build_ifrs9_position_scope_table",
        "_build_ifrs9_source_family_table",
        "_build_account_category_comparison_table",
        "_build_rule_reference_table",
    ),
    "_analysis_tables": (
        "_build_currency_split_table",
        "_build_rating_table",
        "_rating_bucket_label",
        "_is_interest_rate_bond",
        "_build_rate_distribution_table",
        "_build_industry_table",
        "_build_counterparty_type_table",
        "_build_campisi_table",
        "_build_cross_analysis_table",
        "_build_interest_mode_table",
        "_build_decision_items_table",
        "_maturity_full_scope_gap_value",
        "_build_event_calendar_table",
    ),
    "_risk_tables": (
        "_regulatory_metric_status",
        "_build_regulatory_limits_table",
        "_build_overdue_credit_quality_detail_table",
        "_build_overdue_credit_quality_rating_table",
        "_build_risk_alerts_table",
        "_maturity_full_scope_gap_value",
    ),
}


def _zqtz_asset() -> FormalZqtzBalanceFactRow:
    return FormalZqtzBalanceFactRow(
        report_date=RD,
        instrument_code="A0001",
        instrument_name="资产债券",
        portfolio_name="P",
        cost_center="C",
        account_category="可供出售类资产",
        asset_class="信用债",
        bond_type="企业债",
        issuer_name="发行人A",
        industry_name="工业",
        rating="AAA",
        invest_type_std="A",
        accounting_basis="FVOCI",
        position_scope="asset",
        currency_basis="native",
        currency_code="CNY",
        face_value_amount=Decimal("100000000"),
        market_value_amount=Decimal("101000000"),
        amortized_cost_amount=Decimal("100000000"),
        accrued_interest_amount=Decimal("0"),
        # 百分数口径：3.20 = 3.20%（2026-07-19 审计取证 1 裁决）
        coupon_rate=Decimal("3.20"),
        ytm_value=Decimal("3.20"),
        maturity_date=MAT,
        interest_mode="固定",
        is_issuance_like=False,
    )


def _zqtz_issuance() -> FormalZqtzBalanceFactRow:
    return FormalZqtzBalanceFactRow(
        report_date=RD,
        instrument_code="L0001",
        instrument_name="发行类债券",
        portfolio_name="P",
        cost_center="C",
        account_category="发行类债劵",
        asset_class="债券类",
        bond_type="同业存单",
        issuer_name="本行",
        industry_name="金融业",
        rating="",
        invest_type_std="A",
        accounting_basis="FVOCI",
        position_scope="liability",
        currency_basis="native",
        currency_code="CNY",
        face_value_amount=Decimal("40000000"),
        market_value_amount=Decimal("40000000"),
        amortized_cost_amount=Decimal("40000000"),
        accrued_interest_amount=Decimal("0"),
        coupon_rate=Decimal("2.50"),
        ytm_value=Decimal("2.50"),
        maturity_date=MAT,
        interest_mode="固定",
        is_issuance_like=True,
    )


def _tyw_asset() -> FormalTywBalanceFactRow:
    return FormalTywBalanceFactRow(
        report_date=RD,
        position_id="TA1",
        product_type="拆放同业",
        position_side="asset",
        counterparty_name="银行A",
        account_type="资产账户",
        special_account_type="一般",
        core_customer_type="股份制银行",
        invest_type_std="H",
        accounting_basis="AC",
        position_scope="asset",
        currency_basis="native",
        currency_code="CNY",
        principal_amount=Decimal("30000000"),
        accrued_interest_amount=Decimal("0"),
        funding_cost_rate=Decimal("2.10"),
        maturity_date=MAT,
    )


def _tyw_liability() -> FormalTywBalanceFactRow:
    return FormalTywBalanceFactRow(
        report_date=RD,
        position_id="TL1",
        product_type="同业存放",
        position_side="liability",
        counterparty_name="银行B",
        account_type="负债账户",
        special_account_type="一般",
        core_customer_type="股份制银行",
        invest_type_std="H",
        accounting_basis="AC",
        position_scope="liability",
        currency_basis="native",
        currency_code="CNY",
        principal_amount=Decimal("50000000"),
        accrued_interest_amount=Decimal("0"),
        funding_cost_rate=Decimal("1.80"),
        maturity_date=MAT,
    )


def _table(payload: dict, key: str) -> dict:
    for table in payload["tables"]:
        if table["key"] == key:
            return table
    raise AssertionError(f"table {key} not found")


@pytest.mark.parametrize("scope", ["all", "asset", "liability"])
@pytest.mark.parametrize("missing_rate", [None, Decimal("NaN"), Decimal("Infinity"), Decimal("-Infinity"), "", "BROKEN"])
def test_rate_distribution_discloses_missing_rates_without_losing_balances(scope, missing_rate):
    zqtz_rows = [
        replace(_zqtz_asset(), coupon_rate=rate)
        for rate in (missing_rate, Decimal("0"), Decimal("3.20"))
    ]
    tyw_rows = [
        replace(row, funding_cost_rate=rate)
        for row in (_tyw_asset(), _tyw_liability())
        for rate in (missing_rate, Decimal("0"), Decimal("3.20"))
    ]
    payload = build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope=scope,
        currency_basis="native",
        zqtz_rows=[row for row in zqtz_rows if scope == "all" or row.position_scope == scope],
        tyw_rows=[row for row in tyw_rows if scope == "all" or row.position_scope == scope],
        zqtz_full_rows=zqtz_rows,
        tyw_full_rows=tyw_rows,
    )
    distribution = _table(payload, "rate_distribution")
    buckets = {row["bucket"]: row for row in distribution["rows"]}
    assert distribution["rows"][0]["bucket"] == "利率缺失"
    assert buckets["零息/无息"]["bond_count"] == 1
    assert "利率缺失" in buckets
    for count_key, amount_key, expected_amount in (
        ("bond_count", "bond_amount", Decimal("10000")),
        ("interbank_asset_count", "interbank_asset_amount", Decimal("3000")),
        ("interbank_liability_count", "interbank_liability_amount", Decimal("5000")),
    ):
        for label in ("利率缺失", "零息/无息", "3.0%-3.5%"):
            assert buckets[label][count_key] == 1
            assert buckets[label][amount_key] == expected_amount
        assert sum(row[count_key] for row in distribution["rows"]) == 3
        assert sum(row[amount_key] for row in distribution["rows"]) == expected_amount * 3


@pytest.mark.parametrize("negative_rate", [Decimal("-0.50"), "-0.50"])
def test_rate_distribution_keeps_negative_rates_separate_from_observed_zero(negative_rate):
    payload = build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope="all",
        currency_basis="native",
        zqtz_rows=[replace(_zqtz_asset(), coupon_rate=rate) for rate in (negative_rate, Decimal("0"))],
        tyw_rows=[
            replace(row, funding_cost_rate=rate)
            for row in (_tyw_asset(), _tyw_liability())
            for rate in (negative_rate, Decimal("0"))
        ],
    )
    buckets = {row["bucket"]: row for row in _table(payload, "rate_distribution")["rows"]}
    for count_key, amount_key, amount in (
        ("bond_count", "bond_amount", Decimal("10000")),
        ("interbank_asset_count", "interbank_asset_amount", Decimal("3000")),
        ("interbank_liability_count", "interbank_liability_amount", Decimal("5000")),
    ):
        assert buckets["利率缺失"][count_key] == 0
        for label in ("0%以下", "零息/无息"):
            assert buckets[label][count_key] == 1
            assert buckets[label][amount_key] == amount


def test_rate_distribution_excel_preserves_missing_rate_amounts():
    from io import BytesIO

    from openpyxl import load_workbook

    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    payload = build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope="all",
        currency_basis="native",
        zqtz_rows=[replace(_zqtz_asset(), coupon_rate=None)],
        tyw_rows=[replace(row, funding_cost_rate=None) for row in (_tyw_asset(), _tyw_liability())],
    )
    workbook = load_workbook(BytesIO(_build_balance_analysis_workbook_xlsx_bytes(payload)), data_only=True)
    try:
        missing_row = next(
            row for row in workbook["利率分布"].iter_rows(values_only=True) if row[0] == "利率缺失"
        )
        assert missing_row[:7] == ("利率缺失", 1, 10000, 1, 3000, 1, 5000)
    finally:
        workbook.close()


def _all_zqtz() -> list[FormalZqtzBalanceFactRow]:
    return [_zqtz_asset(), _zqtz_issuance()]


def _all_tyw() -> list[FormalTywBalanceFactRow]:
    return [_tyw_asset(), _tyw_liability()]


def _build(*, position_scope, zqtz_rows, tyw_rows, zqtz_full_rows, tyw_full_rows):
    return build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope=position_scope,
        currency_basis="native",
        zqtz_rows=zqtz_rows,
        tyw_rows=tyw_rows,
        zqtz_full_rows=zqtz_full_rows,
        tyw_full_rows=tyw_full_rows,
    )


def test_asset_scope_cross_scope_tables_match_all_scope():
    reference = build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope="all",
        currency_basis="native",
        zqtz_rows=_all_zqtz(),
        tyw_rows=_all_tyw(),
    )
    scoped = _build(
        position_scope="asset",
        zqtz_rows=[_zqtz_asset()],
        tyw_rows=[_tyw_asset()],
        zqtz_full_rows=_all_zqtz(),
        tyw_full_rows=_all_tyw(),
    )

    for key in CROSS_SCOPE_TABLE_KEYS:
        assert _table(scoped, key)["rows"] == _table(reference, key)["rows"], key
    for section in scoped["tables"]:
        if section["key"] in CROSS_SCOPE_SECTION_KEYS:
            ref_section = next(t for t in reference["tables"] if t["key"] == section["key"])
            assert section["rows"] == ref_section["rows"], section["key"]
    assert scoped["cards"] == reference["cards"]


def test_public_modular_entrypoint_delegates_to_authoritative_full_scope_builder():
    kwargs = {
        "report_date": RD,
        "position_scope": "asset",
        "currency_basis": "native",
        "zqtz_rows": [_zqtz_asset()],
        "tyw_rows": [_tyw_asset()],
        "zqtz_full_rows": _all_zqtz(),
        "tyw_full_rows": _all_tyw(),
    }

    assert build_modular_balance_workbook_payload(**kwargs) == (
        build_balance_analysis_workbook_payload(**kwargs)
    )


def test_asset_scope_maturity_gap_exposes_liability_columns_not_zero():
    scoped = _build(
        position_scope="asset",
        zqtz_rows=[_zqtz_asset()],
        tyw_rows=[_tyw_asset()],
        zqtz_full_rows=_all_zqtz(),
        tyw_full_rows=_all_tyw(),
    )
    gap = _table(scoped, "maturity_gap")
    liability_total = sum(Decimal(str(r["interbank_liabilities_amount"])) for r in gap["rows"])
    issuance_total = sum(Decimal(str(r["issuance_amount"])) for r in gap["rows"])
    assert liability_total == Decimal("5000")  # 50,000,000 元 -> 万元
    assert issuance_total == Decimal("4000")  # 40,000,000 元 -> 万元


def test_missing_maturity_falls_into_shortest_bucket_with_disclosure():
    # B10-1（2026-08 审计）：缺失 maturity_date 的行（典型为活期类同业）不得落
    # "已到期/逾期"桶——与负债分析兼容链（maturity_bucket → "3个月以内"、
    # monthly_v1_bucket_name → "0-3M"）的缺失兜底口径统一，归入"3个月以内"，
    # 并保留 bal_wb_risk_maturity_missing_001 披露。旧锁定（落"已到期/逾期"、
    # "按 0 年处理"）属跨模块矛盾口径，在本测试中被有意修正。
    missing_zqtz = [
        replace(_zqtz_asset(), maturity_date=None),
        replace(_zqtz_issuance(), maturity_date=None),
    ]
    missing_tyw = [
        replace(_tyw_asset(), maturity_date=None),
        replace(_tyw_liability(), maturity_date=None),
        replace(
            _tyw_asset(),
            position_id="T-UNMAPPED",
            position_scope="all",
            maturity_date=None,
        ),
    ]
    payload = build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope="all",
        currency_basis="native",
        zqtz_rows=missing_zqtz,
        tyw_rows=missing_tyw,
    )

    maturity_gap = _table(payload, "maturity_gap")
    shortest = next(row for row in maturity_gap["rows"] if row["bucket"] == "3个月以内")
    assert shortest["bond_assets_amount"] == Decimal("10000")
    assert shortest["issuance_amount"] == Decimal("4000")
    assert shortest["interbank_assets_amount"] == Decimal("3000")
    assert shortest["interbank_liabilities_amount"] == Decimal("5000")

    expired = next(row for row in maturity_gap["rows"] if row["bucket"] == "已到期/逾期")
    assert expired["bond_assets_amount"] == Decimal("0")
    assert expired["issuance_amount"] == Decimal("0")
    assert expired["interbank_assets_amount"] == Decimal("0")
    assert expired["interbank_liabilities_amount"] == Decimal("0")

    duration_proxy = next(
        row
        for row in _table(payload, "regulatory_limits")["rows"]
        if row["metric_key"] == "portfolio_modified_duration"
    )
    assert duration_proxy["current_value"] == Decimal("0")

    bond_business_types = _table(payload, "bond_business_types")["rows"]
    assert bond_business_types[0]["weighted_term_years"] is None
    issuance_business_types = _table(payload, "issuance_business_types")["rows"]
    assert issuance_business_types[0]["weighted_term_years"] is None

    cashflow_calendar = _table(payload, "cashflow_calendar")
    count_fields = (
        "bond_maturity_count",
        "interbank_asset_maturity_count",
        "interbank_liability_maturity_count",
        "issuance_maturity_count",
    )
    assert all(
        sum(row[field] for row in cashflow_calendar["rows"]) == 0
        for field in count_fields
    )
    assert _table(payload, "event_calendar")["rows"] == []

    alerts = _table(payload, "risk_alerts")["rows"]
    alert = alerts[-1]
    BalanceAnalysisRiskAlertRow.model_validate(alert)
    assert alert["rule_id"] == "bal_wb_risk_maturity_missing_001"
    assert alert["severity"] == "medium"
    assert alert["title"] == "到期日缺失口径披露"
    assert "共有 4 条正式事实行缺失 maturity_date" in alert["reason"]
    assert "债券投资资产 1" in alert["reason"]
    assert "发行类负债 1" in alert["reason"]
    assert "同业资产 1" in alert["reason"]
    assert "同业负债 1" in alert["reason"]
    assert "涉及期限分析面值/本金 2.20 亿元" in alert["reason"]
    assert "占该口径余额 100.00%" in alert["reason"]
    assert "四类行在期限缺口中归入「3个月以内」桶" in alert["reason"]
    assert "不落「已到期/逾期」" in alert["reason"]
    assert "工作簿的短期限代理" in alert["reason"]
    assert "负债分析兼容页则单列「到期日未提供」" in alert["reason"]
    assert "两者用途和口径不同" in alert["reason"]
    assert "债券投资资产和同业资产同时按 0 年进入组合剩余期限 proxy" in alert["reason"]
    assert "加权期限及现金流、事件日历剔除缺失值" in alert["reason"]


def test_balance_workbook_compatibility_modules_reuse_authoritative_calculations():
    authority = import_module("backend.app.core_finance.balance_analysis_workbook")

    for module_name, expected_exports in _COMPATIBILITY_CALCULATION_EXPORTS.items():
        compatibility = import_module(
            f"backend.app.core_finance.balance_workbook.{module_name}"
        )
        # Several service-contract tests intentionally reload the authority module.
        # Refresh the compatibility module so this assertion observes the current
        # authoritative module instance rather than an earlier test's import cache.
        compatibility = reload(compatibility)
        assert tuple(compatibility.__all__) == expected_exports
        locally_defined_functions = {
            name
            for name, value in vars(compatibility).items()
            if isfunction(value) and value.__module__ == compatibility.__name__
        }
        assert locally_defined_functions == set()
        for symbol_name in expected_exports:
            assert getattr(compatibility, symbol_name) is getattr(
                authority,
                symbol_name,
            ), f"{module_name}.{symbol_name} must reuse the authoritative calculation"


def test_missing_maturity_caliber_package_copies_match_monolith():
    # B10-1 兼容路径钉住：历史私有 import 必须复用单体权威实现，并保持
    # 缺失 maturity_date 归入"3个月以内"桶、输出治理披露的业务语义。
    from backend.app.core_finance.balance_analysis_workbook import (
        _build_maturity_gap_table as monolith_maturity_gap,
    )
    from backend.app.core_finance.balance_analysis_workbook import (
        _build_risk_alerts_table as monolith_risk_alerts,
    )
    from backend.app.core_finance.balance_workbook._bond_tables import (
        _build_maturity_gap_table as package_maturity_gap,
    )
    from backend.app.core_finance.balance_workbook._risk_tables import (
        _build_risk_alerts_table as package_risk_alerts,
    )

    missing_zqtz = [
        replace(_zqtz_asset(), maturity_date=None),
        replace(_zqtz_issuance(), maturity_date=None),
    ]
    missing_tyw = [
        replace(_tyw_asset(), maturity_date=None),
        replace(_tyw_liability(), maturity_date=None),
    ]

    monolith_gap_rows = monolith_maturity_gap(RD, missing_zqtz, missing_tyw)["rows"]
    package_gap_rows = package_maturity_gap(RD, missing_zqtz, missing_tyw)["rows"]
    assert package_gap_rows == monolith_gap_rows
    shortest = next(row for row in package_gap_rows if row["bucket"] == "3个月以内")
    assert shortest["bond_assets_amount"] == Decimal("10000")
    expired = next(row for row in package_gap_rows if row["bucket"] == "已到期/逾期")
    assert expired["bond_assets_amount"] == Decimal("0")

    monolith_alert_rows = monolith_risk_alerts(RD, missing_zqtz, missing_tyw)["rows"]
    package_alert_rows = package_risk_alerts(RD, missing_zqtz, missing_tyw)["rows"]
    assert package_alert_rows == monolith_alert_rows
    assert any(
        row["rule_id"] == "bal_wb_risk_maturity_missing_001" for row in package_alert_rows
    )


def test_complete_maturity_dates_do_not_emit_missing_maturity_alert():
    payload = build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope="all",
        currency_basis="native",
        zqtz_rows=_all_zqtz(),
        tyw_rows=_all_tyw(),
    )

    alerts = _table(payload, "risk_alerts")["rows"]
    assert all(
        row["rule_id"] != "bal_wb_risk_maturity_missing_001"
        for row in alerts
    )


def test_asset_scope_single_scope_liability_table_stays_scoped_empty():
    scoped = _build(
        position_scope="asset",
        zqtz_rows=[_zqtz_asset()],
        tyw_rows=[_tyw_asset()],
        zqtz_full_rows=_all_zqtz(),
        tyw_full_rows=_all_tyw(),
    )
    assert _table(scoped, "issuance_business_types")["rows"] == []


def test_service_fetches_full_scope_rows_and_wires_them_under_asset_scope():
    from tests.helpers import load_module

    workbook_service = load_module(
        "backend.app.services.balance_analysis_workbook_service",
        "backend/app/services/balance_analysis_workbook_service.py",
    )

    def _native_zqtz(scope: str) -> dict:
        return {
            "report_date": "2026-03-31",
            "instrument_code": "A0001" if scope != "liability" else "L0001",
            "instrument_name": "债券",
            "portfolio_name": "P",
            "cost_center": "C",
            "account_category": "可供出售类资产",
            "asset_class": "信用债",
            "bond_type": "企业债",
            "sub_type": "",
            "business_type_primary": "",
            "issuer_name": "发行人A",
            "industry_name": "工业",
            "rating": "AAA",
            "invest_type_std": "A",
            "accounting_basis": "FVOCI",
            "position_scope": scope,
            "currency_basis": "native",
            "currency_code": "CNY",
            "face_value_amount": "100.00000000",
            "market_value_amount": "100.00000000",
            "amortized_cost_amount": "90.00000000",
            "accrued_interest_amount": "0.00000000",
            "coupon_rate": "3.00000000",
            "ytm_value": "3.10000000",
            "maturity_date": "2027-03-31",
            "interest_mode": "固定",
            "is_issuance_like": scope == "liability",
            "overdue_principal_days": 0,
            "overdue_interest_days": 0,
            "value_date": "2026-03-31",
            "customer_attribute": "internal",
            "source_version": "sv",
            "rule_version": "rv",
            "ingest_batch_id": "ib",
            "trace_id": "tr",
        }

    calls: list[tuple[str, str, str]] = []

    class FakeRepo:
        def __init__(self, duckdb_path: str) -> None:
            self.duckdb_path = duckdb_path

        def list_report_dates(self):
            return ["2026-03-31"]

        def fetch_formal_zqtz_rows(self, **kwargs):
            calls.append(("zqtz", kwargs["position_scope"], kwargs["currency_basis"]))
            basis = kwargs["currency_basis"]
            if kwargs["position_scope"] == "all":
                rows = [_native_zqtz("asset"), _native_zqtz("liability")]
            else:
                rows = [_native_zqtz(kwargs["position_scope"])]
            # Echo requested basis so callers can assert currency_basis honor.
            for row in rows:
                row["currency_basis"] = basis
            return rows

        def fetch_formal_tyw_rows(self, **kwargs):
            calls.append(("tyw", kwargs["position_scope"], kwargs["currency_basis"]))
            return []

    captured: dict[str, object] = {}

    class FakeWorkbookModule:
        @staticmethod
        def build_balance_analysis_workbook_payload(**kwargs):
            captured.update(kwargs)
            return {
                "report_date": str(kwargs["report_date"]),
                "position_scope": kwargs["position_scope"],
                "currency_basis": kwargs["currency_basis"],
                "cards": [],
                "tables": [],
            }

    workbook_service._build_balance_workbook_payload(
        duckdb_path="ignored.duckdb",
        governance_dir="ignored-governance",
        report_date="2026-03-31",
        position_scope="asset",
        currency_basis="CNY",
        cache_key="ck",
        job_name="job",
        resolve_completed_formal_build_lineage_fn=lambda **_kwargs: None,
        repo_cls=FakeRepo,
        import_module_fn=lambda _name: FakeWorkbookModule,
        reload_module_fn=lambda module: module,
    )

    # H-2: requested CNY must drive main/full-row fetches (not hardcoded native).
    assert ("zqtz", "asset", "CNY") in calls
    assert ("zqtz", "all", "CNY") in calls
    assert ("tyw", "all", "CNY") in calls
    assert ("zqtz", "all", "native") not in calls
    # Full rows handed to the builder include both asset and liability zqtz rows.
    full_scopes = {row.position_scope for row in captured["zqtz_full_rows"]}
    assert full_scopes == {"asset", "liability"}
    # Scoped rows stay asset-only and keep CNY basis.
    assert {row.position_scope for row in captured["zqtz_rows"]} == {"asset"}
    assert {row.currency_basis for row in captured["zqtz_rows"]} == {"CNY"}
    assert captured["currency_basis"] == "CNY"


def test_liability_scope_cross_scope_tables_match_all_scope():
    reference = build_balance_analysis_workbook_payload(
        report_date=RD,
        position_scope="all",
        currency_basis="native",
        zqtz_rows=_all_zqtz(),
        tyw_rows=_all_tyw(),
    )
    scoped = _build(
        position_scope="liability",
        zqtz_rows=[_zqtz_issuance()],
        tyw_rows=[_tyw_liability()],
        zqtz_full_rows=_all_zqtz(),
        tyw_full_rows=_all_tyw(),
    )
    for key in CROSS_SCOPE_TABLE_KEYS:
        assert _table(scoped, key)["rows"] == _table(reference, key)["rows"], key
    assert scoped["cards"] == reference["cards"]
    assert _table(scoped, "bond_business_types")["rows"] == []
