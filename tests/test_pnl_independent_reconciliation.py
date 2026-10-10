from __future__ import annotations

import json
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import duckdb
from openpyxl import Workbook, load_workbook
import pytest

from backend.app.core_finance.pnl_independent_reconciliation import (
    DailyBalanceObservation,
    IndependentPnlReference,
    calculate_complete_daily_average_cny,
    reconcile_pnl_overview_to_independent_reference,
)
from backend.app.repositories.governance_repo import GovernanceRepository
from backend.app.repositories.pnl_independent_reference_repo import (
    PNL_INDEPENDENT_REFERENCE_STREAM,
    PnlFormalFactDependencyBinding,
    read_formal_fact_dependency_binding,
    read_formal_fact_dependency_revision,
)
from backend.app.repositories.pnl_repo import PnlRepository
from backend.app.services.pnl_independent_reference_service import (
    LedgerPnlComponentScope,
    build_independent_ledger_pnl_reference,
    load_and_reconcile_prepared_independent_pnl_reference,
    load_prepared_independent_pnl_reference,
    pending_component_reference,
    prepare_and_store_independent_ledger_pnl_reference,
    ready_adjustment_reference,
)
from backend.app.services.pnl_service import pnl_overview_envelope
from backend.app.services.product_category_source_service import discover_source_pairs
from backend.app.tasks.pnl_materialize import materialize_pnl_facts


REPORT_DATES = ("2026-07-31", "2026-08-31")
MONTH_BY_DATE = {"2026-07-31": "202607", "2026-08-31": "202608"}
LEDGER_COMPONENT_VALUES = {
    "202607": {
        "CNX": {"514": Decimal("29"), "516": Decimal("6"), "517": Decimal("190")},
        "CNY": {"514": Decimal("15"), "516": Decimal("6"), "517": Decimal("120")},
    },
    "202608": {
        "CNX": {"514": Decimal("38"), "516": Decimal("7"), "517": Decimal("215")},
        "CNY": {"514": Decimal("17"), "516": Decimal("7"), "517": Decimal("131")},
    },
}
MANUAL_ADJUSTMENTS = {"202607": Decimal("0"), "202608": Decimal("2")}
ADJUSTMENT_VERSIONS = {
    "202607": "av-synthetic-202607-empty",
    "202608": "av-synthetic-202608-approved-v1",
}


def _fresh_replay_root(tmp_path: Path, name: str) -> Path:
    root = tmp_path / name
    root.mkdir()
    return root


def _write_ledger_workbook(
    source_dir: Path,
    month: str,
    *,
    component_values: dict[str, dict[str, Decimal]] | None = None,
) -> Path:
    path = source_dir / f"总账对账{month}.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    year, number = int(month[:4]), int(month[4:])
    month_end = date.fromisoformat(f"{year:04d}-{number:02d}-01")
    if number == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, number + 1, 1)
    end = date.fromordinal(next_month.toordinal() - 1)
    values = component_values or LEDGER_COMPONENT_VALUES[month]
    for sheet_name, currency in (("综本", "CNX"), ("人民币", "CNY")):
        sheet = workbook.create_sheet(sheet_name)
        sheet["A2"] = (
            f"会计期间： {month_end.isoformat()}--{end.isoformat()} 币种： {currency}"
        )
        for row_number, component in enumerate(("514", "516", "517"), start=7):
            credit = values[currency][component]
            sheet.cell(row=row_number, column=1, value=f"{component}90000001")
            sheet.cell(row=row_number, column=2, value=f"合成-{component}")
            sheet.cell(row=row_number, column=3, value=currency)
            sheet.cell(row=row_number, column=4, value=Decimal("0"))
            sheet.cell(row=row_number, column=5, value=Decimal("0"))
            sheet.cell(row=row_number, column=6, value=credit)
            sheet.cell(row=row_number, column=7, value=-credit)
    workbook.save(path)
    workbook.close()
    return path


def _write_average_placeholder(source_dir: Path, month: str) -> None:
    workbook = Workbook()
    workbook.active.title = "年日均"
    workbook.create_sheet("月日均")
    workbook.save(source_dir / f"日均{month}.xlsx")
    workbook.close()


def _write_source_workbooks(source_dir: Path) -> None:
    source_dir.mkdir(parents=True)
    for month in ("202607", "202608"):
        _write_ledger_workbook(source_dir, month)
        _write_average_placeholder(source_dir, month)


def _seed_fx(duckdb_path: Path) -> None:
    with duckdb.connect(str(duckdb_path)) as connection:
        connection.execute(
            """
            create table fx_daily_mid (
              trade_date varchar,
              base_currency varchar,
              quote_currency varchar,
              mid_rate decimal(24, 8),
              source_name varchar,
              is_business_day boolean,
              is_carry_forward boolean,
              source_version varchar,
              observed_trade_date varchar
            )
            """
        )
        connection.executemany(
            "insert into fx_daily_mid values (?, 'USD', 'CNY', 7, 'CFETS', true, false, ?, ?)",
            [(report_date, f"sv-fx-{report_date}", report_date) for report_date in REPORT_DATES],
        )


def _fi_rows(
    report_date: str,
    *,
    h_interest: str | None = None,
    include_a: bool = True,
    adjustment_status: str = "approved",
) -> list[dict[str, object]]:
    if report_date == "2026-07-31":
        values = {
            "H": ("10", "-106", "0"),
            "A": ("2", "-10.6", "0"),
            "T": ("5", "-21.2", "6"),
        }
    else:
        values = {
            "H": ("11", "-116.6", "0"),
            "A": ("3", "-12.72", "0"),
            "T": ("6", "-22.26", "7"),
        }
    if h_interest is not None:
        values["H"] = (h_interest, values["H"][1], values["H"][2])
    rows: list[dict[str, object]] = []
    for invest_type, (interest, capital_gain, fair_value) in values.items():
        if invest_type == "A" and not include_a:
            continue
        row: dict[str, object] = {
            "report_date": report_date,
            "instrument_code": f"SYN-{invest_type}-{report_date}",
            "portfolio_name": "Synthetic FI",
            "cost_center": "SYN-CC",
            "invest_type_raw": invest_type,
            "asset_class": "合成企业债",
            "interest_income_514": interest,
            "fair_value_change_516": fair_value,
            "capital_gain_517": capital_gain,
            "event_type": "fi_cumulative_realized_517",
            "currency_basis": "CNY",
            "source_version": f"sv-fi-{report_date}",
            "ingest_batch_id": f"batch-{report_date}",
            "trace_id": f"trace-{invest_type}-{report_date}",
        }
        if invest_type == "A":
            row["fx_base_currency"] = "USD"
        if report_date == "2026-08-31" and invest_type == "H":
            row["manual_adjustment"] = "2"
            row["approval_status"] = adjustment_status
        rows.append(row)
    return rows


def _materialize_two_months(
    duckdb_path: Path, governance_dir: Path
) -> dict[str, dict[str, object]]:
    receipts: dict[str, dict[str, object]] = {}
    for report_date in REPORT_DATES:
        receipts[report_date] = materialize_pnl_facts.fn(
            report_date=report_date,
            is_month_end=True,
            fi_rows=_fi_rows(report_date),
            nonstd_rows_by_type={},
            duckdb_path=str(duckdb_path),
            governance_dir=str(governance_dir),
            formal_pnl_enabled=True,
            formal_pnl_scope_json='["*"]',
        )
    return receipts


def _reference_for_month(
    source_dir: Path,
    month: str,
    *,
    governance_dir: Path | None = None,
    manual_adjustment: Decimal | None = None,
    formal_source_version: str = "sv-formal-synthetic",
    formal_rule_version: str = "rv-formal-synthetic",
    approved_adjustment_version: str | None = None,
    formal_dependency_binding: PnlFormalFactDependencyBinding | None = None,
) -> IndependentPnlReference:
    pair = next(item for item in discover_source_pairs(source_dir) if item.month_key == month)
    workbook_ref = pair.ledger_path.name
    scopes = {
        "interest_income_514": LedgerPnlComponentScope(
            status="ready",
            account_prefixes=("514",),
            evidence_refs=(f"{workbook_ref}#综本!F7", "hat-event-scope://synthetic-fi-v1"),
        ),
        "fair_value_change_516": LedgerPnlComponentScope(
            status="ready",
            account_prefixes=("516",),
            evidence_refs=(f"{workbook_ref}#综本!F8", "hat-event-scope://synthetic-fi-v1"),
        ),
        "capital_gain_517": LedgerPnlComponentScope(
            status="ready",
            account_prefixes=("517",),
            evidence_refs=(f"{workbook_ref}#综本!F9", "realized-event-scope://synthetic-fi-v1"),
        ),
    }
    binding = formal_dependency_binding or PnlFormalFactDependencyBinding(
        report_date=pair.report_date.isoformat(),
        formal_source_version=formal_source_version,
        formal_rule_version=formal_rule_version,
        approved_adjustment_version=(
            approved_adjustment_version or ADJUSTMENT_VERSIONS[month]
        ),
        formal_dependency_revision=1,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
        formal_fi_row_count=1,
        nonstd_bridge_row_count=0,
    )
    kwargs = {
        "component_scopes": scopes,
        "manual_adjustment": ready_adjustment_reference(
            MANUAL_ADJUSTMENTS[month] if manual_adjustment is None else manual_adjustment,
            evidence_refs=(f"approval-register://synthetic/{month}",),
        ),
        "formal_dependency_binding": binding,
    }
    if governance_dir is None:
        return build_independent_ledger_pnl_reference(pair, **kwargs)
    return prepare_and_store_independent_ledger_pnl_reference(
        pair,
        governance_dir=governance_dir,
        **kwargs,
    )


def test_reference_missing_scope_stays_pending_instead_of_zero() -> None:
    result = reconcile_pnl_overview_to_independent_reference(
        {
            "interest_income_514": Decimal("1"),
            "fair_value_change_516": Decimal("2"),
            "capital_gain_517": Decimal("3"),
            "manual_adjustment": Decimal("0"),
            "total_pnl": Decimal("6"),
        },
        None,
        report_date="2026-08-31",
        formal_dependency_revision=1,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
    )
    assert result["status"] == "pending"
    assert result["breached"] is None
    assert result["reference_total_yuan"] is None


def test_pending_516_scope_blocks_total_but_preserves_ready_component_checks() -> None:
    source = IndependentPnlReference(
        report_date="2026-08-31",
        currency_basis="CNY",
        ledger_currency_code="CNX",
        source_version="sv-synthetic",
        rule_version="rv-synthetic",
        formal_source_version="sv-formal",
        formal_rule_version="rv-formal",
        approved_adjustment_version="av-approved",
        formal_dependency_revision=1,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
        components={
            "interest_income_514": ready_adjustment_reference(
                Decimal("1"), evidence_refs=("ledger://514",)
            ),
            "fair_value_change_516": pending_component_reference(
                "ledger account code does not prove H/A/T eligibility"
            ),
            "capital_gain_517": ready_adjustment_reference(
                Decimal("3"), evidence_refs=("ledger://517",)
            ),
            "manual_adjustment": ready_adjustment_reference(
                Decimal("0"), evidence_refs=("approval://empty",)
            ),
        },
    )
    result = reconcile_pnl_overview_to_independent_reference(
        {
            "interest_income_514": Decimal("1"),
            "fair_value_change_516": Decimal("2"),
            "capital_gain_517": Decimal("3"),
            "manual_adjustment": Decimal("0"),
            "total_pnl": Decimal("6"),
        },
        source,
        report_date="2026-08-31",
        formal_dependency_revision=1,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
    )
    assert result["status"] == "pending"
    assert result["pending_components"] == ["fair_value_change_516"]
    assert result["reference_total_yuan"] is None
    assert [row["status"] for row in result["components"]] == [
        "pass",
        "pending",
        "pass",
        "pass",
    ]

    stale = reconcile_pnl_overview_to_independent_reference(
        {
            "interest_income_514": Decimal("1"),
            "fair_value_change_516": Decimal("2"),
            "capital_gain_517": Decimal("3"),
            "manual_adjustment": Decimal("0"),
            "total_pnl": Decimal("6"),
        },
        source,
        report_date="2026-08-31",
        formal_dependency_revision=2,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
    )
    assert stale["status"] == "pending"
    assert "formal_dependency_revision" in str(stale["reason"])


def test_zero_dependency_revision_is_valid_for_a_current_protocol_reference() -> None:
    ready_component = lambda value, ref: ready_adjustment_reference(  # noqa: E731
        Decimal(value), evidence_refs=(ref,)
    )
    reference = IndependentPnlReference(
        report_date="2026-08-31",
        currency_basis="CNY",
        ledger_currency_code="CNX",
        source_version="sv-ledger",
        rule_version="rv-ledger",
        formal_source_version="sv-formal",
        formal_rule_version="rv-formal",
        approved_adjustment_version="av-fact-materialized",
        formal_dependency_revision=0,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
        components={
            "interest_income_514": ready_component("1", "ledger://514"),
            "fair_value_change_516": ready_component("2", "ledger://516"),
            "capital_gain_517": ready_component("3", "ledger://517"),
            "manual_adjustment": ready_component("0", "approval://empty"),
        },
    )
    result = reconcile_pnl_overview_to_independent_reference(
        {
            "interest_income_514": Decimal("1"),
            "fair_value_change_516": Decimal("2"),
            "capital_gain_517": Decimal("3"),
            "manual_adjustment": Decimal("0"),
            "total_pnl": Decimal("6"),
        },
        reference,
        report_date="2026-08-31",
        formal_dependency_revision=0,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
    )
    assert result["status"] == "pass"

    at_tolerance = reconcile_pnl_overview_to_independent_reference(
        {
            "interest_income_514": Decimal("1.01"),
            "fair_value_change_516": Decimal("2"),
            "capital_gain_517": Decimal("3"),
            "manual_adjustment": Decimal("0"),
            "total_pnl": Decimal("6.01"),
        },
        reference,
        report_date="2026-08-31",
        formal_dependency_revision=0,
        formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        adjustment_stage="fact_materialized",
    )
    assert at_tolerance["status"] == "pass"


def test_corrupt_latest_prepared_reference_raises_instead_of_falling_back(
    tmp_path: Path,
) -> None:
    root = _fresh_replay_root(tmp_path, "corrupt-prepared-reference")
    source_dir = root / "source"
    governance_dir = root / "governance"
    _write_source_workbooks(source_dir)
    reference = _reference_for_month(
        source_dir,
        "202608",
        governance_dir=governance_dir,
    )
    corrupt = reference.to_record()
    corrupt["formal_dependency_revision"] = "not-an-integer"
    GovernanceRepository(base_dir=governance_dir).append(
        PNL_INDEPENDENT_REFERENCE_STREAM,
        corrupt,
    )
    with pytest.raises(RuntimeError, match="Invalid prepared independent PnL reference"):
        load_prepared_independent_pnl_reference(
            governance_dir=governance_dir,
            report_date="2026-08-31",
            formal_dependency_revision=1,
            formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
        )


def test_empty_ledger_match_requires_dedicated_zero_observation_evidence(
    tmp_path: Path,
) -> None:
    root = _fresh_replay_root(tmp_path, "empty-ledger-scope")
    source_dir = root / "source"
    _write_source_workbooks(source_dir)
    pair = next(item for item in discover_source_pairs(source_dir) if item.month_key == "202608")
    base_scopes = {
        "interest_income_514": LedgerPnlComponentScope(
            status="ready",
            account_prefixes=("999",),
            evidence_refs=("scope-contract://synthetic/999",),
        ),
        "fair_value_change_516": LedgerPnlComponentScope(
            status="pending",
            reason="not under test",
        ),
        "capital_gain_517": LedgerPnlComponentScope(
            status="pending",
            reason="not under test",
        ),
    }
    reference = build_independent_ledger_pnl_reference(
        pair,
        component_scopes=base_scopes,
        manual_adjustment=ready_adjustment_reference(
            Decimal("0"), evidence_refs=("approval-register://synthetic/empty",)
        ),
        formal_dependency_binding=PnlFormalFactDependencyBinding(
            report_date="2026-08-31",
            formal_source_version="sv-formal",
            formal_rule_version="rv-formal",
            approved_adjustment_version="av-empty",
            formal_dependency_revision=1,
            formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
            adjustment_stage="fact_materialized",
            formal_fi_row_count=1,
            nonstd_bridge_row_count=0,
        ),
    )
    assert reference.components["interest_income_514"].status == "pending"
    assert reference.components["interest_income_514"].value_yuan is None

    explicit_zero_scopes = dict(base_scopes)
    explicit_zero_scopes["interest_income_514"] = LedgerPnlComponentScope(
        status="ready",
        account_prefixes=("999",),
        evidence_refs=("scope-contract://synthetic/999",),
        zero_observation_evidence_refs=(
            "ledger-full-population://synthetic/202608/999-empty",
        ),
    )
    zero_reference = build_independent_ledger_pnl_reference(
        pair,
        component_scopes=explicit_zero_scopes,
        manual_adjustment=ready_adjustment_reference(
            Decimal("0"), evidence_refs=("approval-register://synthetic/empty",)
        ),
        formal_dependency_binding=PnlFormalFactDependencyBinding(
            report_date="2026-08-31",
            formal_source_version="sv-formal",
            formal_rule_version="rv-formal",
            approved_adjustment_version="av-empty",
            formal_dependency_revision=1,
            formal_dependency_protocol_version="pnl_by_business_precompute_state_v1",
            adjustment_stage="fact_materialized",
            formal_fi_row_count=1,
            nonstd_bridge_row_count=0,
        ),
    )
    zero_component = zero_reference.components["interest_income_514"]
    assert zero_component.status == "ready"
    assert zero_component.value_yuan == Decimal("0")
    assert "ledger-full-population://synthetic/202608/999-empty" in zero_component.evidence_refs


def test_daily_average_is_hand_calculated_from_complete_cny_and_usd_calendar_days() -> None:
    start = date(2026, 7, 1)
    observations: list[DailyBalanceObservation] = []
    for offset in range(31):
        report_date = start + timedelta(days=offset)
        observations.extend(
            [
                DailyBalanceObservation(
                    report_date=report_date,
                    observation_id="CNY-AC",
                    amount_native=Decimal("100"),
                    fx_mid_rate=Decimal("1"),
                    evidence_refs=(f"daily-balance://{report_date}/CNY-AC",),
                ),
                DailyBalanceObservation(
                    report_date=report_date,
                    observation_id="USD-FVTPL",
                    amount_native=Decimal("20" if offset == 30 else "10"),
                    fx_mid_rate=Decimal("7"),
                    evidence_refs=(
                        f"daily-balance://{report_date}/USD-FVTPL",
                        f"fx-mid://{report_date}/USD-CNY",
                    ),
                ),
            ]
        )
    result = calculate_complete_daily_average_cny(
        observations,
        period_start=start,
        period_end=date(2026, 7, 31),
    )
    hand_calculated = (Decimal("170") * Decimal("30") + Decimal("240")) / Decimal("31")
    assert result["status"] == "ready"
    assert result["calendar_days"] == 31
    assert result["observation_count"] == 62
    assert result["average_balance_cny"] == hand_calculated

    missing = calculate_complete_daily_average_cny(
        observations[:-2],
        period_start=start,
        period_end=date(2026, 7, 31),
    )
    assert missing["status"] == "pending"
    assert missing["average_balance_cny"] is None
    assert missing["missing_dates"] == ["2026-07-31"]


def test_two_month_real_materialization_roundtrip_and_single_ledger_cell_mismatch(
    tmp_path: Path,
) -> None:
    root = _fresh_replay_root(tmp_path, "baseline-and-mismatch")
    source_dir = root / "source"
    duckdb_path = root / "moss.duckdb"
    governance_dir = root / "governance"
    _write_source_workbooks(source_dir)
    _seed_fx(duckdb_path)
    assert (
        read_formal_fact_dependency_revision(
            duckdb_path,
            report_date="2026-07-31",
        )
        is None
    )
    _materialize_two_months(duckdb_path, governance_dir)

    repository = PnlRepository(str(duckdb_path))
    receipts: dict[str, object] = {"report_dates": list(REPORT_DATES), "checks": {}}
    for report_date in REPORT_DATES:
        month = MONTH_BY_DATE[report_date]
        formal_binding = read_formal_fact_dependency_binding(
            duckdb_path,
            report_date=report_date,
        )
        assert formal_binding is not None
        assert read_formal_fact_dependency_revision(
            duckdb_path,
            report_date=report_date,
        ) == (
            formal_binding.formal_dependency_revision,
            formal_binding.formal_dependency_protocol_version,
        )
        _reference_for_month(
            source_dir,
            month,
            governance_dir=governance_dir,
            formal_dependency_binding=formal_binding,
        )
        loaded = load_prepared_independent_pnl_reference(
            governance_dir=governance_dir,
            report_date=report_date,
            formal_dependency_revision=formal_binding.formal_dependency_revision,
            formal_dependency_protocol_version=(
                formal_binding.formal_dependency_protocol_version
            ),
        )
        assert loaded is not None
        assert (
            load_prepared_independent_pnl_reference(
                governance_dir=governance_dir,
                report_date=report_date,
                formal_dependency_revision=formal_binding.formal_dependency_revision + 1,
                formal_dependency_protocol_version=(
                    formal_binding.formal_dependency_protocol_version
                ),
            )
            is None
        )
        totals = pnl_overview_envelope(
            duckdb_path=str(duckdb_path),
            governance_dir=str(governance_dir),
            report_date=report_date,
        )["result"]
        result = load_and_reconcile_prepared_independent_pnl_reference(
            totals,
            governance_dir=governance_dir,
            report_date=report_date,
            formal_dependency_revision=formal_binding.formal_dependency_revision,
            formal_dependency_protocol_version=(
                formal_binding.formal_dependency_protocol_version
            ),
        )
        assert result["status"] == "pass"
        assert Decimal(str(result["total_diff_yuan"])) == Decimal("0")
        receipts["checks"][month] = result

    july_book = load_workbook(source_dir / "总账对账202607.xlsx", data_only=True)
    try:
        assert Decimal(str(july_book["综本"]["F7"].value)) == Decimal("29")
        assert Decimal(str(july_book["人民币"]["F7"].value)) == Decimal("15")
    finally:
        july_book.close()

    august_path = source_dir / "总账对账202608.xlsx"
    august_book = load_workbook(august_path)
    try:
        august_book["综本"]["F7"] = Decimal("39")
        august_book["综本"]["G7"] = Decimal("-39")
        august_book.save(august_path)
    finally:
        august_book.close()
    august_binding = read_formal_fact_dependency_binding(
        duckdb_path,
        report_date="2026-08-31",
    )
    assert august_binding is not None
    _reference_for_month(
        source_dir,
        "202608",
        governance_dir=governance_dir,
        formal_dependency_binding=august_binding,
    )
    mismatch = load_and_reconcile_prepared_independent_pnl_reference(
        repository.overview_totals("2026-08-31"),
        governance_dir=governance_dir,
        report_date="2026-08-31",
        formal_dependency_revision=august_binding.formal_dependency_revision,
        formal_dependency_protocol_version=(
            august_binding.formal_dependency_protocol_version
        ),
    )
    assert mismatch["status"] == "fail"
    assert mismatch["total_diff_yuan"] == "-1.00000000"
    assert mismatch["components"][0]["status"] == "fail"
    receipts["single_cell_mismatch"] = mismatch
    (root / "replay-receipt.json").write_text(
        json.dumps(receipts, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _formal_rows(duckdb_path: Path) -> dict[str, list[tuple]]:
    with duckdb.connect(str(duckdb_path), read_only=True) as connection:
        return {
            table: connection.execute(f"select * from {table} order by all").fetchall()
            for table in ("fact_formal_pnl_fi", "fact_nonstd_pnl_bridge")
        }


def test_historical_correction_delete_repeat_and_revoke_match_clean_full_replay(
    tmp_path: Path,
) -> None:
    root = _fresh_replay_root(tmp_path, "incremental-full")
    incremental_db = root / "incremental.duckdb"
    clean_db = root / "clean-full.duckdb"
    incremental_governance = root / "incremental-governance"
    clean_governance = root / "clean-governance"
    _seed_fx(incremental_db)
    _seed_fx(clean_db)
    _materialize_two_months(incremental_db, incremental_governance)
    initial_bindings = {
        report_date: read_formal_fact_dependency_binding(
            incremental_db,
            report_date=report_date,
        )
        for report_date in REPORT_DATES
    }
    assert all(binding is not None for binding in initial_bindings.values())

    before_repeat = _formal_rows(incremental_db)
    materialize_pnl_facts.fn(
        report_date="2026-08-31",
        is_month_end=True,
        fi_rows=_fi_rows("2026-08-31"),
        nonstd_rows_by_type={},
        duckdb_path=str(incremental_db),
        governance_dir=str(incremental_governance),
        formal_pnl_enabled=True,
        formal_pnl_scope_json='["*"]',
    )
    assert _formal_rows(incremental_db) == before_repeat
    repeated_august_binding = read_formal_fact_dependency_binding(
        incremental_db,
        report_date="2026-08-31",
    )
    assert repeated_august_binding is not None
    initial_august_binding = initial_bindings["2026-08-31"]
    assert initial_august_binding is not None
    assert (
        repeated_august_binding.formal_source_version
        == initial_august_binding.formal_source_version
    )
    assert (
        repeated_august_binding.formal_dependency_revision
        > initial_august_binding.formal_dependency_revision
    )

    final_rows = {
        "2026-07-31": _fi_rows(
            "2026-07-31",
            h_interest="12",
            include_a=False,
        ),
        "2026-08-31": _fi_rows(
            "2026-08-31",
            adjustment_status="rejected",
        ),
    }
    final_incremental_receipts: dict[str, dict[str, object]] = {}
    for report_date in REPORT_DATES:
        final_incremental_receipts[report_date] = materialize_pnl_facts.fn(
            report_date=report_date,
            is_month_end=True,
            fi_rows=final_rows[report_date],
            nonstd_rows_by_type={},
            duckdb_path=str(incremental_db),
            governance_dir=str(incremental_governance),
            formal_pnl_enabled=True,
            formal_pnl_scope_json='["*"]',
        )
        materialize_pnl_facts.fn(
            report_date=report_date,
            is_month_end=True,
            fi_rows=final_rows[report_date],
            nonstd_rows_by_type={},
            duckdb_path=str(clean_db),
            governance_dir=str(clean_governance),
            formal_pnl_enabled=True,
            formal_pnl_scope_json='["*"]',
        )

    assert _formal_rows(incremental_db) == _formal_rows(clean_db)
    with duckdb.connect(str(incremental_db), read_only=True) as connection:
        assert connection.execute(
            "select count(*) from fact_formal_pnl_fi where report_date='2026-07-31'"
        ).fetchone()[0] == 2
        assert connection.execute(
            "select manual_adjustment from fact_formal_pnl_fi "
            "where report_date='2026-08-31' and invest_type_std='H'"
        ).fetchone()[0] == Decimal("0")

    source_dir = root / "final-ledger-source"
    source_dir.mkdir()
    final_ledger_values = {
        "202607": {
            "CNX": {"514": Decimal("17"), "516": Decimal("6"), "517": Decimal("120")},
            "CNY": {"514": Decimal("17"), "516": Decimal("6"), "517": Decimal("120")},
        },
        "202608": LEDGER_COMPONENT_VALUES["202608"],
    }
    for month in ("202607", "202608"):
        _write_ledger_workbook(
            source_dir,
            month,
            component_values=final_ledger_values[month],
        )
        _write_average_placeholder(source_dir, month)

    repository = PnlRepository(str(incremental_db))
    receipt: dict[str, object] = {"incremental_equals_clean_full": True, "checks": {}}
    for report_date in REPORT_DATES:
        month = MONTH_BY_DATE[report_date]
        formal_binding = read_formal_fact_dependency_binding(
            incremental_db,
            report_date=report_date,
        )
        assert formal_binding is not None
        initial_binding = initial_bindings[report_date]
        assert initial_binding is not None
        assert formal_binding.formal_source_version != initial_binding.formal_source_version
        result = reconcile_pnl_overview_to_independent_reference(
            repository.overview_totals(report_date),
            _reference_for_month(
                source_dir,
                month,
                manual_adjustment=Decimal("0"),
                formal_dependency_binding=formal_binding,
            ),
            report_date=report_date,
            formal_dependency_revision=formal_binding.formal_dependency_revision,
            formal_dependency_protocol_version=(
                formal_binding.formal_dependency_protocol_version
            ),
            adjustment_stage=formal_binding.adjustment_stage,
        )
        assert result["status"] == "pass"
        receipt["checks"][month] = result
    (root / "replay-receipt.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
