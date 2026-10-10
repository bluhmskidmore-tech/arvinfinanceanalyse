from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import duckdb
from openpyxl import Workbook, load_workbook
import pytest

from backend.app.core_finance.bond_analytics.dv01 import (
    build_dv01_shock_scenario_payloads,
)
from backend.app.core_finance.pnl_independent_reconciliation import IndependentPnlReference
from backend.app.repositories.pnl_independent_reference_repo import (
    PnlFormalFactDependencyBinding,
    read_formal_fact_dependency_binding,
)
from backend.app.repositories.pnl_repo import PnlRepository
from backend.app.repositories.risk_tensor_repo import RiskTensorRepository
from backend.app.repositories.snapshot_repo import ensure_snapshot_tables
from backend.app.repositories.task_write_guard import repository_task_write_scope
from backend.app.repositories.yield_curve_repo import YieldCurveRepository
from backend.app.schemas.yield_curve import YieldCurvePoint, YieldCurveSnapshot
from backend.app.services.pnl_independent_reference_service import (
    LedgerPnlComponentScope,
    load_and_reconcile_prepared_independent_pnl_reference,
    prepare_and_store_independent_ledger_pnl_reference,
    ready_adjustment_reference,
)
from backend.app.services.product_category_source_service import discover_source_pairs
from backend.app.tasks.balance_analysis_materialize import materialize_balance_analysis_facts
from backend.app.tasks.bond_analytics_materialize import materialize_bond_analytics_facts
from backend.app.tasks.pnl_materialize import materialize_pnl_facts
from backend.app.tasks.risk_tensor_materialize import materialize_risk_tensor_facts
from backend.app.tasks.yield_curve_materialize import RULE_VERSION as YIELD_CURVE_RULE_VERSION


pytestmark = [pytest.mark.integration, pytest.mark.materialize]

REPORT_DATES = ("2026-03-31", "2026-04-30")
FX_BY_DATE = {
    "2026-03-31": Decimal("7.0"),
    "2026-04-30": Decimal("7.1"),
}
MATURITIES = {
    "H-CNY": date(2027, 3, 31),
    "A-USD": date(2028, 3, 31),
    "T-CNY": date(2029, 3, 31),
}
PNL_EXPECTED = {
    "2026-03-31": {
        # T/CNY interbank-CD 514 is VAT-inclusive in this approved window:
        # 10 + (2 * 7.0) + (3 / 1.06), stored to the formal 8-decimal scale.
        "interest_income_514": Decimal("26.83018868"),
        "fair_value_change_516": Decimal("4"),
        "capital_gain_517": Decimal("27"),
        "manual_adjustment": Decimal("0"),
        "total_pnl": Decimal("57.83018868"),
    },
    "2026-04-30": {
        # 10 + (2 * 7.1) + (3 / 1.06), stored to the formal 8-decimal scale.
        "interest_income_514": Decimal("27.03018868"),
        "fair_value_change_516": Decimal("4"),
        "capital_gain_517": Decimal("27.2"),
        "manual_adjustment": Decimal("1.5"),
        "total_pnl": Decimal("59.73018868"),
    },
}

# The ledger evidence is deliberately fixed separately from PNL_EXPECTED.  These
# are raw account/value observations for the synthetic CNX ledger; changing the
# formal PnL oracle cannot silently change the independent reconciliation input.
INDEPENDENT_LEDGER_ROWS = {
    "2026-03-31": (
        ("51490000001", Decimal("26.83018868")),
        ("51690000001", Decimal("4")),
        ("51790000001", Decimal("27")),
    ),
    "2026-04-30": (
        ("51490000001", Decimal("27.03018868")),
        ("51690000001", Decimal("4")),
        ("51790000001", Decimal("27.2")),
    ),
}
INDEPENDENT_LEDGER_ADJUSTMENTS = {
    "2026-03-31": Decimal("0"),
    "2026-04-30": Decimal("1.5"),
}
CORRECTED_MARCH_LEDGER_ROWS = (
    ("51490000001", Decimal("28.83018868")),
    ("51690000001", Decimal("4")),
    ("51790000001", Decimal("27")),
)


def _seed_snapshots_and_fx(duckdb_path: Path) -> None:
    with duckdb.connect(str(duckdb_path)) as connection:
        ensure_snapshot_tables(connection)
        connection.execute(
            """
            create table if not exists fx_daily_mid (
              trade_date varchar,
              base_currency varchar,
              quote_currency varchar,
              mid_rate decimal(24, 8),
              source_name varchar,
              is_business_day boolean,
              is_carry_forward boolean,
              source_version varchar,
              observed_trade_date varchar
            )
            """
        )
        connection.executemany(
            """
            insert into fx_daily_mid (
              trade_date, base_currency, quote_currency, mid_rate, source_name,
              is_business_day, is_carry_forward, source_version, observed_trade_date
            ) values (?, 'USD', 'CNY', ?, 'CFETS', true, false, ?, ?)
            """,
            [
                (
                    report_date,
                    fx_rate,
                    f"sv_joint_fx_{report_date.replace('-', '')}",
                    report_date,
                )
                for report_date, fx_rate in FX_BY_DATE.items()
            ],
        )

        zqtz_rows: list[list[object]] = []
        tyw_rows: list[list[object]] = []
        for report_date in REPORT_DATES:
            source_version = f"sv_joint_position_{report_date.replace('-', '')}"
            ingest_batch_id = f"ib_joint_{report_date.replace('-', '')}"
            zqtz_rows.extend(
                [
                    _zqtz_row(
                        report_date,
                        "H-CNY",
                        "持有至到期投资",
                        "利率债",
                        "国债",
                        "CNY",
                        Decimal("100"),
                        Decimal("100"),
                        Decimal("98"),
                        Decimal("1"),
                        MATURITIES["H-CNY"],
                        False,
                        source_version,
                        ingest_batch_id,
                    ),
                    _zqtz_row(
                        report_date,
                        "A-USD",
                        "可供出售类资产",
                        "信用债",
                        "企业债",
                        "USD",
                        Decimal("20"),
                        Decimal("19"),
                        Decimal("18"),
                        Decimal("0.5"),
                        MATURITIES["A-USD"],
                        False,
                        source_version,
                        ingest_batch_id,
                    ),
                    _zqtz_row(
                        report_date,
                        "T-CNY",
                        "交易性金融资产",
                        "信用债",
                        "公司债",
                        "CNY",
                        Decimal("50"),
                        Decimal("55"),
                        Decimal("52"),
                        Decimal("0.5"),
                        MATURITIES["T-CNY"],
                        False,
                        source_version,
                        ingest_batch_id,
                    ),
                    _zqtz_row(
                        report_date,
                        "ISS-CNY",
                        "发行类债券",
                        "发行类债券",
                        "同业存单",
                        "CNY",
                        Decimal("80"),
                        Decimal("80"),
                        Decimal("80"),
                        Decimal("0"),
                        date(2027, 12, 31),
                        True,
                        source_version,
                        ingest_batch_id,
                    ),
                ]
            )
            tyw_rows.append(
                [
                    report_date,
                    "LIAB-CNY",
                    "持有至到期同业存单",
                    "liability",
                    "合成交易对手",
                    "负债账户",
                    "一般",
                    "银行",
                    "CNY",
                    Decimal("30"),
                    Decimal("0.2"),
                    Decimal("0.015"),
                    "2026-05-31",
                    None,
                    source_version,
                    "rv_joint_snapshot_v1",
                    ingest_batch_id,
                    f"trace_liab_{report_date.replace('-', '')}",
                ]
            )

        connection.executemany(
            """
            insert into zqtz_bond_daily_snapshot (
              report_date, instrument_code, instrument_name, portfolio_name, cost_center,
              account_category, asset_class, bond_type, issuer_name, industry_name, rating,
              currency_code, face_value_native, market_value_native, amortized_cost_native,
              accrued_interest_native, coupon_rate, ytm_value, maturity_date, next_call_date,
              overdue_days, is_issuance_like, interest_mode, source_version, rule_version,
              ingest_batch_id, trace_id
            ) values (
              ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
            )
            """,
            zqtz_rows,
        )
        connection.executemany(
            """
            insert into tyw_interbank_daily_snapshot (
              report_date, position_id, product_type, position_side, counterparty_name,
              account_type, special_account_type, core_customer_type, currency_code,
              principal_native, accrued_interest_native, funding_cost_rate, maturity_date,
              pledged_bond_code, source_version, rule_version, ingest_batch_id, trace_id
            ) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            tyw_rows,
        )


def _zqtz_row(
    report_date: str,
    instrument_code: str,
    account_category: str,
    asset_class: str,
    bond_type: str,
    currency_code: str,
    face_value: Decimal,
    market_value: Decimal,
    amortized_cost: Decimal,
    accrued_interest: Decimal,
    maturity_date: date,
    is_issuance_like: bool,
    source_version: str,
    ingest_batch_id: str,
) -> list[object]:
    return [
        report_date,
        instrument_code,
        f"合成-{instrument_code}",
        "联合回放组合",
        "CC-JOINT",
        account_category,
        asset_class,
        bond_type,
        f"发行人-{instrument_code}",
        "合成行业",
        "AAA",
        currency_code,
        face_value,
        market_value,
        amortized_cost,
        accrued_interest,
        Decimal("0"),
        Decimal("2.0"),
        maturity_date.isoformat(),
        None,
        0,
        is_issuance_like,
        "到期一次还本付息",
        source_version,
        "rv_joint_snapshot_v1",
        ingest_batch_id,
        f"trace_{instrument_code.lower()}_{report_date.replace('-', '')}",
    ]


def _seed_curves(duckdb_path: Path) -> None:
    points = [
        YieldCurvePoint("3M", Decimal("1.0")),
        YieldCurvePoint("6M", Decimal("1.1")),
        YieldCurvePoint("1Y", Decimal("1.5")),
        YieldCurvePoint("2Y", Decimal("1.6")),
        YieldCurvePoint("3Y", Decimal("1.8")),
        YieldCurvePoint("5Y", Decimal("2.0")),
        YieldCurvePoint("7Y", Decimal("2.2")),
        YieldCurvePoint("10Y", Decimal("2.5")),
        YieldCurvePoint("20Y", Decimal("2.8")),
        YieldCurvePoint("30Y", Decimal("3.0")),
    ]
    repo = YieldCurveRepository(str(duckdb_path))
    for trade_date in ("2026-03-01", "2026-03-31", "2026-04-01", "2026-04-30"):
        snapshots = [
            YieldCurveSnapshot(
                curve_type=curve_type,
                trade_date=trade_date,
                points=points,
                vendor_name="joint_replay",
                vendor_version="vv_joint_curve_v1",
                source_version=f"sv_joint_curve_{trade_date.replace('-', '')}",
            )
            for curve_type in ("treasury", "cdb", "aaa_credit")
        ]
        with repository_task_write_scope("backend.app.tasks.joint_replay_curve_seed"):
            repo.replace_curve_snapshots(
                trade_date=trade_date,
                snapshots=snapshots,
                rule_version=YIELD_CURVE_RULE_VERSION,
            )


def _fi_rows(report_date: str) -> list[dict[str, object]]:
    fx_rate = FX_BY_DATE[report_date]
    return [
        {
            "report_date": report_date,
            "instrument_code": "H-CNY",
            "portfolio_name": "联合回放组合",
            "cost_center": "CC-JOINT",
            "invest_type_raw": "H",
            "asset_class": "国债",
            "interest_income_514": "10",
            "fair_value_change_516": "5",
            "capital_gain_517": "-10.6",
            "event_type": "fi_cumulative_realized_517",
            "currency_basis": "CNY",
            "source_version": f"sv_joint_pnl_{report_date.replace('-', '')}",
            "ingest_batch_id": f"ib_joint_{report_date.replace('-', '')}",
            "trace_id": f"trace_pnl_h_{report_date.replace('-', '')}",
        },
        {
            "report_date": report_date,
            "instrument_code": "A-USD",
            "portfolio_name": "联合回放组合",
            "cost_center": "CC-JOINT",
            "invest_type_raw": "A",
            "asset_class": "政策性金融债",
            "interest_income_514": "2",
            "fair_value_change_516": "1",
            "capital_gain_517": "-2.12",
            "event_type": "fi_cumulative_realized_517",
            "currency_basis": "CNY",
            "fx_base_currency": "USD",
            "source_version": f"sv_joint_pnl_{report_date.replace('-', '')}",
            "ingest_batch_id": f"ib_joint_{report_date.replace('-', '')}",
            "trace_id": f"trace_pnl_a_{report_date.replace('-', '')}",
            "expected_fx_rate": fx_rate,
        },
        {
            "report_date": report_date,
            "instrument_code": "T-CNY",
            "portfolio_name": "联合回放组合",
            "cost_center": "CC-JOINT",
            "invest_type_raw": "T",
            "asset_class": "同业存单",
            "interest_income_514": "3",
            "fair_value_change_516": "4",
            "capital_gain_517": "-3.18",
            "event_type": "fi_cumulative_realized_517",
            "currency_basis": "CNY",
            "source_version": f"sv_joint_pnl_{report_date.replace('-', '')}",
            "ingest_batch_id": f"ib_joint_{report_date.replace('-', '')}",
            "trace_id": f"trace_pnl_t_{report_date.replace('-', '')}",
            **(
                {"manual_adjustment": "1.5", "approval_status": "approved"}
                if report_date == "2026-04-30"
                else {}
            ),
        },
    ]


def _write_independent_ledger_pair(
    source_dir: Path,
    report_date: str,
    ledger_rows: tuple[tuple[str, Decimal], ...],
) -> None:
    report_day = date.fromisoformat(report_date)
    month_key = report_day.strftime("%Y%m")
    month_start = report_day.replace(day=1)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, currency in (("综本", "CNX"), ("人民币", "CNY")):
        sheet = workbook.create_sheet(sheet_name)
        sheet["A2"] = (
            f"会计期间： {month_start.isoformat()}--{report_date} 币种： {currency}"
        )
        for row_number, (account_code, raw_value) in enumerate(ledger_rows, start=7):
            value = raw_value
            if currency == "CNY":
                value = value / Decimal("2")
            sheet.cell(row=row_number, column=1, value=account_code)
            sheet.cell(row=row_number, column=2, value=f"联合回放-{account_code[:3]}")
            sheet.cell(row=row_number, column=3, value=currency)
            sheet.cell(row=row_number, column=4, value=Decimal("0"))
            sheet.cell(row=row_number, column=5, value=Decimal("0"))
            sheet.cell(row=row_number, column=6, value=value)
            sheet.cell(row=row_number, column=7, value=-value)
    workbook.save(source_dir / f"总账对账{month_key}.xlsx")
    workbook.close()

    average = Workbook()
    average.active.title = "年日均"
    average.create_sheet("月日均")
    average.save(source_dir / f"日均{month_key}.xlsx")
    average.close()


def _prepare_independent_reference(
    *,
    duckdb_path: Path,
    governance_dir: Path,
    source_dir: Path,
    report_date: str,
    ledger_rows: tuple[tuple[str, Decimal], ...] | None = None,
    ledger_adjustment: Decimal | None = None,
    write_ledger: bool = True,
) -> tuple[PnlFormalFactDependencyBinding, IndependentPnlReference]:
    resolved_ledger_rows = ledger_rows or INDEPENDENT_LEDGER_ROWS[report_date]
    resolved_adjustment = (
        INDEPENDENT_LEDGER_ADJUSTMENTS[report_date]
        if ledger_adjustment is None
        else ledger_adjustment
    )
    if write_ledger:
        _write_independent_ledger_pair(
            source_dir,
            report_date,
            resolved_ledger_rows,
        )
    month_key = report_date.replace("-", "")[:6]
    pair = next(
        item for item in discover_source_pairs(source_dir) if item.month_key == month_key
    )
    binding = read_formal_fact_dependency_binding(
        duckdb_path,
        report_date=report_date,
    )
    assert binding is not None
    scopes = {
        component: LedgerPnlComponentScope(
            status="ready",
            account_prefixes=(prefix,),
            evidence_refs=(
                f"synthetic-scope-only://joint-replay/{month_key}/{prefix}",
            ),
        )
        for component, prefix in (
            ("interest_income_514", "514"),
            ("fair_value_change_516", "516"),
            ("capital_gain_517", "517"),
        )
    }
    reference = prepare_and_store_independent_ledger_pnl_reference(
        pair,
        governance_dir=governance_dir,
        component_scopes=scopes,
        manual_adjustment=ready_adjustment_reference(
            resolved_adjustment,
            evidence_refs=(f"synthetic-approval://joint-replay/{month_key}",),
        ),
        formal_dependency_binding=binding,
    )
    return binding, reference


def _run_real_tasks(
    *,
    duckdb_path: Path,
    governance_dir: Path,
    report_date: str,
) -> dict[str, dict[str, object]]:
    balance = materialize_balance_analysis_facts.fn(
        report_date=report_date,
        duckdb_path=str(duckdb_path),
        governance_dir=str(governance_dir),
        use_existing_fx_only=True,
    )
    pnl = materialize_pnl_facts.fn(
        report_date=report_date,
        is_month_end=True,
        fi_rows=_fi_rows(report_date),
        nonstd_rows_by_type={},
        duckdb_path=str(duckdb_path),
        governance_dir=str(governance_dir),
        formal_pnl_enabled=True,
        formal_pnl_scope_json='["*"]',
    )
    bond = materialize_bond_analytics_facts.fn(
        report_date=report_date,
        duckdb_path=str(duckdb_path),
        governance_dir=str(governance_dir),
    )
    risk = materialize_risk_tensor_facts.fn(
        report_date=report_date,
        duckdb_path=str(duckdb_path),
        governance_dir=str(governance_dir),
    )
    return {"balance": balance, "pnl": pnl, "bond": bond, "risk": risk}


def _zero_coupon_expected(
    *, report_date: str, maturity_date: date, face_value_cny: Decimal
) -> tuple[Decimal, Decimal, Decimal]:
    years = Decimal((maturity_date - date.fromisoformat(report_date)).days) / Decimal(365)
    modified = years / Decimal("1.02")
    dv01 = face_value_cny * modified / Decimal("10000")
    quantum = Decimal("0.00000001")
    return years.quantize(quantum), modified.quantize(quantum), dv01.quantize(quantum)


def _liability_outflow_90d(report_date: str) -> Decimal:
    days = (date(2026, 5, 31) - date.fromisoformat(report_date)).days
    principal = Decimal("30")
    annual_rate_decimal = Decimal("0.015") / Decimal("100")
    value = principal + principal * annual_rate_decimal * Decimal(days) / Decimal("365")
    return value.quantize(Decimal("0.00000001"))


def test_two_month_joint_replay_runs_real_tasks_and_matches_independent_expectations(
    tmp_path: Path,
) -> None:
    duckdb_path = tmp_path / "joint-replay.duckdb"
    governance_dir = tmp_path / "governance"
    ledger_dir = tmp_path / "independent-ledger"
    ledger_dir.mkdir()
    _seed_snapshots_and_fx(duckdb_path)
    _seed_curves(duckdb_path)

    task_receipts: dict[str, dict[str, dict[str, object]]] = {}
    bindings: dict[str, PnlFormalFactDependencyBinding] = {}
    curve_repo = YieldCurveRepository(str(duckdb_path))
    for report_date in REPORT_DATES:
        curve_snapshot = curve_repo.fetch_curve_snapshot(report_date, "treasury")
        assert curve_snapshot is not None
        assert curve_snapshot["curve"]["1Y"] == Decimal("1.5")
        assert curve_snapshot["source_version"] == (
            f"sv_joint_curve_{report_date.replace('-', '')}"
        )
        assert curve_snapshot["rule_version"] == YIELD_CURVE_RULE_VERSION
    for report_date in REPORT_DATES:
        task_receipts[report_date] = _run_real_tasks(
            duckdb_path=duckdb_path,
            governance_dir=governance_dir,
            report_date=report_date,
        )
        bindings[report_date], _ = _prepare_independent_reference(
            duckdb_path=duckdb_path,
            governance_dir=governance_dir,
            source_dir=ledger_dir,
            report_date=report_date,
        )

    assert {
        (report_date, task_name, receipt["status"])
        for report_date, receipts in task_receipts.items()
        for task_name, receipt in receipts.items()
    } == {
        (report_date, task_name, "completed")
        for report_date in REPORT_DATES
        for task_name in ("balance", "pnl", "bond", "risk")
    }
    assert all(
        str(receipt.get("source_version") or "").strip()
        and str(receipt.get("rule_version") or "").strip()
        for receipts in task_receipts.values()
        for receipt in receipts.values()
    )

    with duckdb.connect(str(duckdb_path), read_only=True) as connection:
        for report_date in REPORT_DATES:
            fx_rate = FX_BY_DATE[report_date]
            balance_rows = connection.execute(
                """
                select instrument_code, invest_type_std, accounting_basis, position_scope,
                       face_value_amount, market_value_amount, amortized_cost_amount,
                       accrued_interest_amount
                from fact_formal_zqtz_balance_daily
                where report_date = ? and currency_basis = 'CNY'
                order by instrument_code
                """,
                [report_date],
            ).fetchall()
            assert balance_rows == [
                (
                    "A-USD",
                    "A",
                    "FVOCI",
                    "asset",
                    Decimal("20") * fx_rate,
                    Decimal("19") * fx_rate,
                    Decimal("18") * fx_rate,
                    Decimal("0.5") * fx_rate,
                ),
                (
                    "H-CNY",
                    "H",
                    "AC",
                    "asset",
                    Decimal("100"),
                    Decimal("100"),
                    Decimal("98"),
                    Decimal("1"),
                ),
                (
                    "ISS-CNY",
                    "H",
                    "AC",
                    "liability",
                    Decimal("80"),
                    Decimal("80"),
                    Decimal("80"),
                    Decimal("0"),
                ),
                (
                    "T-CNY",
                    "T",
                    "FVTPL",
                    "asset",
                    Decimal("50"),
                    Decimal("55"),
                    Decimal("52"),
                    Decimal("0.5"),
                ),
            ]

            bond_rows = connection.execute(
                """
                select instrument_code, accounting_class, face_value, market_value,
                       macaulay_duration, modified_duration, dv01
                from fact_formal_bond_analytics_daily
                where report_date = ?
                order by instrument_code
                """,
                [report_date],
            ).fetchall()
            assert [row[0] for row in bond_rows] == ["A-USD", "H-CNY", "T-CNY"]
            assert [row[1] for row in bond_rows] == ["OCI", "AC", "TPL"]
            expected_face = {
                "A-USD": Decimal("20") * fx_rate,
                "H-CNY": Decimal("100"),
                "T-CNY": Decimal("50"),
            }
            expected_market = {
                "A-USD": Decimal("19") * fx_rate,
                "H-CNY": Decimal("100"),
                "T-CNY": Decimal("55"),
            }
            expected_dv01_total = Decimal("0")
            expected_dv01_by_instrument: dict[str, Decimal] = {}
            for (
                instrument_code,
                _accounting_class,
                face_value,
                market_value,
                macaulay,
                modified,
                dv01,
            ) in bond_rows:
                expected = _zero_coupon_expected(
                    report_date=report_date,
                    maturity_date=MATURITIES[instrument_code],
                    face_value_cny=expected_face[instrument_code],
                )
                assert face_value == expected_face[instrument_code]
                assert market_value == expected_market[instrument_code]
                assert (macaulay, modified, dv01) == expected
                expected_dv01_total += expected[2]
                expected_dv01_by_instrument[instrument_code] = expected[2]

            risk = RiskTensorRepository(str(duckdb_path)).fetch_risk_tensor_row(report_date)
            assert risk is not None
            assert risk["bond_count"] == 3
            assert risk["total_market_value"] == sum(
                expected_market.values(), Decimal("0")
            )
            assert risk["portfolio_dv01"] == expected_dv01_total
            assert risk["regulatory_dv01"] == expected_dv01_total
            assert risk["krd_1y"] == expected_dv01_by_instrument["H-CNY"]
            assert risk["krd_3y"] == (
                expected_dv01_by_instrument["A-USD"]
                + expected_dv01_by_instrument["T-CNY"]
            )
            assert {
                field: risk[field]
                for field in ("krd_5y", "krd_7y", "krd_10y", "krd_30y")
            } == {
                "krd_5y": Decimal("0"),
                "krd_7y": Decimal("0"),
                "krd_10y": Decimal("0"),
                "krd_30y": Decimal("0"),
            }
            assert sum(
                (
                    risk[field]
                    for field in (
                        "krd_1y",
                        "krd_3y",
                        "krd_5y",
                        "krd_7y",
                        "krd_10y",
                        "krd_30y",
                    )
                ),
                Decimal("0"),
            ) == expected_dv01_total
            shock_scenarios = {
                scenario["shock_bp"]: scenario["estimated_pnl"]
                for scenario in build_dv01_shock_scenario_payloads(
                    total_dv01=expected_dv01_total,
                    shocks=[Decimal("1")],
                )
            }
            assert shock_scenarios == {
                Decimal("1"): -expected_dv01_total,
                Decimal("-1"): expected_dv01_total,
            }

            expected_liability_90d = _liability_outflow_90d(report_date)
            assert risk["asset_cashflow_30d"] == Decimal("0")
            assert risk["asset_cashflow_90d"] == Decimal("0")
            assert risk["liability_cashflow_30d"] == Decimal("0")
            assert risk["liability_cashflow_90d"] == expected_liability_90d
            assert risk["liquidity_gap_30d"] == Decimal("0")
            assert risk["liquidity_gap_90d"] == -expected_liability_90d

            pnl_repo = PnlRepository(str(duckdb_path))
            pnl_rows = {
                str(row["instrument_code"]): row
                for row in pnl_repo.fetch_formal_fi_rows(report_date)
            }
            adjustment = (
                Decimal("1.5") if report_date == "2026-04-30" else Decimal("0")
            )
            expected_pnl_rows = {
                "H-CNY": (
                    "H",
                    "AC",
                    "CNY",
                    Decimal("10"),
                    Decimal("0"),
                    Decimal("10"),
                    Decimal("0"),
                    Decimal("20"),
                ),
                "A-USD": (
                    "A",
                    "FVOCI",
                    "CNY",
                    Decimal("2") * fx_rate,
                    Decimal("0"),
                    Decimal("2") * fx_rate,
                    Decimal("0"),
                    Decimal("4") * fx_rate,
                ),
                "T-CNY": (
                    "T",
                    "FVTPL",
                    "CNY",
                    Decimal("2.83018868"),
                    Decimal("4"),
                    Decimal("3"),
                    adjustment,
                    Decimal("9.83018868") + adjustment,
                ),
            }
            assert set(pnl_rows) == set(expected_pnl_rows)
            for instrument_code, expected_row in expected_pnl_rows.items():
                row = pnl_rows[instrument_code]
                assert (
                    row["invest_type_std"],
                    row["accounting_basis"],
                    row["currency_basis"],
                    row["interest_income_514"],
                    row["fair_value_change_516"],
                    row["capital_gain_517"],
                    row["manual_adjustment"],
                    row["total_pnl"],
                ) == expected_row

            pnl_totals = pnl_repo.overview_totals(report_date)
            for field_name, expected_value in PNL_EXPECTED[report_date].items():
                assert pnl_totals[field_name] == expected_value
            binding = bindings[report_date]
            reconciliation = load_and_reconcile_prepared_independent_pnl_reference(
                pnl_totals,
                governance_dir=governance_dir,
                report_date=report_date,
                formal_dependency_revision=binding.formal_dependency_revision,
                formal_dependency_protocol_version=(
                    binding.formal_dependency_protocol_version
                ),
            )
            assert reconciliation["status"] == "pass"
            assert Decimal(str(reconciliation["total_diff_yuan"])) == Decimal("0")

            # Cross-module comparisons stop at identical asset/date/CNY identity.
            # Closing balance (stock), monthly PnL (flow), and positive CNY DV01
            # sensitivity magnitude per 1 bp are distinct measures. Shock PnL
            # direction is checked independently above.
            # They intentionally remain separate and are not forced to equal.
            assert {row[0] for row in balance_rows if row[3] == "asset"} == {
                row[0] for row in bond_rows
            }


def test_joint_replay_detects_ledger_mutation_and_replaces_corrected_history(
    tmp_path: Path,
) -> None:
    duckdb_path = tmp_path / "joint-replay-correction.duckdb"
    governance_dir = tmp_path / "governance"
    ledger_dir = tmp_path / "independent-ledger"
    ledger_dir.mkdir()
    _seed_snapshots_and_fx(duckdb_path)
    _seed_curves(duckdb_path)

    bindings: dict[str, PnlFormalFactDependencyBinding] = {}
    for report_date in REPORT_DATES:
        _run_real_tasks(
            duckdb_path=duckdb_path,
            governance_dir=governance_dir,
            report_date=report_date,
        )
        bindings[report_date], _ = _prepare_independent_reference(
            duckdb_path=duckdb_path,
            governance_dir=governance_dir,
            source_dir=ledger_dir,
            report_date=report_date,
        )

    april_date = "2026-04-30"
    april_totals_before = PnlRepository(str(duckdb_path)).overview_totals(april_date)
    with duckdb.connect(str(duckdb_path), read_only=True) as connection:
        april_balance_before = connection.execute(
            """
            select count(*), coalesce(sum(market_value_amount), 0)
            from fact_formal_zqtz_balance_daily
            where report_date = ? and currency_basis = 'CNY'
            """,
            [april_date],
        ).fetchone()
        april_tyw_before = connection.execute(
            """
            select count(*), coalesce(sum(principal_amount), 0)
            from fact_formal_tyw_balance_daily
            where report_date = ? and currency_basis = 'CNY'
            """,
            [april_date],
        ).fetchone()
        april_bond_before = connection.execute(
            """
            select count(*), coalesce(sum(market_value), 0)
            from fact_formal_bond_analytics_daily
            where report_date = ?
            """,
            [april_date],
        ).fetchone()
    april_risk_before = RiskTensorRepository(str(duckdb_path)).fetch_risk_tensor_row(
        april_date
    )
    assert april_risk_before is not None
    april_ledger_path = ledger_dir / "总账对账202604.xlsx"
    april_book = load_workbook(april_ledger_path)
    try:
        original_ledger_514 = Decimal(str(april_book["综本"]["F7"].value))
        april_book["综本"]["F7"] = original_ledger_514 + Decimal("1")
        april_book["综本"]["G7"] = -april_book["综本"]["F7"].value
        april_book.save(april_ledger_path)
    finally:
        april_book.close()
    april_binding, _ = _prepare_independent_reference(
        duckdb_path=duckdb_path,
        governance_dir=governance_dir,
        source_dir=ledger_dir,
        report_date=april_date,
        write_ledger=False,
    )
    april_mismatch = load_and_reconcile_prepared_independent_pnl_reference(
        april_totals_before,
        governance_dir=governance_dir,
        report_date=april_date,
        formal_dependency_revision=april_binding.formal_dependency_revision,
        formal_dependency_protocol_version=(
            april_binding.formal_dependency_protocol_version
        ),
    )
    assert april_mismatch["status"] == "fail"
    assert april_mismatch["total_diff_yuan"] == "-1.00000000"
    assert april_mismatch["components"][0]["status"] == "fail"
    assert april_totals_before["total_pnl"] == sum(
        (
            april_totals_before[field]
            for field in (
                "interest_income_514",
                "fair_value_change_516",
                "capital_gain_517",
                "manual_adjustment",
            )
        ),
        Decimal("0"),
    )

    march_date = "2026-03-31"
    old_march_binding = bindings[march_date]
    with duckdb.connect(str(duckdb_path)) as connection:
        connection.execute(
            """
            update zqtz_bond_daily_snapshot
            set market_value_native = 21,
                source_version = 'sv_joint_position_20260331_corrected'
            where report_date = ? and instrument_code = 'A-USD'
            """,
            [march_date],
        )
    corrected_pnl_rows = _fi_rows(march_date)
    corrected_pnl_rows[0]["interest_income_514"] = "12"
    for row in corrected_pnl_rows:
        row["source_version"] = "sv_joint_pnl_20260331_corrected"

    correction_receipts = {
        "balance": materialize_balance_analysis_facts.fn(
            report_date=march_date,
            duckdb_path=str(duckdb_path),
            governance_dir=str(governance_dir),
            use_existing_fx_only=True,
        ),
        "pnl": materialize_pnl_facts.fn(
            report_date=march_date,
            is_month_end=True,
            fi_rows=corrected_pnl_rows,
            nonstd_rows_by_type={},
            duckdb_path=str(duckdb_path),
            governance_dir=str(governance_dir),
            formal_pnl_enabled=True,
            formal_pnl_scope_json='["*"]',
        ),
        "bond": materialize_bond_analytics_facts.fn(
            report_date=march_date,
            duckdb_path=str(duckdb_path),
            governance_dir=str(governance_dir),
        ),
        "risk": materialize_risk_tensor_facts.fn(
            report_date=march_date,
            duckdb_path=str(duckdb_path),
            governance_dir=str(governance_dir),
        ),
    }
    assert {name: receipt["status"] for name, receipt in correction_receipts.items()} == {
        "balance": "completed",
        "pnl": "completed",
        "bond": "completed",
        "risk": "completed",
    }

    corrected_binding = read_formal_fact_dependency_binding(
        duckdb_path,
        report_date=march_date,
    )
    assert corrected_binding is not None
    assert (
        corrected_binding.formal_dependency_revision
        > old_march_binding.formal_dependency_revision
    )
    assert (
        corrected_binding.formal_source_version
        != old_march_binding.formal_source_version
    )
    stale_reconciliation = load_and_reconcile_prepared_independent_pnl_reference(
        PnlRepository(str(duckdb_path)).overview_totals(march_date),
        governance_dir=governance_dir,
        report_date=march_date,
        formal_dependency_revision=corrected_binding.formal_dependency_revision,
        formal_dependency_protocol_version=(
            corrected_binding.formal_dependency_protocol_version
        ),
    )
    assert stale_reconciliation["status"] == "pending"
    assert stale_reconciliation["reason"] == "independent ledger reference is not prepared"

    corrected_expected = dict(PNL_EXPECTED[march_date])
    corrected_expected["interest_income_514"] += Decimal("2")
    corrected_expected["total_pnl"] += Decimal("2")
    rebuilt_binding, _ = _prepare_independent_reference(
        duckdb_path=duckdb_path,
        governance_dir=governance_dir,
        source_dir=ledger_dir,
        report_date=march_date,
        ledger_rows=CORRECTED_MARCH_LEDGER_ROWS,
        ledger_adjustment=Decimal("0"),
    )
    corrected_totals = PnlRepository(str(duckdb_path)).overview_totals(march_date)
    assert corrected_totals["formal_fi_row_count"] == 3
    assert corrected_totals["nonstd_bridge_row_count"] == 0
    for field_name, expected_value in corrected_expected.items():
        assert corrected_totals[field_name] == expected_value
    rebuilt_reconciliation = load_and_reconcile_prepared_independent_pnl_reference(
        corrected_totals,
        governance_dir=governance_dir,
        report_date=march_date,
        formal_dependency_revision=rebuilt_binding.formal_dependency_revision,
        formal_dependency_protocol_version=(
            rebuilt_binding.formal_dependency_protocol_version
        ),
    )
    assert rebuilt_reconciliation["status"] == "pass"

    with duckdb.connect(str(duckdb_path), read_only=True) as connection:
        corrected_march_market = connection.execute(
            """
            select market_value_amount
            from fact_formal_zqtz_balance_daily
            where report_date = ? and instrument_code = 'A-USD'
              and currency_basis = 'CNY'
            """,
            [march_date],
        ).fetchone()[0]
        corrected_bond_market = connection.execute(
            """
            select market_value
            from fact_formal_bond_analytics_daily
            where report_date = ? and instrument_code = 'A-USD'
            """,
            [march_date],
        ).fetchone()[0]
        assert corrected_march_market == Decimal("147")
        assert corrected_bond_market == Decimal("147")
        assert connection.execute(
            "select count(*) from fact_formal_bond_analytics_daily where report_date = ?",
            [march_date],
        ).fetchone()[0] == 3
        assert connection.execute(
            """
            select count(*), count(distinct instrument_code)
            from fact_formal_zqtz_balance_daily
            where report_date = ? and currency_basis = 'CNY'
            """,
            [march_date],
        ).fetchone() == (4, 4)
        assert connection.execute(
            """
            select count(*), coalesce(sum(market_value_amount), 0)
            from fact_formal_zqtz_balance_daily
            where report_date = ? and currency_basis = 'CNY'
            """,
            [april_date],
        ).fetchone() == april_balance_before
        assert connection.execute(
            """
            select count(*), coalesce(sum(principal_amount), 0)
            from fact_formal_tyw_balance_daily
            where report_date = ? and currency_basis = 'CNY'
            """,
            [april_date],
        ).fetchone() == april_tyw_before
        assert connection.execute(
            """
            select count(*), coalesce(sum(market_value), 0)
            from fact_formal_bond_analytics_daily
            where report_date = ?
            """,
            [april_date],
        ).fetchone() == april_bond_before

    corrected_risk = RiskTensorRepository(str(duckdb_path)).fetch_risk_tensor_row(march_date)
    assert corrected_risk is not None
    assert corrected_risk["total_market_value"] == Decimal("302")
    assert RiskTensorRepository(str(duckdb_path)).fetch_risk_tensor_row(
        april_date
    ) == april_risk_before
    assert PnlRepository(str(duckdb_path)).overview_totals(april_date) == april_totals_before
