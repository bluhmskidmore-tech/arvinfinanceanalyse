"""Bounded analytical QDB contract: unavailable source values are not zero.

Synthetic inputs only. These cases do not promote this analytical surface or
certify completeness for accounts absent from every source index.
"""
from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.core_finance import qdb_gl_monthly_analysis as q
from backend.app.services.qdb_gl_input_validation_service import validate_qdb_gl_baseline_source

pytestmark = [pytest.mark.excluded_surface_acceptance, pytest.mark.surface_qdb_gl]
D = Decimal
YI = D('100000000')


def synthetic_merged(amounts=None, *, missing=()):
    """Derive complete independent source totals, then omit explicit tuples."""
    amounts = amounts if amounts is not None else {'12303000001': YI}
    ledger = []
    totals = {3: {}, 5: {}, 11: {}}
    for code, amount in amounts.items():
        ledger.append({'科目代码': code, '科目名称': 'synthetic only', '币种': 'CNX',
                       '期初余额': amount, '本期借方': D(0), '本期贷方': D(0),
                       '期末余额': amount, '变动额': D(0), '本期净额': D(0),
                       '大类1': code[:1], '大类3': code[:3], '大类5': code[:5]})
        for size, index in totals.items():
            index[code[:size]] = index.get(code[:size], D(0)) + amount
    averages = {
        f'{period}日均_CNX_{size}d': [
            {'科目代码': code, '日均余额': value}
            for code, value in index.items() if (period, code) not in missing
        ]
        for period in ('月', '年') for size, index in totals.items()
    }
    return q.merge_all({'综本': ledger, '人民币': ledger}, averages)


def metric(rows, name):
    return next(row for row in rows if row['指标'] == name)


def sheet(workbook, key):
    return next(item for item in workbook['sheets'] if item['key'] == key)


def workbook(merged, prior=None):
    return q.build_qdb_gl_monthly_analysis_workbook(
        report_month='202601', merged_data=merged,
        comparison_data={'prior_month': prior, 'prior_year': prior} if prior is not None else None,
    )


FAMILIES = [
    ('segment_base_scale', q._segment_base_scale_raw_rows, '公司贷款合计', '12303000001', '123'),
    ('company_scale', q._company_scale_raw_rows, '公司贷款合计', '12303000001', '123'),
    ('retail_scale', q._retail_scale_raw_rows, '参考：个人贷款合计', '12203000001', '122'),
    ('financial_market_scale', q._financial_market_scale_raw_rows, '生息债券投资', '14203000001', '142'),
]


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
@pytest.mark.parametrize('period', ['月', '年'])
def test_present_component_missing_average_preserves_spot_and_other_period(key, builder, label, code, average_code, period):
    merged = synthetic_merged({code: YI}, missing=[(period, average_code)])
    row = metric(builder(merged), label)
    assert row[period + '日均'] is None
    assert row['时点余额'] == YI
    assert row[('年' if period == '月' else '月') + '日均'] == YI
    assert row['口径来源'].startswith('source_missing:')
    assert average_code in row['口径来源'] and period + '日均' in row['口径来源']


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
@pytest.mark.parametrize('value', [D(0), YI * D('1.25'), -YI * D('1.25')])
def test_observed_zero_and_signed_values_remain_known(key, builder, label, code, average_code, value):
    row = metric(builder(synthetic_merged({code: value})), label)
    assert [row[field] for field in ('时点余额', '年日均', '月日均')] == [value] * 3
    assert not row['口径来源'].startswith('source_missing:')


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
def test_absent_business_zero_convention_is_unchanged(key, builder, label, code, average_code):
    row = metric(builder(synthetic_merged({'10101000001': YI})), label)
    assert [row[field] for field in ('时点余额', '年日均', '月日均')] == [D(0)] * 3
    assert not row['口径来源'].startswith('source_missing:')


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
@pytest.mark.parametrize('unknown_side', ['current', 'prior'])
def test_comparisons_decline_missing_inputs_and_disclose_source(key, builder, label, code, average_code, unknown_side):
    healthy = synthetic_merged({code: YI})
    missing = synthetic_merged({code: YI}, missing=[('月', average_code)])
    current, prior = (missing, healthy) if unknown_side == 'current' else (healthy, missing)
    compare_key = 'segment_scale_compare' if key == 'segment_base_scale' else key + '_compare'
    rows = sheet(workbook(current, prior), compare_key)['rows']
    row = next(row for row in rows if row['指标'] == label and row['口径'] == '月日均环比')
    assert row['本期' if unknown_side == 'current' else '对比期'] is None
    assert row['增减额'] is None and row['增减幅%'] is None
    assert row['口径来源'].startswith('source_missing:')
    assert average_code in row['口径来源'] and '月日均' in row['口径来源']
    spot = next(row for row in rows if row['指标'] == label and row['口径'] == '时点环比')
    assert spot['本期'] == spot['对比期'] == 1 and spot['增减额'] == spot['增减幅%'] == 0


@pytest.mark.parametrize('current,prior,expected_delta,expected_pct', [
    (D(0), YI, -1, -100), (YI, D(0), 1, None),
    (D(0), D(0), 0, None), (YI * D('1.2'), YI, D('0.2'), 20),
])
def test_comparison_zero_and_nonzero_denominators(current, prior, expected_delta, expected_pct):
    result = workbook(synthetic_merged({'12303000001': current}), synthetic_merged({'12303000001': prior}))
    row = next(row for row in sheet(result, 'company_scale_compare')['rows'] if row['指标'] == '公司贷款合计' and row['口径'] == '月日均环比')
    assert D(str(row['增减额'])) == expected_delta
    assert row['增减幅%'] == expected_pct


@pytest.mark.parametrize('period', ['月', '年'])
@pytest.mark.parametrize('builder', [q._company_scale_raw_rows, q._segment_base_scale_raw_rows])
def test_known_component_does_not_hide_missing_component(builder, period):
    merged = synthetic_merged({'12303000001': YI, '12903000001': YI}, missing=[(period, '129')])
    row = metric(builder(merged), '公司贷款合计')
    assert row[period + '日均'] is None
    assert row['时点余额'] == 2 * YI
    assert '129' in row['口径来源']


@pytest.mark.parametrize('builder,label,code,average_code', [
    (q._company_scale_raw_rows, '公司贷款合计', '13604000001', '13604'),
    (q._retail_scale_raw_rows, '参考：信用卡', '13604000001', '13604'),
    (q._financial_market_scale_raw_rows, '同业资产', '14004000001', '14004'),
    (q._company_scale_raw_rows, '公司存款-结构性', '21601020001', '21601'),
    (q._retail_scale_raw_rows, '零售存款-结构性', '21602020001', '21602'),
])
@pytest.mark.parametrize('absent_periods', [('月',), ('月', '年')])
def test_missing_5d_average_is_unknown_when_gl_prefix_exists(builder, label, code, average_code, absent_periods):
    merged = synthetic_merged({code: YI}, missing=[(p, average_code) for p in absent_periods])
    row = metric(builder(merged), label)
    assert row['月日均'] is None
    assert (row['年日均'] is None) == ('年' in absent_periods)
    assert row['时点余额'] is not None
    assert average_code in row['口径来源'] and '月日均' in row['口径来源']


@pytest.mark.parametrize('builder,label,average_code', [
    (q._company_scale_raw_rows, '公司贷款合计', '13604'),
    (q._retail_scale_raw_rows, '参考：信用卡', '13604'),
    (q._financial_market_scale_raw_rows, '同业资产', '14004'),
])
def test_other_period_average_index_is_evidence_of_missing_5d_input(builder, label, average_code):
    merged = synthetic_merged({'10101000001': YI})
    merged['日均_5位'] = [{'科目代码': average_code, '月日均': None, '年日均': YI}]
    row = metric(builder(merged), label)
    assert row['月日均'] is None
    assert average_code in row['口径来源']


@pytest.mark.parametrize('builder,label,code', [
    (q._segment_base_scale_raw_rows, '个人贷款合计', '13604000001'),
    (q._segment_base_scale_raw_rows, '公司贷款合计', '13001000002'),
    (q._company_scale_raw_rows, '公司贷款合计', '13001000002'),
    (q._retail_scale_raw_rows, '参考：个人贷款合计', '13001000001'),
    (q._financial_market_scale_raw_rows, '生息债券投资', '14301010001'),
    (q._financial_market_scale_raw_rows, '同业负债', '27205000001'),
])
def test_missing_exact_or_prefix_11d_average_null_propagates_signed_term(builder, label, code):
    row = metric(builder(synthetic_merged({code: YI}, missing=[('月', code)])), label)
    assert row['月日均'] is None
    assert row['时点余额'] is not None
    assert code in row['口径来源']


@pytest.mark.parametrize('value', [None, '', 'invalid', D('NaN'), D('Infinity'), D('-Infinity')])
@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
def test_present_nonfinite_or_unparseable_average_remains_unknown(key, builder, label, code, average_code, value):
    merged = synthetic_merged({code: YI})
    for bucket in ('3位', '日均_3位'):
        next(row for row in merged[bucket] if row['科目代码'] == average_code)['月日均'] = value
    row = metric(builder(merged), label)
    assert row['月日均'] is None
    assert row['时点余额'] == YI


def test_top11_missing_year_and_trend_are_null_without_changing_selection():
    merged = synthetic_merged({'12303000001': 2 * YI}, missing=[('年', '12303000001')])
    rows = q.build_11d_top_rows(q.compute_deviation(merged['11位']))
    assert len(rows) == 1
    assert rows[0]['月日均'] == 2 and rows[0]['期末余额'] == 2
    assert rows[0]['年日均'] is None and rows[0]['趋势额'] is None and rows[0]['趋势%'] is None
    assert q.build_11d_top_rows(q.compute_deviation(synthetic_merged({'12303000001': YI})['11位'])) == []


def test_missing_scale_cannot_create_partial_company_yield_or_attribution():
    amounts = {'12303000001': YI, '12903000001': YI, '50103000001': -YI / 100}
    result = workbook(synthetic_merged(amounts, missing=[('年', '129')]), synthetic_merged(amounts))
    row = metric(sheet(result, 'income_rate_analysis')['rows'], '公司贷款利息收入')
    assert row['总账收益/支出'] == 0.01
    assert row['年日均规模'] is None and row['年化收益率/付息率%'] is None
    assert row['口径来源'].startswith('source_missing:') and '129' in row['口径来源']
    attr = metric(sheet(result, 'income_rate_attribution')['rows'], '公司贷款利息收入')
    assert all(attr[key] is None for key in ('规模贡献', '利率贡献', '恒等式自检（非独立对账）'))
    assert attr['增减额'] == 0
    assert attr['口径来源'].startswith('source_missing:') and '129' in attr['口径来源']


def test_deposit_split_dependent_rates_keep_missing_scale_and_source():
    result = workbook(synthetic_merged({'20103000001': -YI, '52101000001': YI / 100}, missing=[('月', '201'), ('年', '201')]))
    rows = sheet(result, 'deposit_interest_split')['rows']
    for label in ('公司存款', '公司存款-活期', '存款利息支出合计'):
        row = metric(rows, label)
        assert row['本期年日均'] is None and row['本月月日均'] is None
        assert row['年化付息率%'] is None and row['本月付息率%'] is None
        assert row['口径来源'].startswith('source_missing:') and '201' in row['口径来源']


def canonical_pair(tmp_path, *, month_missing=True):
    avg_path, ledger_path = tmp_path / '日均202601.xlsx', tmp_path / '总账对账202601.xlsx'
    book = Workbook()
    book.remove(book.active)
    for period in ('月', '年'):
        ws = book.create_sheet(period)
        for code_col, value_col, currency, level in q.DAILY_AVG_COLUMN_MAP:
            for col, label in zip((code_col - 1, code_col, value_col), q._DAILY_AVG_EXPECTED_BLOCK_HEADERS):
                ws.cell(row=3, column=col + 1, value=label)
            if level not in {'3d', '11d'}:
                continue
            prefix = '129' if month_missing and period == '月' else '123'
            code = prefix if level == '3d' else prefix + '03000001'
            for col, value in ((code_col - 1, currency), (code_col, code), (value_col, 0 if prefix == '129' else int(YI))):
                ws.cell(row=4, column=col + 1, value=value)
    book.save(avg_path)
    book.close()
    book = Workbook()
    book.remove(book.active)
    for name, currency in (('综本', 'CNX'), ('人民币', 'CNY')):
        ws = book.create_sheet(name)
        ws.append([])
        for col, label in enumerate(q._LEDGER_EXPECTED_HEADERS, 1):
            ws.cell(row=6, column=col, value=label)
        for col, value in enumerate(('12303000001', 'synthetic only', currency, int(YI), 0, 0, int(YI)), 1):
            ws.cell(row=7, column=col, value=value)
    book.save(ledger_path)
    book.close()
    return ledger_path, avg_path


def test_validated_absent_tuple_preserves_unknown_through_merge_summary_comparison_and_xlsx(tmp_path):
    ledger_path, avg_path = canonical_pair(tmp_path)
    assert validate_qdb_gl_baseline_source(ledger_path).admissible
    assert validate_qdb_gl_baseline_source(avg_path).admissible
    merged = q.merge_all(q.parse_general_ledger(ledger_path), q.parse_daily_avg(avg_path))
    assert merged['3位'][0]['月日均'] is None
    result = workbook(merged, synthetic_merged())
    assert sheet(result, 'summary_3d')['rows'][0]['月日均'] is None
    assert metric(sheet(result, 'company_scale')['rows'], '公司贷款合计')['月日均'] is None
    excel = load_workbook(BytesIO(q.export_qdb_gl_monthly_analysis_workbook_xlsx_bytes(result)), data_only=True)
    try:
        for title, label in (('公司规模', '公司贷款合计'), ('分部基础规模', '公司贷款合计')):
            rows = list(excel[title].values)
            row = next(row for row in rows[1:] if row[0] == label)
            assert row[rows[0].index('月日均')] is None
            assert row[rows[0].index('时点余额')] == 1
            assert row[rows[0].index('口径来源')].startswith('source_missing:')
    finally:
        excel.close()
    # Per-file admissibility does reject a present tuple with a blank amount.
    book = load_workbook(avg_path)
    book['月']['C4'] = None
    book.save(avg_path)
    book.close()
    assert not validate_qdb_gl_baseline_source(avg_path).admissible


def test_unrelated_missing_sources_do_not_taint_healthy_rows_or_formal_markers():
    result = workbook(synthetic_merged({'12303000001': YI, '14203000001': YI}, missing=[('月', '142')]))
    company = metric(sheet(result, 'company_scale')['rows'], '公司贷款合计')
    assert company['月日均'] == 1 and not company['口径来源'].startswith('source_missing:')
    statuses = sheet(result, 'financial_indicator_status')['rows']
    for name in q.FORMAL_FINANCIAL_INDICATOR_PENDING_NAMES:
        row = metric(statuses, name)
        assert row['当前值'] is None
        assert row['口径来源'] == q.FORMAL_FINANCIAL_INDICATOR_PENDING_SOURCE


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
def test_other_period_average_index_is_evidence_of_missing_3d_input(key, builder, label, code, average_code):
    merged = synthetic_merged({'10101000001': YI})
    merged['日均_3位'].append({'科目代码': average_code, '月日均': None, '年日均': YI})
    row = metric(builder(merged), label)
    assert row['月日均'] is None and row['时点余额'] == 0
    # The segment source convention is GL-backed; other scale families read the average index.
    assert row['年日均'] == (0 if key == 'segment_base_scale' else YI)
    assert average_code in row['口径来源']


@pytest.mark.parametrize('builder,label,code,coefficient', [
    (q._company_scale_raw_rows, '公司存款-活期', '20103000001', -1),
    (q._company_scale_raw_rows, '公司存款-活期', '20301000001', -1),
    (q._company_scale_raw_rows, '公司存款-定期', '20250000001', 0),
    (q._company_scale_raw_rows, '公司存款合计', '21601020001', -1),
    (q._company_scale_raw_rows, '公司存款合计', '21601030001', 0),
    (q._company_scale_raw_rows, '公司贷款合计', '13604000001', 0),
    (q._company_scale_raw_rows, '公司贷款合计', '13603000001', 1),
    (q._company_scale_raw_rows, '公司贷款合计', '13003000002', 1),
    (q._retail_scale_raw_rows, '零售存款合计', '20250000001', -1),
    (q._retail_scale_raw_rows, '零售存款合计', '21702000001', -1),
    (q._retail_scale_raw_rows, '零售存款合计', '21602020001', -1),
    (q._retail_scale_raw_rows, '参考：个人贷款合计', '13604000001', 1),
    (q._retail_scale_raw_rows, '参考：个人贷款合计', '13003000001', 1),
    (q._financial_market_scale_raw_rows, '生息债券投资', '14301010001', 0),
    (q._financial_market_scale_raw_rows, '生息债券投资', '14301010002', 0),
    (q._financial_market_scale_raw_rows, '生息债券投资', '14302000001', 1),
    (q._financial_market_scale_raw_rows, '同业资产', '14004000001', 0),
    (q._financial_market_scale_raw_rows, '同业资产', '14005000001', 0),
    (q._financial_market_scale_raw_rows, '同业资产', '14003000001', 1),
    (q._financial_market_scale_raw_rows, '同业负债', '27205000001', -1),
    (q._financial_market_scale_raw_rows, '同业负债', '27206000001', -1),
    (q._segment_base_scale_raw_rows, '公司贷款合计', '13604000001', 0),
    (q._segment_base_scale_raw_rows, '个人贷款合计', '13604000001', 1),
    (q._segment_base_scale_raw_rows, '公司存款合计', '20250000001', 0),
    (q._segment_base_scale_raw_rows, '储蓄存款合计', '20250000001', -1),
])
@pytest.mark.parametrize('value', [YI * D('1.23456'), -YI * D('1.23456')])
def test_existing_signed_component_contracts(builder, label, code, coefficient, value):
    row = metric(builder(synthetic_merged({code: value})), label)
    assert [row[field] for field in ('时点余额', '年日均', '月日均')] == [coefficient * value] * 3


@pytest.mark.parametrize('family,code,label', [
    ('company', '20103000001', '披露：公司存款组件和-合计残差'),
    ('retail', '21103000001', '披露：零售存款组件和-合计残差'),
])
def test_missing_deposit_residual_preserves_source_and_accounting_owner_disclaimer(family, code, label):
    result = workbook(synthetic_merged({code: -YI}, missing=[('月', code[:3])]))
    row = metric(sheet(result, family + '_scale')['rows'], label)
    assert row['时点余额'] == row['年日均'] == 0 and row['月日均'] is None
    assert row['口径来源'].startswith('source_missing:') and code[:3] in row['口径来源']
    assert 'BAL-P1-07' in row['口径来源'] and '待会计owner裁决' in row['口径来源']


def test_month_only_gap_does_not_mark_year_yield_unknown():
    result = workbook(synthetic_merged({'12303000001': YI, '50103000001': -YI / 100}, missing=[('月', '123')]))
    row = metric(sheet(result, 'income_rate_analysis')['rows'], '公司贷款利息收入')
    assert row['年日均规模'] == 1 and row['年化收益率/付息率%'] == 11.77
    assert row['口径来源'] == q.INCOME_RATE_ANALYSIS_SOURCE


def test_nullable_readers_do_not_mutate_merged_sources():
    merged = synthetic_merged(missing=[('月', '123')])
    before = deepcopy(merged)
    workbook(merged, merged)
    assert merged == before


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
def test_comparison_source_uses_relevant_period_and_keeps_known_spot_source(key, builder, label, code, average_code):
    current = synthetic_merged({code: YI}, missing=[('月', average_code)])
    prior = synthetic_merged({code: YI}, missing=[('年', average_code)])
    compare_key = 'segment_scale_compare' if key == 'segment_base_scale' else key + '_compare'
    rows = [row for row in sheet(workbook(current, prior), compare_key)['rows'] if row['指标'] == label]
    year = next(row for row in rows if row['口径'] == '年日均同比')
    assert year['本期'] == 1 and year['对比期'] is None and year['增减额'] is None
    assert '年日均' in year['口径来源'] and '月日均' not in year['口径来源']
    month = next(row for row in rows if row['口径'] == '月日均环比')
    assert month['本期'] is None and month['对比期'] == 1
    assert '月日均' in month['口径来源'] and '年日均' not in month['口径来源']
    for row in rows:
        if row['口径'].startswith('时点'):
            assert row['本期'] == row['对比期'] == 1
            assert not row['口径来源'].startswith('source_missing:')


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
@pytest.mark.parametrize('missing_period,healthy_basis', [('月', '年日均同比'), ('年', '月日均环比')])
def test_other_period_gap_does_not_taint_complete_comparison(
    key, builder, label, code, average_code, missing_period, healthy_basis,
):
    merged = synthetic_merged({code: YI}, missing=[(missing_period, average_code)])
    compare_key = 'segment_scale_compare' if key == 'segment_base_scale' else key + '_compare'
    row = next(
        row for row in sheet(workbook(merged, merged), compare_key)['rows']
        if row['指标'] == label and row['口径'] == healthy_basis
    )
    assert row['本期'] == row['对比期'] == 1
    assert row['增减额'] == row['增减幅%'] == 0
    assert not row['口径来源'].startswith('source_missing:')


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
def test_3d_source_precedence_is_unchanged_when_indexes_disagree(key, builder, label, code, average_code):
    merged = synthetic_merged({code: YI})
    for field in ('月日均', '年日均'):
        next(row for row in merged['3位'] if row['科目代码'] == average_code)[field] = 2 * YI
        next(row for row in merged['日均_3位'] if row['科目代码'] == average_code)[field] = 7 * YI
    row = metric(builder(merged), label)
    expected = (2 if key == 'segment_base_scale' else 7) * YI
    assert row['月日均'] == row['年日均'] == expected
    assert row['时点余额'] == YI


def test_top11_missing_month_average_keeps_existing_threshold_selection():
    merged = synthetic_merged({'12303000001': 2 * YI}, missing=[('月', '12303000001')])
    assert q.build_11d_top_rows(q.compute_deviation(merged['11位'])) == []


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
def test_gl_present_with_both_average_tuples_absent_keeps_both_unknown(key, builder, label, code, average_code):
    row = metric(builder(synthetic_merged({code: YI}, missing=[('月', average_code), ('年', average_code)])), label)
    assert row['时点余额'] == YI and row['月日均'] is None and row['年日均'] is None
    assert average_code in row['口径来源'] and '月日均' in row['口径来源'] and '年日均' in row['口径来源']


@pytest.mark.parametrize('builder,label,code,average_code,expected', [
    (q._company_scale_raw_rows, '公司贷款合计', '13604000001', '13604', YI),
    (q._retail_scale_raw_rows, '参考：信用卡', '13604000001', '13604', D(0)),
    (q._financial_market_scale_raw_rows, '同业资产', '14004000001', '14004', YI),
])
def test_explicit_zero_5d_average_is_known_even_with_nonzero_gl(builder, label, code, average_code, expected):
    merged = synthetic_merged({code: YI})
    next(row for row in merged['日均_5位'] if row['科目代码'] == average_code)['月日均'] = D(0)
    row = metric(builder(merged), label)
    assert row['月日均'] == expected
    assert not row['口径来源'].startswith('source_missing:')


@pytest.mark.parametrize('key,builder,label,code,average_code', FAMILIES)
@pytest.mark.parametrize('period', ['月', '年'])
def test_all_scale_and_comparison_exports_keep_unknown_cells_blank(key, builder, label, code, average_code, period):
    result = workbook(synthetic_merged({code: 2 * YI}, missing=[(period, average_code)]), synthetic_merged({code: YI}))
    scale = sheet(result, key)
    compare = sheet(result, 'segment_scale_compare' if key == 'segment_base_scale' else key + '_compare')
    book = load_workbook(BytesIO(q.export_qdb_gl_monthly_analysis_workbook_xlsx_bytes(result)), data_only=True)
    try:
        rows = list(book[scale['title']].values)
        row = next(row for row in rows[1:] if row[0] == label)
        assert row[rows[0].index(period + '日均')] is None
        assert row[rows[0].index('时点余额')] == 2
        rows = list(book[compare['title']].values)
        row = next(row for row in rows[1:] if row[0] == label and row[1] == ('月日均环比' if period == '月' else '年日均同比'))
        assert all(row[rows[0].index(field)] is None for field in ('本期', '增减额', '增减幅%'))
        assert row[rows[0].index('对比期')] == 1
    finally:
        book.close()


def test_top11_unknown_year_and_trend_export_as_blank_cells():
    result = workbook(synthetic_merged({'12303000001': 2 * YI}, missing=[('年', '12303000001')]))
    book = load_workbook(BytesIO(q.export_qdb_gl_monthly_analysis_workbook_xlsx_bytes(result)), data_only=True)
    try:
        rows = list(book[sheet(result, 'top_11d')['title']].values)
        assert rows[1][rows[0].index('月日均')] == 2
        assert all(rows[1][rows[0].index(field)] is None for field in ('年日均', '趋势额', '趋势%'))
    finally:
        book.close()
