"""Independent XLSX boundary tests: reopen bytes, inspect OOXML, use real producers.

All records are synthetic. These tests establish serialization, not financial
calibration, spreadsheet execution, historical alias schemas, or UI acceptance.
"""

from __future__ import annotations

import io
import json
import socket
import zipfile
from datetime import date
from decimal import Decimal
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook


TOKENS = [
    pytest.param("=1+1", id="equals-formula"),
    pytest.param("#N/A", id="error-na"),
    pytest.param("#REF!", id="error-ref"),
    pytest.param("#VALUE!", id="error-value"),
    pytest.param("#DIV/0!", id="error-div-zero"),
    pytest.param("001234", id="leading-zero-id"),
    pytest.param("1234567890123456789", id="nineteen-digit-id"),
    pytest.param("'原样", id="apostrophe"),
    pytest.param("  中文 合成  ", id="unicode-surrounding-space"),
    pytest.param(" \t ", id="whitespace-only"),
    pytest.param("+1", id="plus-numeric-text"),
    pytest.param("-01", id="minus-numeric-text"),
    pytest.param("@SUM(A1)", id="at-text"),
]

TABLES = (
    ("bond_business_types", "bond_type", "债券持仓", (
        "count", "balance_amount", "share", "weighted_rate_pct", "weighted_term_years",
        "amortized_cost_amount", "market_value_amount", "floating_pnl_amount")),
    ("counterparty_types", "counterparty_type", "同业持仓", (
        "asset_count", "asset_amount", "asset_weighted_rate_pct", "liability_count",
        "liability_amount", "liability_weighted_rate_pct", "net_position_amount")),
    ("maturity_gap", "bucket", "期限分布", (
        "bond_assets_amount", "interbank_assets_amount", "asset_total_amount",
        "issuance_amount", "interbank_liabilities_amount", "full_scope_liability_amount",
        "gap_amount", "full_scope_gap_amount", "cumulative_gap_amount", "asset_share",
        "liability_share", "asset_weighted_rate_pct", "liability_weighted_rate_pct", "spread_bp")),
    ("rate_distribution", "bucket", "利率分布", (
        "bond_count", "bond_amount", "interbank_asset_count", "interbank_asset_amount",
        "interbank_liability_count", "interbank_liability_amount")),
)
SHEETS = ["概览", *(item[2] for item in TABLES)]
LEDGER_NUMBERS = {
    "face_amount", "fair_value", "amortized_cost", "accrued_interest",
    "interest_receivable_payable", "quantity", "latest_face_value", "coupon_rate",
    "yield_to_maturity",
}
XML_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@pytest.fixture(autouse=True)
def deny_network(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("Independent XLSX review forbids all network/provider access")

    for name in ("connect", "connect_ex", "sendto", "sendmsg"):
        if hasattr(socket.socket, name):
            monkeypatch.setattr(socket.socket, name, denied)
    monkeypatch.setattr(socket, "create_connection", denied)
    monkeypatch.setattr(socket, "getaddrinfo", denied)


def _reopen(content):
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml_sheets = [
            ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            for index in range(1, len(workbook.worksheets) + 1)
        ]
    return workbook, dict(zip(workbook.sheetnames, xml_sheets, strict=True))


def _literal(cell, expected, xml):
    assert cell.value == expected, (cell.coordinate, expected, cell.value)
    assert cell.data_type == "s", (cell.coordinate, expected, cell.data_type)
    node = xml.find(f".//s:c[@r='{cell.coordinate}']", XML_NS)
    assert node is not None
    assert node.get("t") in {"s", "inlineStr"}, ET.tostring(node)
    assert node.find("s:f", XML_NS) is None
    if node.get("t") == "inlineStr":
        assert "".join(node.itertext()) == expected


def _no_formula_or_error(xml):
    assert not xml.findall(".//s:f", XML_NS)
    assert not xml.findall(".//s:c[@t='e']", XML_NS)


def _seed_ledger(path, token, *, active=False):
    import duckdb
    from backend.app.repositories.ledger_analytics_repo import POSITION_EXPORT_COLUMNS

    columns = (*POSITION_EXPORT_COLUMNS, "source_version", "rule_version")
    row = dict.fromkeys(POSITION_EXPORT_COLUMNS, "合成")
    row.update(dict.fromkeys(LEDGER_NUMBERS, 0))
    row.update(
        batch_id=31 if active else 29, row_no=1,
        as_of_date="2026-10-01" if active else "2026-09-30",
        position_key="1234567890123456789", bond_code="001234",
        bond_name="ACTIVE MUST NOT LEAK" if active else token,
        counterparty_name_cn="ACTIVE MUST NOT LEAK" if active else token,
        legal_customer_name=None, group_customer_name="",
        portfolio="合成筛选", direction="ASSET", currency=" cny ",
        interest_start_date="2025-01-02", maturity_date="2027-03-04",
        face_amount=90 if active else Decimal("-120.25"), fair_value=0,
        amortized_cost=None, source_version="active" if active else "selected-source",
        rule_version="synthetic-rule",
    )
    with duckdb.connect(str(path)) as conn:
        conn.execute("CREATE TABLE position_snapshot (" + ", ".join(
            f"{key} " + ("BIGINT" if key in {"batch_id", "row_no"}
                        else "DECIMAL(24,8)" if key in LEDGER_NUMBERS else "VARCHAR")
            for key in columns
        ) + ")")
        conn.execute(
            "INSERT INTO position_snapshot VALUES (" + ",".join("?" for _ in columns) + ")",
            [row[key] for key in columns],
        )


@pytest.mark.parametrize("token", TOKENS)
def test_ledger_real_repository_pin_literal_cells(tmp_path, token):
    from backend.app.repositories.duckdb_read_context import DuckDBReadSelection, duckdb_read_scope
    from backend.app.repositories.ledger_analytics_repo import POSITION_EXPORT_COLUMNS
    from backend.app.services.ledger_analytics_service import LedgerAnalyticsService

    active, selected = tmp_path / "active.duckdb", tmp_path / "selected.duckdb"
    _seed_ledger(active, token, active=True)
    _seed_ledger(selected, token)
    filters = {"portfolio": "合成筛选", "currency": "CNY", "bond_code": "001234"}
    service = LedgerAnalyticsService(str(active))
    with duckdb_read_scope(DuckDBReadSelection(active, selected, "synthetic-selected"), required_online=True):
        result = service.repo.list_positions(
            requested_as_of_date="2026-10-01", filters=filters, limit=None, offset=0)
        assert result["items"][0]["bond_name"] == token
        assert result["items"][0]["counterparty_name_cn"] == token
        assert result["items"][0]["currency"] == "CNY"
        filename, content, headers = service.export_positions(
            requested_as_of_date="2026-10-01", filters=filters)
    assert filename == "ledger-positions-2026-09-30.xlsx"
    assert headers == {
        "X-Ledger-Source-Version": "selected-source", "X-Ledger-Rule-Version": "synthetic-rule",
        "X-Ledger-Batch-Id": "29", "X-Ledger-Stale": "true",
        "X-Ledger-Fallback": "true", "X-Ledger-No-Data": "false",
    }
    workbook, xml = _reopen(content)
    try:
        assert workbook.sheetnames == ["positions", "metadata"]
        sheet = workbook["positions"]
        assert tuple(cell.value for cell in sheet[1]) == POSITION_EXPORT_COLUMNS
        for column, key in enumerate(POSITION_EXPORT_COLUMNS, 1):
            _literal(sheet.cell(1, column), key, xml["positions"])
            expected = result["items"][0][key]
            cell = sheet.cell(2, column)
            if expected in (None, ""):
                assert cell.value is None
            elif isinstance(expected, str):
                _literal(cell, expected, xml["positions"])
            else:
                assert cell.value == expected
                assert cell.data_type == "n"
        metadata = {row[0].value: row[1] for row in workbook["metadata"].iter_rows(min_row=2)}
        for key, expected in {
            "source_version": "selected-source", "requested_as_of_date": "2026-10-01",
            "resolved_as_of_date": "2026-09-30", "rule_version": "synthetic-rule",
            "filters": json.dumps(filters, ensure_ascii=False, sort_keys=True),
        }.items():
            _literal(metadata[key], expected, xml["metadata"])
        assert metadata["stale"].value is True and metadata["stale"].data_type == "b"
        assert metadata["batch_id"].value == 29 and metadata["batch_id"].data_type == "n"
        assert metadata["total"].value == 1 and metadata["total"].data_type == "n"
        for tree in xml.values():
            _no_formula_or_error(tree)
    finally:
        workbook.close()


@pytest.mark.parametrize("token", ["=1+1", "#NUM!", "#NULL!", "#NAME?", "#N/A"])
def test_ledger_metadata_only_literal_values(token):
    from backend.app.services.ledger_analytics_service import _positions_workbook, _workbook_bytes

    source = _positions_workbook([], metadata={"source_version": token, "rule_version": token},
                                 filters={"bond_code": token, "portfolio": "001234"})
    content = _workbook_bytes(source)
    source.close()
    workbook, xml = _reopen(content)
    try:
        cells = {row[0].value: row[1] for row in workbook["metadata"].iter_rows(min_row=2)}
        _literal(cells["source_version"], token, xml["metadata"])
        _literal(cells["rule_version"], token, xml["metadata"])
        _literal(cells["filters"], json.dumps({"bond_code": token, "portfolio": "001234"},
                                            ensure_ascii=False, sort_keys=True), xml["metadata"])
        _no_formula_or_error(xml["metadata"])
    finally:
        workbook.close()


@pytest.mark.parametrize("label", ["=1+1", "#N/A", "001234"])
def test_ledger_writer_header_boundary(monkeypatch, label):
    from backend.app.services import ledger_analytics_service as service

    # Supplementary writer case only: current repository headers are constants.
    monkeypatch.setattr(service, "POSITION_EXPORT_COLUMNS", (label,))
    source = service._positions_workbook([{label: "ordinary"}], metadata={}, filters={})
    content = service._workbook_bytes(source)
    source.close()
    workbook, xml = _reopen(content)
    try:
        _literal(workbook["positions"]["A1"], label, xml["positions"])
        for row in workbook["metadata"].iter_rows():
            _literal(row[0], row[0].value, xml["metadata"])
    finally:
        workbook.close()


def test_ledger_no_data_preserves_blank_metadata_and_response(tmp_path):
    from backend.app.services.ledger_analytics_service import LedgerAnalyticsService

    filename, content, headers = LedgerAnalyticsService(str(tmp_path / "absent.duckdb")).export_positions(
        requested_as_of_date="2026-09-30", filters={"bond_code": "=1+1"})
    assert filename == "ledger-positions-2026-09-30.xlsx"
    assert headers == {
        "X-Ledger-Source-Version": "", "X-Ledger-Rule-Version": "", "X-Ledger-Batch-Id": "",
        "X-Ledger-Stale": "false", "X-Ledger-Fallback": "false", "X-Ledger-No-Data": "true",
    }
    workbook, xml = _reopen(content)
    try:
        assert workbook["positions"].max_row == 1
        cells = {row[0].value: row[1] for row in workbook["metadata"].iter_rows(min_row=2)}
        assert cells["batch_id"].value is None
        assert cells["resolved_as_of_date"].value is None
        _literal(cells["requested_as_of_date"], "2026-09-30", xml["metadata"])
        _literal(cells["filters"], '{"bond_code": "=1+1"}', xml["metadata"])
    finally:
        workbook.close()


def _formal_balance(token="普通合成", route="primary"):
    from backend.app.core_finance.balance_analysis import FormalTywBalanceFactRow, FormalZqtzBalanceFactRow
    from backend.app.core_finance.balance_analysis_workbook import build_balance_analysis_workbook_payload
    from backend.app.schemas.balance_analysis import BalanceAnalysisWorkbookPayload

    tyw = FormalTywBalanceFactRow(
        report_date=date(2026, 9, 30), position_id="SYNTHETIC-TYW", product_type="synthetic",
        position_side="asset", counterparty_name=token if route == "fallback" else "unused-fallback",
        account_type="synthetic", special_account_type="", core_customer_type=token if route == "primary" else "",
        invest_type_std="H", accounting_basis="AC", position_scope="asset", currency_basis="CNY",
        currency_code="CNY", principal_amount=Decimal("12500.25"), accrued_interest_amount=Decimal("0"),
        funding_cost_rate=Decimal("2.125"), maturity_date=date(2027, 3, 31),
    )
    bond = FormalZqtzBalanceFactRow(
        report_date=date(2026, 9, 30), instrument_code="SYNTHETIC-ZQTZ", instrument_name="合成国债",
        portfolio_name="合成", cost_center="合成", account_category="合成", asset_class="债券",
        bond_type="国债", issuer_name="合成", industry_name="合成", rating="AAA",
        invest_type_std="H", accounting_basis="AC", position_scope="asset", currency_basis="CNY",
        currency_code="CNY", face_value_amount=Decimal("20000"), market_value_amount=Decimal("19000.5"),
        amortized_cost_amount=Decimal("20000"), accrued_interest_amount=Decimal("0"),
        coupon_rate=None, ytm_value=None, maturity_date=date(2027, 9, 30),
        interest_mode="固定", is_issuance_like=False,
    )
    core = build_balance_analysis_workbook_payload(
        report_date=date(2026, 9, 30), position_scope="all", currency_basis="CNY",
        zqtz_rows=[bond], tyw_rows=[tyw])
    # The actual response schema validates each normal/operational section and
    # performs the same Decimal JSON conversion as the service response path.
    payload = BalanceAnalysisWorkbookPayload.model_validate({
        **{key: core[key] for key in ("report_date", "position_scope", "currency_basis", "cards")},
        "tables": [table for table in core["tables"] if table["section_kind"] == "table"],
        "operational_sections": [table for table in core["tables"] if table["section_kind"] != "table"],
    }).model_dump(mode="json")
    return core, payload


@pytest.mark.parametrize("route", ["primary", "fallback"])
@pytest.mark.parametrize("token", TOKENS)
def test_balance_formal_core_schema_dimension_survives(token, route):
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    core, payload = _formal_balance(token, route)
    before = next(table for table in core["tables"] if table["key"] == "counterparty_types")
    after = next(table for table in payload["tables"] if table["key"] == "counterparty_types")
    assert before["rows"][0]["counterparty_type"] == token
    assert after["rows"][0]["counterparty_type"] == token
    assert isinstance(before["rows"][0]["asset_amount"], Decimal)
    assert isinstance(before["rows"][0]["asset_weighted_rate_pct"], Decimal)
    assert after["rows"][0]["asset_amount"] == "1.250025"
    assert after["rows"][0]["asset_weighted_rate_pct"] == str(before["rows"][0]["asset_weighted_rate_pct"])
    assert all(isinstance(card["value"], Decimal) for card in core["cards"])
    assert all(isinstance(card["value"], str) for card in payload["cards"])
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        assert workbook.sheetnames == SHEETS
        _literal(workbook["同业持仓"]["A2"], token, xml["同业持仓"])
        assert workbook["同业持仓"]["C2"].value == 1.250025
        assert workbook["同业持仓"]["C2"].data_type == "n"
        assert workbook["同业持仓"]["C2"].number_format == "#,##0.00000000"
        assert workbook["同业持仓"]["D2"].value == 2.125
        assert workbook["同业持仓"]["D2"].number_format == "0.00000000"
        for tree in xml.values():
            _no_formula_or_error(tree)
    finally:
        workbook.close()


def test_balance_real_core_all_tables_and_cards_keep_numeric_values_formats():
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    core, payload = _formal_balance()
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        assert workbook.sheetnames == SHEETS
        for key, dimension, name, numeric_keys in TABLES:
            table = next(table for table in payload["tables"] if table["key"] == key)
            assert [column["key"] for column in table["columns"]] == [dimension, *numeric_keys]
            sheet = workbook[name]
            assert sheet.max_row == len(table["rows"]) + 1
            assert [cell.value for cell in sheet[1]] == [column["label"] for column in table["columns"]]
            for row_index, row in enumerate(table["rows"], 2):
                _literal(sheet.cell(row_index, 1), row[dimension], xml[name])
                for column_index, column_key in enumerate(numeric_keys, 2):
                    cell, value = sheet.cell(row_index, column_index), row[column_key]
                    if value is None:
                        assert cell.value is None
                        continue
                    expected = Decimal(str(value))
                    assert cell.data_type == "n" and cell.value == pytest.approx(float(expected))
                    if column_key.endswith("_amount"):
                        assert cell.number_format == "#,##0.00000000"
                    elif expected == expected.to_integral_value():
                        assert cell.number_format == "0"
                    else:
                        assert cell.number_format == "0.00000000"
        for index, card in enumerate(payload["cards"], 2):
            _literal(workbook["概览"].cell(index, 1), card["label"], xml["概览"])
            cell = workbook["概览"].cell(index, 2)
            assert cell.value == pytest.approx(float(Decimal(card["value"])))
            assert cell.data_type == "n" and cell.number_format == "#,##0.00000000"
        assert workbook["债券持仓"]["I2"].value == -0.09995
    finally:
        workbook.close()


def _minimal_payload():
    return {"cards": [], "tables": [
        {"key": key, "title": "not exported", "columns": [{"key": dimension, "label": "维度"}], "rows": []}
        for key, dimension, _, _ in TABLES
    ]}


@pytest.mark.parametrize("table_index", range(4), ids=[item[0] for item in TABLES])
@pytest.mark.parametrize("token", ["001234", "1234567890123456789", " \t ", "=1+1", "#N/A"])
def test_balance_writer_all_known_dimensions(table_index, token):
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    payload = _minimal_payload()
    table = payload["tables"][table_index]
    _, dimension, name, _ = TABLES[table_index]
    table["rows"] = [{dimension: token}]
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        _literal(workbook[name]["A2"], token, xml[name])
    finally:
        workbook.close()


@pytest.mark.parametrize("token", ["=1+1", "#N/A", "#REF!", "001234", "  中文  "])
def test_balance_headers_card_labels_and_nonnumeric_card_values(token):
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    payload = _minimal_payload()
    for table in payload["tables"]:
        table["columns"][0]["label"] = token
    payload["cards"] = [{"label": token, "value": "#VALUE!"}, {"label": "numeric", "value": "001234"}]
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        for name in SHEETS[1:]:
            _literal(workbook[name]["A1"], token, xml[name])
            assert workbook[name].max_row == 1
        for coord, expected in [("A1", "label"), ("B1", "value"), ("A2", token), ("B2", "#VALUE!")]:
            _literal(workbook["概览"][coord], expected, xml["概览"])
        assert workbook["概览"]["B3"].value == 1234
        assert workbook["概览"]["B3"].data_type == "n"
        assert workbook["概览"]["B3"].number_format == "#,##0.00000000"
    finally:
        workbook.close()


@pytest.mark.parametrize("value, expected, data_type", [
    ("=1+1", "=1+1", "s"), ("#N/A", "#N/A", "s"), ("#NUM!", "#NUM!", "s"),
    ("'001234", "'001234", "s"), ("-2.125", -2.125, "n"), ("0", 0, "n"), ("2.000", 2, "n"),
])
def test_balance_card_value_boundary(value, expected, data_type):
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    payload = _minimal_payload()
    payload["cards"] = [{"label": "synthetic", "value": value}]
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        cell = workbook["概览"]["B2"]
        if data_type == "s":
            _literal(cell, expected, xml["概览"])
        else:
            assert cell.value == expected and cell.data_type == data_type
            assert cell.number_format == "#,##0.00000000"
    finally:
        workbook.close()


@pytest.mark.parametrize("table_index", range(3), ids=["zqtz_balance", "tyw_balance", "maturity_distribution"])
@pytest.mark.parametrize("token", ["=1+1", "#DIV/0!", "plain"])
def test_balance_unknown_alias_priority_and_coercion_limit(table_index, token):
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    payload = _minimal_payload()
    aliases = ["zqtz_balance", "tyw_balance", "maturity_distribution"]
    dimension = TABLES[table_index][1]
    payload["tables"][table_index]["rows"] = [{dimension: "CURRENT MUST NOT WIN"}]
    payload["tables"].append({
        "key": aliases[table_index], "columns": [
            {"key": dimension, "label": "historical unknown dimension"},
            {"key": "unknown_text", "label": "literal"},
            {"key": "unknown_amount", "label": "amount"},
        ],
        "rows": [{dimension: "001234", "unknown_text": token, "unknown_amount": "-0.125"}],
    })
    name = TABLES[table_index][2]
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        # No proven alias dimension contract: preserve existing numeric coercion.
        assert workbook[name]["A2"].value == 1234 and workbook[name]["A2"].data_type == "n"
        _literal(workbook[name]["B2"], token, xml[name])
        assert workbook[name]["C2"].value == -0.125
        assert workbook[name]["C2"].number_format == "#,##0.00000000"
    finally:
        workbook.close()


@pytest.mark.parametrize("table_index", range(4), ids=[item[0] for item in TABLES])
def test_balance_empty_and_missing_table_contract(table_index):
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    payload = _minimal_payload()
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        assert workbook.sheetnames == SHEETS
        assert all(workbook[name].max_row == 1 for name in SHEETS)
    finally:
        workbook.close()
    payload["tables"].pop(table_index)
    with pytest.raises(RuntimeError, match="export table unavailable"):
        _build_balance_analysis_workbook_xlsx_bytes(payload)


def test_balance_blank_dimensions_and_numeric_fallback_controls():
    from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes

    payload = _minimal_payload()
    table = payload["tables"][0]
    table["columns"] += [
        {"key": key, "label": key} for key in ("balance_amount", "count", "weighted_rate_pct", "unknown_key")
    ]
    table["rows"] = [
        {"bond_type": None, "balance_amount": "-12.125", "count": "0", "weighted_rate_pct": "2.5", "unknown_key": "001234"},
        {"bond_type": "", "balance_amount": None, "count": 0, "weighted_rate_pct": None, "unknown_key": True},
        {"bond_type": "normal", "balance_amount": "0", "count": "2", "weighted_rate_pct": "2.000", "unknown_key": "#NAME?"},
    ]
    workbook, xml = _reopen(_build_balance_analysis_workbook_xlsx_bytes(payload))
    try:
        sheet = workbook["债券持仓"]
        assert sheet["A2"].value is None and sheet["A3"].value is None
        assert sheet["B2"].value == -12.125 and sheet["B2"].number_format == "#,##0.00000000"
        assert sheet["C2"].value == 0 and sheet["C2"].number_format == "0"
        assert sheet["D2"].value == 2.5 and sheet["D2"].number_format == "0.00000000"
        assert sheet["E2"].value == 1234 and sheet["E2"].data_type == "n"
        assert sheet["B3"].value is None and sheet["D3"].value is None
        assert sheet["E3"].value is True and sheet["E3"].data_type == "b"
        assert sheet["B4"].value == 0 and sheet["B4"].number_format == "#,##0.00000000"
        assert sheet["D4"].value == 2 and sheet["D4"].number_format == "0"
        _literal(sheet["E4"], "#NAME?", xml["债券持仓"])
    finally:
        workbook.close()
