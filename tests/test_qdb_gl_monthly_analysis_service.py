from __future__ import annotations

from threading import Event, Lock, Thread

from openpyxl import load_workbook

from tests.helpers import load_module
from tests.test_qdb_gl_monthly_analysis_core import _write_month_pair

import pytest

pytestmark = [
    pytest.mark.excluded_surface_acceptance,
    pytest.mark.surface_qdb_gl,
]

def test_service_discovers_available_report_months(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    envelope = module.qdb_gl_monthly_analysis_dates_envelope(source_dir=str(source_dir))

    assert envelope["result_meta"]["basis"] == "analytical"
    assert envelope["result_meta"]["formal_use_allowed"] is False
    assert envelope["result"]["report_months"] == ["202602"]

def test_service_blocks_generation_when_input_contract_validation_fails(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    source_dir.mkdir(parents=True)
    avg_path, _ledger_path = _write_month_pair(source_dir, "202602")
    avg_path.rename(source_dir / "日均错误202602.xlsx")

    try:
        module.qdb_gl_monthly_analysis_workbook_envelope(
            source_dir=str(source_dir),
            report_month="202602",
        )
    except ValueError as exc:
        assert "202602" in str(exc)
    else:
        raise AssertionError("Expected month-pair generation to fail when the canonical source file is missing.")

def test_service_workbook_envelope_includes_qdb_source_evidence_metadata(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_鎬昏处瀵硅处-鏃ュ潎"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    envelope = module.qdb_gl_monthly_analysis_workbook_envelope(
        source_dir=str(source_dir),
        report_month="202602",
    )

    meta = envelope["result_meta"]
    assert meta["tables_used"] == [
        "qdb_gl_average_balance_workbook",
        "qdb_gl_ledger_reconciliation_workbook",
    ]
    assert meta["evidence_rows"] == 2
    assert meta["filters_applied"] == {
        "report_month": "202602",
        "comparison_months": {
            "prior_month": {"report_month": "202601", "status": "missing"},
            "two_months_ago": {"report_month": "202512", "status": "missing"},
            "prior_year": {"report_month": "202502", "status": "missing"},
        },
    }
    assert meta["requested_report_date"] == "202602"
    assert meta["resolved_report_date"] == "202602"
    assert meta["date_basis"] == "qdb_gl_monthly_analysis_report_month"

def test_service_blocks_workbook_when_rebuilt_payload_resolves_different_month(tmp_path, monkeypatch):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_鎬昏处瀵硅处-鏃ュ潎"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")
    original_build = module.build_qdb_gl_monthly_analysis_workbook

    def build_wrong_month_workbook(**kwargs):
        payload = original_build(**kwargs)
        payload["report_month"] = "202601"
        return payload

    monkeypatch.setattr(module, "build_qdb_gl_monthly_analysis_workbook", build_wrong_month_workbook)

    try:
        module.qdb_gl_monthly_analysis_workbook_envelope(
            source_dir=str(source_dir),
            report_month="202602",
        )
    except ValueError as exc:
        assert "202602" in str(exc)
        assert "202601" in str(exc)
    else:
        raise AssertionError("Expected workbook generation to fail when the rebuilt report_month mismatches the request.")

def test_service_blocks_workbook_when_source_evidence_lacks_table_anchor(tmp_path, monkeypatch):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_鎬昏处瀵硅处-鏃ュ潎"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")
    original_validate = module.validate_qdb_gl_baseline_source

    def validate_without_average_table_anchor(path):
        evidence = original_validate(path)
        if evidence.source_kind == "average_balance":
            return evidence.model_copy(update={"source_kind": "unknown"})
        return evidence

    monkeypatch.setattr(module, "validate_qdb_gl_baseline_source", validate_without_average_table_anchor)

    try:
        module.qdb_gl_monthly_analysis_workbook_envelope(
            source_dir=str(source_dir),
            report_month="202602",
        )
    except ValueError as exc:
        assert "source evidence" in str(exc)
        assert "202602" in str(exc)
    else:
        raise AssertionError("Expected workbook generation to fail when QDB source evidence lacks table anchors.")

def test_service_workbook_envelope_includes_segment_scale_compare_when_history_exists(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202601")
    _write_month_pair(source_dir, "202602")

    envelope = module.qdb_gl_monthly_analysis_workbook_envelope(
        source_dir=str(source_dir),
        report_month="202602",
    )

    assert "prior_month:202601" in envelope["result_meta"]["source_version"]
    assert envelope["result_meta"]["evidence_rows"] == 4
    assert envelope["result_meta"]["tables_used"] == [
        "qdb_gl_average_balance_workbook",
        "qdb_gl_ledger_reconciliation_workbook",
    ]
    assert envelope["result_meta"]["filters_applied"]["comparison_months"]["prior_month"] == {
        "report_month": "202601",
        "status": "loaded",
    }
    assert envelope["result_meta"]["filters_applied"]["comparison_months"]["two_months_ago"] == {
        "report_month": "202512",
        "status": "missing",
    }
    segment_sheet = next(
        sheet for sheet in envelope["result"]["sheets"] if sheet["key"] == "segment_scale_compare"
    )
    assert segment_sheet["title"] == "分部规模同比环比"
    assert segment_sheet["columns"] == ["指标", "口径", "本期", "对比期", "增减额", "增减幅%", "口径来源"]
    assert any(row["口径"] == "时点环比" for row in segment_sheet["rows"])
    market_sheet = next(
        sheet for sheet in envelope["result"]["sheets"] if sheet["key"] == "financial_market_scale_compare"
    )
    assert market_sheet["title"] == "金融市场规模同比环比"
    assert market_sheet["columns"] == ["指标", "口径", "本期", "对比期", "增减额", "增减幅%", "口径来源"]
    assert any(row["指标"] == "同业负债" and row["口径"] == "月日均环比" for row in market_sheet["rows"])
    company_sheet = next(
        sheet for sheet in envelope["result"]["sheets"] if sheet["key"] == "company_scale_compare"
    )
    assert company_sheet["title"] == "公司规模同比环比"
    assert company_sheet["columns"] == ["指标", "口径", "本期", "对比期", "增减额", "增减幅%", "口径来源"]
    assert any(row["指标"] == "公司贷款合计" and row["口径"] == "时点环比" for row in company_sheet["rows"])
    retail_sheet = next(
        sheet for sheet in envelope["result"]["sheets"] if sheet["key"] == "retail_scale_compare"
    )
    assert retail_sheet["title"] == "零售规模同比环比"
    assert retail_sheet["columns"] == ["指标", "口径", "本期", "对比期", "增减额", "增减幅%", "口径来源"]
    assert any(row["指标"] == "零售存款合计" and row["口径"] == "时点环比" for row in retail_sheet["rows"])
    income_sheet = next(
        sheet for sheet in envelope["result"]["sheets"] if sheet["key"] == "income_rate_analysis"
    )
    assert income_sheet["title"] == "收益率分析（总账可复算）"
    assert any(row["指标"] == "公司贷款利息收入" for row in income_sheet["rows"])

def test_service_workbook_marks_invalid_comparison_month_without_blocking_requested_month(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_鎬昏处瀵硅处-鏃ュ潎"
    source_dir.mkdir(parents=True)
    _avg_path, ledger_path = _write_month_pair(source_dir, "202601")
    _write_month_pair(source_dir, "202602")
    workbook = load_workbook(ledger_path)
    try:
        workbook.active["A6"] = "invalid-header"
        workbook.save(ledger_path)
    finally:
        workbook.close()

    envelope = module.qdb_gl_monthly_analysis_workbook_envelope(
        source_dir=str(source_dir),
        report_month="202602",
    )

    assert envelope["result"]["report_month"] == "202602"
    assert envelope["result_meta"]["filters_applied"]["comparison_months"]["prior_month"] == {
        "report_month": "202601",
        "status": "invalid",
    }
    assert "prior_month:202601" not in envelope["result_meta"]["source_version"]

def test_service_supports_sync_refresh_and_status_flow(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    governance_dir = tmp_path / "governance"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    refresh_payload = module.refresh_qdb_gl_monthly_analysis(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
    )

    assert refresh_payload["status"] in {"queued", "completed"}
    assert refresh_payload["job_name"] == "qdb_gl_monthly_analysis"
    assert refresh_payload["source_version"].startswith("sv_qdb_gl_")
    assert refresh_payload["source_version"] != "202602"
    assert refresh_payload["report_date"] == "202602"
    assert refresh_payload["sheet_count"] > 0
    assert refresh_payload["tables_used"] == [
        "qdb_gl_average_balance_workbook",
        "qdb_gl_ledger_reconciliation_workbook",
    ]
    assert refresh_payload["evidence_rows"] == 2
    assert refresh_payload["comparison_months"]["prior_month"] == {
        "report_month": "202601",
        "status": "missing",
    }
    status_payload = module.qdb_gl_monthly_analysis_refresh_status(
        governance_dir=str(governance_dir),
        run_id=refresh_payload["run_id"],
    )
    assert status_payload["run_id"] == refresh_payload["run_id"]
    assert status_payload["source_version"] == refresh_payload["source_version"]
    assert status_payload["report_date"] == "202602"

def test_service_refresh_uses_unique_run_id_for_new_attempts(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_鎬昏处瀵硅处-鏃ゅ潎"
    governance_dir = tmp_path / "governance"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    first_payload = module.refresh_qdb_gl_monthly_analysis(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
        idempotency_key="qdb-gl-refresh-first",
    )
    second_payload = module.refresh_qdb_gl_monthly_analysis(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
        idempotency_key="qdb-gl-refresh-second",
    )

    assert first_payload["run_id"] != second_payload["run_id"]
    assert first_payload["idempotency_replay"] is False
    assert second_payload["idempotency_replay"] is False
    assert module.qdb_gl_monthly_analysis_refresh_status(
        governance_dir=str(governance_dir),
        run_id=first_payload["run_id"],
    )["run_id"] == first_payload["run_id"]

def test_service_refresh_serializes_same_idempotency_key_concurrently(tmp_path, monkeypatch):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_鎬昏处瀵硅处-鏃ゅ潎"
    governance_dir = tmp_path / "governance"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    original_rebuild = module._rebuild_workbook_payload
    first_rebuild_entered = Event()
    release_first_rebuild = Event()
    calls_lock = Lock()
    results_lock = Lock()
    calls: list[dict[str, object]] = []
    results: list[dict[str, object]] = []
    errors: list[BaseException] = []

    def slow_rebuild(**kwargs):
        with calls_lock:
            calls.append(kwargs)
            call_index = len(calls)
        if call_index == 1:
            first_rebuild_entered.set()
            if not release_first_rebuild.wait(timeout=5):
                raise TimeoutError("Timed out waiting to release first QDB GL rebuild")
        return original_rebuild(**kwargs)

    monkeypatch.setattr(module, "_rebuild_workbook_payload", slow_rebuild)

    def invoke_refresh() -> None:
        try:
            payload = module.refresh_qdb_gl_monthly_analysis(
                source_dir=str(source_dir),
                governance_dir=str(governance_dir),
                report_month="202602",
                idempotency_key="qdb-gl-refresh-concurrent",
            )
        except BaseException as exc:
            with results_lock:
                errors.append(exc)
            return
        with results_lock:
            results.append(payload)

    first_thread = Thread(target=invoke_refresh)
    second_thread = Thread(target=invoke_refresh)
    first_thread.start()
    assert first_rebuild_entered.wait(timeout=2)
    second_thread.start()
    release_first_rebuild.set()
    first_thread.join(timeout=5)
    second_thread.join(timeout=5)

    assert not first_thread.is_alive()
    assert not second_thread.is_alive()
    assert errors == []
    assert len(results) == 2
    assert len(calls) == 1
    assert {payload["run_id"] for payload in results} == {results[0]["run_id"]}
    replay_flags = [payload["idempotency_replay"] for payload in results]
    assert replay_flags.count(False) == 1
    assert replay_flags.count(True) == 1

def test_service_refresh_records_failed_build_when_requested_month_rebuild_fails(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_鎬昏处瀵硅处-鏃ュ潎"
    governance_dir = tmp_path / "governance"
    source_dir.mkdir(parents=True)
    _avg_path, ledger_path = _write_month_pair(source_dir, "202602")
    workbook = load_workbook(ledger_path)
    try:
        workbook.active["A6"] = "invalid-header"
        workbook.save(ledger_path)
    finally:
        workbook.close()

    refresh_payload = module.refresh_qdb_gl_monthly_analysis(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
    )

    assert refresh_payload["status"] == "failed"
    assert str(refresh_payload["run_id"]).startswith("qdb_gl_monthly_analysis:202602:")
    assert refresh_payload["report_date"] == "202602"
    assert refresh_payload["failure_category"] == "qdb_gl_monthly_analysis_build"
    assert "202602" in refresh_payload["error_message"]
    status_payload = module.qdb_gl_monthly_analysis_refresh_status(
        governance_dir=str(governance_dir),
        run_id=refresh_payload["run_id"],
    )
    assert status_payload["status"] == "failed"
    assert status_payload["failure_category"] == "qdb_gl_monthly_analysis_build"
    assert status_payload["error_message"] == refresh_payload["error_message"]

def test_service_scenario_envelope_uses_analytical_basis_and_override_summary(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    envelope = module.qdb_gl_monthly_analysis_scenario_envelope(
        source_dir=str(source_dir),
        report_month="202602",
        scenario_name="threshold-stress",
        threshold_overrides={
            "DEVIATION_WARN": 6,
            "DEVIATION_ALERT": 12,
        },
    )

    assert envelope["result_meta"]["basis"] == "analytical"
    assert envelope["result_meta"]["formal_use_allowed"] is False
    assert envelope["result_meta"]["result_kind"] == "qdb-gl-monthly-analysis.scenario"
    assert envelope["result"]["scenario_name"] == "threshold-stress"
    assert envelope["result"]["applied_overrides"] == {
        "DEVIATION_WARN": 6,
        "DEVIATION_ALERT": 12,
    }

def test_service_workbook_envelope_applies_approved_analysis_adjustments(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    governance_dir = tmp_path / "governance"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    module.create_qdb_gl_monthly_analysis_manual_adjustment(
        governance_dir=str(governance_dir),
        payload={
            "report_month": "202602",
            "adjustment_class": "analysis_adjustment",
            "target": {
                "section_key": "alerts",
                "row_key": "14001000001",
                "metric_key": "alert_level",
            },
            "operator": "OVERRIDE",
            "value": "manual_override",
            "approval_status": "approved",
        },
    )

    envelope = module.qdb_gl_monthly_analysis_workbook_envelope(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
    )

    alerts_sheet = next(
        sheet for sheet in envelope["result"]["sheets"] if sheet["key"] == "alerts"
    )
    code_key = alerts_sheet["columns"][0]
    level_key = alerts_sheet["columns"][2]
    target_row = next(row for row in alerts_sheet["rows"] if str(row[code_key]) == "14001000001")

    assert target_row[level_key] == "manual_override"

def test_service_workbook_envelope_applies_approved_mapping_adjustments(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    governance_dir = tmp_path / "governance"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    module.create_qdb_gl_monthly_analysis_manual_adjustment(
        governance_dir=str(governance_dir),
        payload={
            "report_month": "202602",
            "adjustment_class": "mapping_adjustment",
            "target": {
                "account_code": "14001000001",
                "field": "account_name",
            },
            "operator": "OVERRIDE",
            "value": "买入返售-人工修正",
            "approval_status": "approved",
        },
    )

    envelope = module.qdb_gl_monthly_analysis_workbook_envelope(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
    )

    alerts_sheet = next(
        sheet for sheet in envelope["result"]["sheets"] if sheet["key"] == "alerts"
    )
    code_key = alerts_sheet["columns"][0]
    name_key = alerts_sheet["columns"][1]
    target_row = next(row for row in alerts_sheet["rows"] if str(row[code_key]) == "14001000001")

    assert target_row[name_key] == "买入返售-人工修正"

def test_service_scenario_envelope_returns_rebuilt_workbook_payload_with_override_effects(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    source_dir = tmp_path / "data_input" / "pnl_总账对账-日均"
    governance_dir = tmp_path / "governance"
    source_dir.mkdir(parents=True)
    _write_month_pair(source_dir, "202602")

    baseline = module.qdb_gl_monthly_analysis_workbook_envelope(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
    )
    scenario = module.qdb_gl_monthly_analysis_scenario_envelope(
        source_dir=str(source_dir),
        governance_dir=str(governance_dir),
        report_month="202602",
        scenario_name="threshold-stress",
        threshold_overrides={
            "DEVIATION_WARN": 80,
            "DEVIATION_ALERT": 90,
            "DEVIATION_CRITICAL": 100,
        },
    )

    baseline_alerts = next(
        sheet for sheet in baseline["result"]["sheets"] if sheet["key"] == "alerts"
    )["rows"]
    scenario_alerts = next(
        sheet for sheet in scenario["result"]["sheets"] if sheet["key"] == "alerts"
    )["rows"]

    assert scenario["result"]["scenario_name"] == "threshold-stress"
    assert scenario["result"]["report_month"] == "202602"
    assert [sheet["key"] for sheet in scenario["result"]["sheets"]] == [
        "overview",
        "financial_indicator_status",
        "summary_3d",
        "asset_structure",
        "liability_structure",
        "loan_industry",
        "deposit_demand_industry",
        "deposit_term_industry",
        "industry_gap",
        "top_11d",
        "alerts",
        "foreign_currency",
        "segment_base_scale",
        "company_scale",
        "retail_scale",
        "financial_market_scale",
        "income_rate_analysis",
        "deposit_interest_split",
        "parent_company_revenue_components",
    ]
    assert scenario["result"]["applied_overrides"] == {
        "DEVIATION_WARN": 80,
        "DEVIATION_ALERT": 90,
        "DEVIATION_CRITICAL": 100,
    }
    assert len(baseline_alerts) > len(scenario_alerts)

def test_service_supports_branch_specific_manual_adjustment_audit_flow(tmp_path):
    module = load_module(
        "backend.app.services.qdb_gl_monthly_analysis_service",
        "backend/app/services/qdb_gl_monthly_analysis_service.py",
    )
    governance_dir = tmp_path / "governance"

    created = module.create_qdb_gl_monthly_analysis_manual_adjustment(
        governance_dir=str(governance_dir),
        payload={
            "report_month": "202602",
            "adjustment_class": "mapping_adjustment",
            "target": {"account_code": "12301", "field": "industry_name"},
            "operator": "OVERRIDE",
            "value": "农业",
            "approval_status": "approved",
        },
    )

    assert created["stream"] == "monthly_operating_analysis_adjustments"
    assert created["adjustment_class"] == "mapping_adjustment"

    listed = module.list_qdb_gl_monthly_analysis_manual_adjustments(
        governance_dir=str(governance_dir),
        report_month="202602",
    )
    assert listed["adjustment_count"] == 1
    assert listed["adjustments"][0]["adjustment_id"] == created["adjustment_id"]

    edited = module.update_qdb_gl_monthly_analysis_manual_adjustment(
        governance_dir=str(governance_dir),
        adjustment_id=created["adjustment_id"],
        payload={
            "report_month": "202602",
            "adjustment_class": "analysis_adjustment",
            "target": {"section_key": "alerts", "row_key": "14001000001", "metric_key": "alert_level"},
            "operator": "OVERRIDE",
            "value": "中度",
            "approval_status": "approved",
        },
    )
    assert edited["adjustment_class"] == "analysis_adjustment"

    revoked = module.revoke_qdb_gl_monthly_analysis_manual_adjustment(
        governance_dir=str(governance_dir),
        adjustment_id=created["adjustment_id"],
    )
    assert revoked["approval_status"] == "rejected"

    restored = module.restore_qdb_gl_monthly_analysis_manual_adjustment(
        governance_dir=str(governance_dir),
        adjustment_id=created["adjustment_id"],
    )
    assert restored["approval_status"] == "approved"

    filename, content = module.export_qdb_gl_monthly_analysis_manual_adjustments_csv(
        governance_dir=str(governance_dir),
        report_month="202602",
    )
    assert filename == "monthly-operating-analysis-audit-202602.csv"
    assert "analysis_adjustment" in content
