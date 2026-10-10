"""BD-026: preserve exporter-input text in the real synthetic Ledger XLSX path.

These repository/service contracts do not start the app, evaluate spreadsheet
formulas, or establish real-data reconciliation or HTTP authorization coverage.
"""
from __future__ import annotations

import io
import json
import socket
import xml.etree.ElementTree as ET
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

import duckdb
import pytest
from openpyxl import load_workbook

from backend.app.governance.ledger_classification import LEDGER_CLASSIFICATION_RULE_VERSION
from backend.app.repositories.duckdb_read_context import DuckDBReadSelection, duckdb_read_scope
from backend.app.repositories.ledger_analytics_repo import POSITION_EXPORT_COLUMNS
from backend.app.services import ledger_analytics_service as ledger

pytestmark = pytest.mark.integration

REPORT_DATE = "2026-09-30"
REQUESTED_DATE = "2026-10-01"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
NUMBERS = {
    "face_amount": Decimal("0"),
    "fair_value": Decimal("-123.456"),
    "amortized_cost": None,
    "accrued_interest": Decimal("1.25"),
    "interest_receivable_payable": Decimal("0"),
    "quantity": Decimal("-2"),
    "latest_face_value": Decimal("100"),
    "coupon_rate": Decimal("0"),
    "yield_to_maturity": Decimal("-1.125"),
}
TEXT_CASES = [
    pytest.param("=1+1", id="formula-looking"),
    pytest.param("#NULL!", id="error-null"),
    pytest.param("#DIV/0!", id="error-div0"),
    pytest.param("#VALUE!", id="error-value"),
    pytest.param("#REF!", id="error-ref"),
    pytest.param("#NAME?", id="error-name"),
    pytest.param("#NUM!", id="error-num"),
    pytest.param("#N/A", id="error-na"),
    pytest.param("中文合成名称", id="chinese"),
    pytest.param("'001234", id="leading-apostrophe"),
    pytest.param("001234", id="leading-zero"),
    pytest.param("1234567890123456789", id="nineteen-digit"),
    pytest.param("  保留空白 \t", id="whitespace"),
    pytest.param("+1+1", id="plus-prefix"),
    pytest.param("-1+1", id="minus-prefix"),
    pytest.param("@SUM(A1:A2)", id="at-prefix"),
    pytest.param("", id="empty"),
    pytest.param(None, id="null"),
]


@pytest.fixture(autouse=True)
def _deny_network(monkeypatch):
    def blocked(*args, **kwargs):
        raise AssertionError("Network access is forbidden in synthetic Ledger XLSX tests")

    for name in ("connect", "connect_ex", "sendto"):
        monkeypatch.setattr(socket.socket, name, blocked)
    for name in ("create_connection", "getaddrinfo"):
        monkeypatch.setattr(socket, name, blocked)


def _seed(path: Path, *, text, active=False, source="sv_synthetic_G1",
          rule=LEDGER_CLASSIFICATION_RULE_VERSION):
    columns = (*POSITION_EXPORT_COLUMNS, "source_version", "rule_version")
    with duckdb.connect(str(path)) as conn:
        conn.execute("CREATE TABLE position_snapshot (" + ", ".join(
            f"{column} {'BIGINT' if column in {'batch_id', 'row_no'} else 'DECIMAL(24,8)' if column in NUMBERS else 'VARCHAR'}"
            for column in columns
        ) + ")")
        # Insert out of order to exercise the repository's row_no ordering.
        for row_no in (2, 1):
            values = {column: "synthetic" for column in POSITION_EXPORT_COLUMNS}
            values.update(NUMBERS)
            values.update(
                position_key=f"synthetic-{row_no}", batch_id=2 if active else 1,
                row_no=row_no, as_of_date=REQUESTED_DATE if active else REPORT_DATE,
                bond_code="001234", bond_name=text if row_no == 1 else None,
                counterparty_name_cn=text if row_no == 1 else "",
                legal_customer_name=None, group_customer_name="",
                portfolio="合成组合", direction="ASSET", currency=" cny ",
                interest_start_date="2026-01-01", maturity_date="2027-09-30",
                source_version=source, rule_version=rule,
            )
            if active:
                values.update(face_amount=Decimal("987"), bond_name="ACTIVE ONLY",
                              counterparty_name_cn="ACTIVE ONLY")
            conn.execute(
                "INSERT INTO position_snapshot VALUES (" + ", ".join("?" for _ in columns) + ")",
                [values[column] for column in columns],
            )


def _reopen(content, tmp_path):
    (tmp_path / "export.xlsx").write_bytes(content)
    with ZipFile(io.BytesIO(content)) as archive:
        roots = {}
        for name in archive.namelist():
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                xml = archive.read(name)
                (tmp_path / name.rsplit("/", 1)[1]).write_bytes(xml)
                roots[name] = ET.fromstring(xml)
    return load_workbook(io.BytesIO(content), data_only=False), roots


def _assert_cell(cell, expected, xml):
    if expected is None or expected == "":
        # openpyxl deliberately serializes empty strings as blank cells.
        assert cell.value is None
        assert cell.data_type not in {"f", "e"}
    elif isinstance(expected, str):
        assert cell.value == expected
        assert cell.data_type == "s", (cell.coordinate, expected, cell.data_type)
        node = xml.find(f".//m:c[@r='{cell.coordinate}']", NS)
        assert node is not None
        assert node.get("t") in {"inlineStr", "s"}
        assert node.find("m:f", NS) is None
        if node.get("t") == "inlineStr":
            assert "".join(node.itertext()) == expected
    else:
        assert cell.value == expected
        assert cell.data_type == ("b" if isinstance(expected, bool) else "n")


def _assert_metadata(workbook, roots, expected):
    sheet = workbook["metadata"]
    xml = roots["xl/worksheets/sheet2.xml"]
    assert list(sheet.values)[0] == ("key", "value")
    assert [row[0] for row in list(sheet.values)[1:]] == list(expected)
    for index, (key, value) in enumerate(expected.items(), start=2):
        _assert_cell(sheet.cell(index, 1), key, xml)
        _assert_cell(sheet.cell(index, 2), value, xml)
    _assert_cell(sheet["A1"], "key", xml)
    _assert_cell(sheet["B1"], "value", xml)


@pytest.mark.parametrize("text", TEXT_CASES)
def test_repository_service_keeps_literal_text_and_selected_generation(tmp_path, text):
    active, snapshot = tmp_path / "active.duckdb", tmp_path / "g1.duckdb"
    _seed(active, text="ACTIVE ONLY", active=True, source="sv_synthetic_G2")
    _seed(snapshot, text=text)
    service = ledger.LedgerAnalyticsService(str(active))
    filters = {"portfolio": "合成组合", "currency": "CNY", "cost_center": None}
    selection = DuckDBReadSelection(active, snapshot, "G1")
    with duckdb_read_scope(selection, required_online=True):
        result = service.repo.list_positions(
            requested_as_of_date=REQUESTED_DATE, filters=filters, limit=None, offset=0,
        )
        assert result is not None
        assert result["source_version"] == "sv_synthetic_G1"
        assert result["as_of_date"] == REPORT_DATE
        assert result["items"][0]["bond_name"] == text
        assert result["items"][0]["counterparty_name_cn"] == text
        assert result["items"][0]["currency"] == "CNY"
        assert result["items"][0]["face_amount"] == 0.0
        filename, content, headers = service.export_positions(
            requested_as_of_date=REQUESTED_DATE, filters=filters,
        )
    assert service.repo.list_positions(
        requested_as_of_date=REQUESTED_DATE, filters=filters, limit=None, offset=0,
    )["items"][0]["bond_name"] == "ACTIVE ONLY"
    assert filename == f"ledger-positions-{REPORT_DATE}.xlsx"
    assert headers == {
        "X-Ledger-Source-Version": "sv_synthetic_G1",
        "X-Ledger-Rule-Version": LEDGER_CLASSIFICATION_RULE_VERSION,
        "X-Ledger-Batch-Id": "1", "X-Ledger-Stale": "true",
        "X-Ledger-Fallback": "true", "X-Ledger-No-Data": "false",
    }
    workbook, roots = _reopen(content, tmp_path)
    try:
        assert workbook.sheetnames == ["positions", "metadata"]
        sheet, xml = workbook["positions"], roots["xl/worksheets/sheet1.xml"]
        assert sheet.max_row == 3
        assert tuple(cell.value for cell in sheet[1]) == POSITION_EXPORT_COLUMNS
        for index, column in enumerate(POSITION_EXPORT_COLUMNS, start=1):
            _assert_cell(sheet.cell(1, index), column, xml)
        # Check numeric/date/null controls independently of dangerous text cells.
        for row_index, row in enumerate(result["items"], start=2):
            assert row["row_no"] == row_index - 1
            for column in (*NUMBERS, "batch_id", "row_no", "as_of_date",
                           "interest_start_date", "maturity_date"):
                _assert_cell(sheet.cell(row_index, POSITION_EXPORT_COLUMNS.index(column) + 1),
                             row[column], xml)
        _assert_metadata(workbook, roots, {
            "batch_id": 1, "requested_as_of_date": REQUESTED_DATE,
            "resolved_as_of_date": REPORT_DATE, "as_of_date": REPORT_DATE,
            "source_version": "sv_synthetic_G1", "rule_version": LEDGER_CLASSIFICATION_RULE_VERSION,
            "stale": True, "fallback": True, "total": 2,
            "filters": json.dumps(filters, ensure_ascii=False, sort_keys=True),
        })
        for row_index, row in enumerate(result["items"], start=2):
            for column_index, column in enumerate(POSITION_EXPORT_COLUMNS, start=1):
                _assert_cell(sheet.cell(row_index, column_index), row[column], xml)
        assert all(not root.findall(".//m:f", NS) for root in roots.values())
    finally:
        workbook.close()


@pytest.mark.parametrize("text", ["=1+1", "#N/A", "001234", "  中文元数据  "])
def test_real_repository_metadata_is_literal_and_headers_unchanged(tmp_path, text):
    database = tmp_path / "metadata.duckdb"
    _seed(database, text="synthetic name", source=text, rule=text)
    filters = {"bond_code": "001234", "portfolio": "合成组合", "cost_center": None}
    filename, content, headers = ledger.LedgerAnalyticsService(str(database)).export_positions(
        requested_as_of_date=REPORT_DATE, filters=filters,
    )
    assert filename == f"ledger-positions-{REPORT_DATE}.xlsx"
    assert headers == {
        "X-Ledger-Source-Version": text, "X-Ledger-Rule-Version": text,
        "X-Ledger-Batch-Id": "1", "X-Ledger-Stale": "false",
        "X-Ledger-Fallback": "false", "X-Ledger-No-Data": "false",
    }
    workbook, roots = _reopen(content, tmp_path)
    try:
        _assert_metadata(workbook, roots, {
            "batch_id": 1, "requested_as_of_date": REPORT_DATE,
            "resolved_as_of_date": REPORT_DATE, "as_of_date": REPORT_DATE,
            "source_version": text, "rule_version": text,
            "stale": False, "fallback": False, "total": 2,
            "filters": json.dumps(filters, ensure_ascii=False, sort_keys=True),
        })
    finally:
        workbook.close()


@pytest.mark.parametrize("header", ["=1+1", "#N/A", "001234"])
def test_builder_header_text_is_literal(tmp_path, monkeypatch, header):
    # Supplementary writer seam: repository column labels are fixed in production.
    monkeypatch.setattr(ledger, "POSITION_EXPORT_COLUMNS", (header, "control"))
    original = ledger._positions_workbook([{header: header, "control": 0}], metadata={}, filters={})
    try:
        content = ledger._workbook_bytes(original)
    finally:
        original.close()
    workbook, roots = _reopen(content, tmp_path)
    try:
        sheet, xml = workbook["positions"], roots["xl/worksheets/sheet1.xml"]
        _assert_cell(sheet["A1"], header, xml)
        _assert_cell(sheet["A2"], header, xml)
        _assert_cell(sheet["B2"], 0, xml)
    finally:
        workbook.close()


@pytest.mark.parametrize("missing_database", [True, False], ids=["missing-database", "no-matching-rows"])
def test_empty_export_keeps_shape_blank_metadata_and_response_contract(tmp_path, missing_database):
    database = tmp_path / "empty.duckdb"
    if not missing_database:
        _seed(database, text="synthetic name")
    filters = {"bond_code": "#N/A", "portfolio": "=1+1", "cost_center": "中文筛选"}
    filename, content, headers = ledger.LedgerAnalyticsService(str(database)).export_positions(
        requested_as_of_date=REPORT_DATE, filters=filters,
    )
    assert filename == f"ledger-positions-{REPORT_DATE}.xlsx"
    assert headers == {
        "X-Ledger-Source-Version": "" if missing_database else "sv_synthetic_G1",
        "X-Ledger-Rule-Version": "" if missing_database else LEDGER_CLASSIFICATION_RULE_VERSION,
        "X-Ledger-Batch-Id": "" if missing_database else "1",
        "X-Ledger-Stale": "false", "X-Ledger-Fallback": "false", "X-Ledger-No-Data": "true",
    }
    workbook, roots = _reopen(content, tmp_path)
    try:
        assert workbook.sheetnames == ["positions", "metadata"]
        assert list(workbook["positions"].values) == [POSITION_EXPORT_COLUMNS]
        _assert_metadata(workbook, roots, {
            "batch_id": None if missing_database else 1,
            "requested_as_of_date": REPORT_DATE,
            "resolved_as_of_date": None if missing_database else REPORT_DATE,
            "as_of_date": None if missing_database else REPORT_DATE,
            "source_version": None if missing_database else "sv_synthetic_G1",
            "rule_version": None if missing_database else LEDGER_CLASSIFICATION_RULE_VERSION,
            "stale": None if missing_database else False,
            "fallback": None if missing_database else False,
            "total": None if missing_database else 0,
            "filters": json.dumps(filters, ensure_ascii=False, sort_keys=True),
        })
    finally:
        workbook.close()
