from __future__ import annotations

import os
from calendar import monthrange
from datetime import date
from decimal import Decimal
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

from backend.app.repositories.governance_repo import CACHE_BUILD_RUN_STREAM, GovernanceRepository
from backend.app.services import ledger_pnl_service
from backend.app.services.product_category_source_service import (
    SourcePair,
    _parse_ledger_workbook,
    build_canonical_facts,
    build_ledger_only_facts,
    discover_source_pairs,
)


def _write_ledger(
    directory: Path,
    month: int,
    *,
    annual: bool = False,
    period_start: date | None = None,
    period_end: date | None = None,
) -> Path:
    report_date = date(2024, month, monthrange(2024, month)[1])
    start = period_start or date(2024, 1 if annual else month, 1)
    end = period_end or report_date
    workbook = Workbook()
    workbook.remove(workbook.active)
    for currency, title, factor in (("CNX", "综本", 1), ("CNY", "人民币", 2)):
        sheet = workbook.create_sheet(f"结转前{title}" if annual else title)
        # The actual December 2025 CNY workbook has a leading column, including B5.
        offset = 1 if currency == "CNY" else 0
        sheet.cell(5, 1 + offset, f"会计期间： {start}--{end} 币种： {currency}")
        sheet.append([None] * offset + [
            "组合科目代码", "组合科目名称", "币种", "期初余额", "本期借方", "本期贷方", "期末余额",
        ])
        for code, cash in (("50204000001", 150 if annual else 10), ("52300030001", -70 if annual else -5)):
            amount = Decimal(cash) / factor
            sheet.append([None] * offset + [
                code, "测试科目", currency, 17, max(-amount, 0), max(amount, 0), 23,
            ])
        if month == 1:
            # An account absent from the annual workbook must reverse its earlier flow.
            sheet.append([None] * offset + ["51702010001", "历史科目", currency, 0, 0, 7, -7])
    path = directory / f"总账对账2024{month:02d}.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def _write_average(directory: Path, month: int) -> Path:
    workbook = Workbook()
    for sheet in (workbook.active, workbook.create_sheet("月")):
        sheet.cell(4, 1, "CNX")
        sheet.cell(4, 2, "120")
        sheet.cell(4, 3, 12345)
    path = directory / f"日均2024{month:02d}.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def _write_year(directory: Path) -> tuple[Path, Path]:
    for month in range(1, 12):
        _write_ledger(directory, month)
    return _write_ledger(directory, 12, annual=True), _write_average(directory, 12)


@pytest.mark.parametrize("builder", [build_canonical_facts, build_ledger_only_facts])
def test_annual_source_is_differenced_for_both_fact_builders(tmp_path: Path, builder):
    ledger_path, avg_path = _write_year(tmp_path)
    pair = SourcePair("202412", date(2024, 12, 31), ledger_path, avg_path, "sv_test")

    facts = {(row.account_code, row.currency): row for row in builder(pair)}

    assert facts[("50204000001", "CNX")].monthly_pnl == Decimal("40")
    assert facts[("50204000001", "CNY")].monthly_pnl == Decimal("20")
    assert facts[("52300030001", "CNX")].monthly_pnl == Decimal("-15")
    assert facts[("52300030001", "CNY")].monthly_pnl == Decimal("-7.5")
    assert facts[("51702010001", "CNX")].monthly_pnl == Decimal("-7")
    assert facts[("51702010001", "CNX")].ending_balance == Decimal("0")
    assert facts[("50204000001", "CNX")].beginning_balance == Decimal("17")
    assert facts[("50204000001", "CNX")].ending_balance == Decimal("23")
    if builder is build_canonical_facts:
        assert facts[("120", "CNX")].daily_avg_balance == Decimal("12345")
        assert facts[("120", "CNX")].annual_avg_balance == Decimal("12345")


def test_monthly_source_keeps_its_own_period_flow_and_shifted_cny_header(tmp_path: Path):
    path = _write_ledger(tmp_path, 12)
    rows = _parse_ledger_workbook(path)
    assert rows[("50204000001", "CNX")]["monthly_pnl"] == Decimal("10")
    assert rows[("52300030001", "CNY")]["monthly_pnl"] == Decimal("-2.5")


@pytest.mark.parametrize("read", [_parse_ledger_workbook, lambda path: discover_source_pairs(path.parent)])
def test_annual_source_rejects_missing_historical_month(tmp_path: Path, read):
    path, _ = _write_year(tmp_path)
    (tmp_path / "总账对账202406.xlsx").unlink()
    with pytest.raises(ValueError, match="202406"):
        read(path)


def test_annual_source_rejects_historical_workbook_with_wrong_period(tmp_path: Path):
    path, _ = _write_year(tmp_path)
    _write_ledger(tmp_path, 6, period_start=date(2024, 1, 1))
    with pytest.raises(ValueError, match="period"):
        _parse_ledger_workbook(path)


@pytest.mark.parametrize("header", [
    None,
    "会计期间： 2024-12-01--2024-12-30",
    "会计期间： 2024-10-01--2024-12-31",
    "会计期间： 2024-12-01--2024-12-32",
])
def test_ledger_period_must_be_present_and_supported(tmp_path: Path, header: str | None):
    path = _write_ledger(tmp_path, 12)
    workbook = load_workbook(path)
    workbook["综本"]["A5"] = header
    workbook["人民币"]["B5"] = header
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="period"):
        _parse_ledger_workbook(path)


def test_ledger_currencies_must_have_the_same_source_period(tmp_path: Path):
    path = _write_ledger(tmp_path, 12)
    workbook = load_workbook(path)
    workbook["人民币"]["B5"] = "会计期间： 2024-01-01--2024-12-31"
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="period"):
        _parse_ledger_workbook(path)


def test_annual_source_version_includes_prior_ledgers_only(tmp_path: Path):
    _write_year(tmp_path)
    _write_average(tmp_path, 1)
    before = {pair.month_key: pair for pair in discover_source_pairs(tmp_path)}
    assert len(before["202412"].ledger_dependencies) == 11
    assert before["202401"].ledger_dependencies == ()

    prior = tmp_path / "总账对账202406.xlsx"
    workbook = load_workbook(prior)
    workbook["综本"]["F7"] = 11
    workbook.save(prior)
    workbook.close()
    after = {pair.month_key: pair for pair in discover_source_pairs(tmp_path)}
    assert after["202412"].source_version != before["202412"].source_version
    assert after["202401"].source_version == before["202401"].source_version
    (tmp_path / "日均202401.xlsx").write_bytes(b"prior average changed")
    after_average = {pair.month_key: pair for pair in discover_source_pairs(tmp_path)}
    assert after_average["202412"].source_version == after["202412"].source_version
    assert after_average["202401"].source_version != after["202401"].source_version


def test_ledger_only_cache_tracks_history_but_not_average_files(tmp_path: Path, monkeypatch):
    _, avg = _write_year(tmp_path)
    pair = discover_source_pairs(tmp_path)[0]
    calls = []

    def fake_build(selected_pair):
        calls.append(selected_pair)
        return [len(calls)]

    monkeypatch.setattr(ledger_pnl_service, "_FACTS_CACHE", {})
    monkeypatch.setattr(ledger_pnl_service, "build_ledger_only_facts", fake_build)
    assert ledger_pnl_service._cached_ledger_only_facts(pair) == [1]
    avg.write_bytes(b"average changed")
    assert ledger_pnl_service._cached_ledger_only_facts(pair) == [1]

    dependency = tmp_path / "总账对账202406.xlsx"
    original_stat = dependency.stat()
    # Hold metadata stable: history content is part of the dependency fingerprint.
    content = dependency.read_bytes()
    dependency.write_bytes(bytes([content[0] ^ 1]) + content[1:])
    os.utime(dependency, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))
    assert ledger_pnl_service._cached_ledger_only_facts(pair) == [2]


def test_invalid_source_period_records_failed_run_and_preserves_read_model(tmp_path: Path):
    # Golden HTTP replays may evict this task after test collection.
    from backend.app.tasks.product_category_pnl import _materialize_product_category_pnl

    source = tmp_path / "source"
    source.mkdir()
    ledger = _write_ledger(source, 1)
    _write_average(source, 1)
    database = tmp_path / "test.duckdb"
    governance = tmp_path / "governance"
    kwargs = dict(duckdb_path=str(database), source_dir=str(source), governance_dir=str(governance))
    assert _materialize_product_category_pnl(**kwargs, run_id="valid-period")["status"] == "completed"
    with duckdb.connect(str(database), read_only=True) as connection:
        before = connection.execute("select * from product_category_pnl_formal_read_model order by view, category_id").fetchall()

    workbook = load_workbook(ledger)
    workbook["综本"]["A5"] = None
    workbook.save(ledger)
    workbook.close()
    with pytest.raises(ValueError, match="source period"):
        _materialize_product_category_pnl(**kwargs, run_id="invalid-period")

    with duckdb.connect(str(database), read_only=True) as connection:
        after = connection.execute("select * from product_category_pnl_formal_read_model order by view, category_id").fetchall()
    assert after == before
    records = [
        row for row in GovernanceRepository(base_dir=governance).read_all(CACHE_BUILD_RUN_STREAM)
        if row.get("run_id") == "invalid-period"
    ]
    assert [row["status"] for row in records] == ["running", "failed"]
    assert "source period" in records[-1]["error_message"]
