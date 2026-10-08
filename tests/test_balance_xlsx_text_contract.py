"""Synthetic core -> response JSON -> XLSX representation contracts; no Excel execution."""

from __future__ import annotations

import socket
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from backend.app.core_finance.balance_analysis import (
    FormalTywBalanceFactRow,
    FormalZqtzBalanceFactRow,
)
from backend.app.core_finance.balance_analysis_workbook import build_balance_analysis_workbook_payload
from backend.app.schemas.balance_analysis import BalanceAnalysisWorkbookPayload
from backend.app.services.balance_analysis_workbook_service import _build_balance_analysis_workbook_xlsx_bytes


SHEETS = ("概览", "债券持仓", "同业持仓", "期限分布", "利率分布")
TABLE_COLUMNS = {
    "bond_business_types": (
        "bond_type", "count", "balance_amount", "share", "weighted_rate_pct",
        "weighted_term_years", "amortized_cost_amount", "market_value_amount", "floating_pnl_amount",
    ),
    "counterparty_types": (
        "counterparty_type", "asset_count", "asset_amount", "asset_weighted_rate_pct",
        "liability_count", "liability_amount", "liability_weighted_rate_pct", "net_position_amount",
    ),
    "maturity_gap": (
        "bucket", "bond_assets_amount", "interbank_assets_amount", "asset_total_amount",
        "issuance_amount", "interbank_liabilities_amount", "full_scope_liability_amount",
        "gap_amount", "full_scope_gap_amount", "cumulative_gap_amount", "asset_share",
        "liability_share", "asset_weighted_rate_pct", "liability_weighted_rate_pct", "spread_bp",
    ),
    "rate_distribution": (
        "bucket", "bond_count", "bond_amount", "interbank_asset_count",
        "interbank_asset_amount", "interbank_liability_count", "interbank_liability_amount",
    ),
}
TEXT_CASES = (
    "=1+1", "#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!",
    "中文类型", "'already quoted", "001234", "1234567890123456789", " 001234 ",
    " \t ", "2026-10-06", "+1", "-1",
)
XML_NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@pytest.fixture(autouse=True)
def _deny_network(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("Network is forbidden in synthetic XLSX tests")

    for method in ("connect", "connect_ex", "sendto"):
        monkeypatch.setattr(socket.socket, method, denied)
    monkeypatch.setattr(socket, "create_connection", denied)
    monkeypatch.setattr(socket, "getaddrinfo", denied)


def _formal_tyw(*, label: str, label_source: str, scope: str = "asset", amount: str = "12345.6789"):
    return FormalTywBalanceFactRow(
        report_date=date(2026, 10, 6), position_id="synthetic-tyw", product_type="同业存款",
        position_side=scope, counterparty_name=label if label_source == "counterparty_name" else "unused fallback",
        account_type="synthetic", special_account_type="",
        core_customer_type=label if label_source == "core_customer_type" else "",
        invest_type_std="H", accounting_basis="AC", position_scope=scope,
        currency_basis="CNY", currency_code="CNY", principal_amount=Decimal(amount),
        accrued_interest_amount=Decimal("0"), funding_cost_rate=Decimal("2.75"),
        maturity_date=date(2027, 4, 6),
    )


def _formal_bond():
    return FormalZqtzBalanceFactRow(
        report_date=date(2026, 10, 6), instrument_code="SYNTHETIC", instrument_name="合成国债",
        portfolio_name="synthetic", cost_center="", account_category="", asset_class="债券",
        bond_type="国债", issuer_name="synthetic", industry_name="", rating="AAA",
        invest_type_std="H", accounting_basis="AC", position_scope="asset", currency_basis="CNY",
        currency_code="CNY", face_value_amount=Decimal("10000"), market_value_amount=Decimal("9999.5"),
        amortized_cost_amount=Decimal("10000"), accrued_interest_amount=Decimal("0"),
        coupon_rate=Decimal("2.25"), ytm_value=None, maturity_date=date(2027, 10, 6),
        interest_mode="固定利率", is_issuance_like=False,
    )


def _core_and_json(*, tyw_rows=(), zqtz_rows=()):
    core = build_balance_analysis_workbook_payload(
        report_date=date(2026, 10, 6), position_scope="all", currency_basis="CNY",
        zqtz_rows=list(zqtz_rows), tyw_rows=list(tyw_rows),
    )
    # The production service splits these sections before the real schema JSON dump.
    model = BalanceAnalysisWorkbookPayload(
        report_date=core["report_date"], position_scope=core["position_scope"],
        currency_basis=core["currency_basis"], cards=core["cards"],
        tables=[table for table in core["tables"] if table["section_kind"] == "table"],
        operational_sections=[table for table in core["tables"] if table["section_kind"] != "table"],
    )
    return core, model.model_dump(mode="json")


def _table(payload, key):
    return next(table for table in payload["tables"] if table["key"] == key)


def _roundtrip(payload, tmp_path):
    raw = _build_balance_analysis_workbook_xlsx_bytes(payload)
    (tmp_path / "balance.xlsx").write_bytes(raw)
    workbook = load_workbook(BytesIO(raw), data_only=False)
    assert tuple(workbook.sheetnames) == SHEETS
    with ZipFile(BytesIO(raw)) as archive:
        xml = {
            name: ElementTree.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            for index, name in enumerate(SHEETS, start=1)
        }
    return workbook, xml


def _literal(workbook, xml, sheet_name, coordinate, expected):
    cell = workbook[sheet_name][coordinate]
    if expected in (None, ""):
        assert cell.value is None
        return
    assert cell.value == expected
    assert cell.data_type == "s"
    element = xml[sheet_name].find(f".//x:c[@r='{coordinate}']", XML_NS)
    assert element is not None and element.get("t") in ("inlineStr", "s")
    assert element.find("x:f", XML_NS) is None
    assert "".join(element.itertext()) == expected


def _numeric(cell, expected, key):
    if expected is None:
        assert cell.value is None
        assert cell.number_format == "General"
        return
    number = Decimal(str(expected))
    assert cell.data_type == "n"
    assert cell.value == pytest.approx(float(number))
    assert cell.alignment.horizontal == "right"
    if key == "value" or key.endswith(("_amount", "_value")):
        assert cell.number_format == "#,##0.00000000"
    elif number == number.to_integral_value():
        assert cell.number_format == "0"
    else:
        assert cell.number_format == "0.00000000"


def _assert_current_tables(payload, workbook, xml):
    for sheet_name, (key, expected_columns) in zip(SHEETS[1:], TABLE_COLUMNS.items(), strict=True):
        table = _table(payload, key)
        assert tuple(column["key"] for column in table["columns"]) == expected_columns
        sheet = workbook[sheet_name]
        assert sheet.max_row == 1 + len(table["rows"])
        assert sheet.max_column == len(expected_columns)
        for column_index, column in enumerate(table["columns"], start=1):
            cell = sheet.cell(1, column_index)
            _literal(workbook, xml, sheet_name, cell.coordinate, column["label"])
            assert cell.font.bold
        for row_index, row in enumerate(table["rows"], start=2):
            for column_index, column_key in enumerate(expected_columns, start=1):
                cell = sheet.cell(row_index, column_index)
                if column_index == 1:
                    _literal(workbook, xml, sheet_name, cell.coordinate, row.get(column_key))
                else:
                    _numeric(cell, row.get(column_key), column_key)
    for root in xml.values():
        assert root.findall(".//x:f", XML_NS) == []
        assert root.findall(".//x:c[@t='e']", XML_NS) == []


@pytest.mark.parametrize("label_source", ("core_customer_type", "counterparty_name"))
@pytest.mark.parametrize("label", TEXT_CASES)
def test_formal_core_schema_json_preserves_counterparty_text(label, label_source, tmp_path):
    formal = _formal_tyw(label=label, label_source=label_source)
    core, payload = _core_and_json(tyw_rows=[formal], zqtz_rows=[_formal_bond()])
    core_row = _table(core, "counterparty_types")["rows"][0]
    json_row = _table(payload, "counterparty_types")["rows"][0]
    assert core_row["counterparty_type"] == json_row["counterparty_type"] == label
    assert isinstance(core_row["asset_amount"], Decimal)
    assert isinstance(core_row["asset_weighted_rate_pct"], Decimal)
    assert json_row["asset_amount"] == "1.23456789"
    assert json_row["asset_weighted_rate_pct"] == "2.75"
    assert formal.report_date == date(2026, 10, 6)
    assert payload["report_date"] == "2026-10-06"
    assert {section["section_kind"] for section in payload["operational_sections"]} == {
        "decision_items", "event_calendar", "risk_alerts",
    }
    workbook, xml = _roundtrip(payload, tmp_path)
    try:
        _assert_current_tables(payload, workbook, xml)
        assert len(core["cards"]) == len(payload["cards"]) == 5
        _literal(workbook, xml, "概览", "A1", "label")
        _literal(workbook, xml, "概览", "B1", "value")
        for index, (core_card, card) in enumerate(zip(core["cards"], payload["cards"], strict=True), start=2):
            assert isinstance(core_card["value"], Decimal)
            assert isinstance(card["value"], str)
            _literal(workbook, xml, "概览", f"A{index}", card["label"])
            _numeric(workbook["概览"][f"B{index}"], card["value"], "value")
    finally:
        workbook.close()


@pytest.mark.parametrize("key", TABLE_COLUMNS)
def test_current_dimension_writer_preserves_exact_strings(key, tmp_path):
    # Bond/bucket text here is supplementary writer input, not a claim about core labels.
    _, payload = _core_and_json()
    table = _table(payload, key)
    dimension = TABLE_COLUMNS[key][0]
    table["rows"] = [{dimension: value} for value in (*TEXT_CASES, "", None)]
    workbook, xml = _roundtrip(payload, tmp_path)
    try:
        _assert_current_tables(payload, workbook, xml)
    finally:
        workbook.close()


def test_all_current_numeric_columns_keep_zero_negative_null_and_formats(tmp_path):
    _, payload = _core_and_json()
    for key, columns in TABLE_COLUMNS.items():
        table = _table(payload, key)
        table["rows"] = [
            {columns[0]: "synthetic", **{column: value for column in columns[1:]}}
            for value in (0, -2, None, "0", "-2.125", "2.00000000")
        ]
    workbook, xml = _roundtrip(payload, tmp_path)
    try:
        _assert_current_tables(payload, workbook, xml)
    finally:
        workbook.close()


@pytest.mark.parametrize("target", ("headers", "card-labels", "card-values"))
def test_headers_card_labels_and_nonnumeric_values_are_literal(target, tmp_path):
    _, payload = _core_and_json()
    text_values = ("=1+1", "#N/A", "#REF!", "中文", "'literal", "2026-10-06")
    payload["cards"] = [
        {
            "key": f"synthetic-{index}",
            "label": value if target == "card-labels" else f"safe-label-{index}",
            "value": value if target == "card-values" else "0",
        }
        for index, value in enumerate(text_values)
    ]
    payload["cards"].extend([
        {"key": "numeric-label", "label": "001234", "value": "-12.125"},
        {"key": "zero", "label": "zero", "value": "0"},
        {"key": "legacy-number", "label": "legacy-number", "value": "001234"},
    ])
    payload = BalanceAnalysisWorkbookPayload(**payload).model_dump(mode="json")
    if target == "headers":
        for table in payload["tables"]:
            for index, column in enumerate(table["columns"]):
                column["label"] = TEXT_CASES[index % len(TEXT_CASES)]
    workbook, xml = _roundtrip(payload, tmp_path)
    try:
        _assert_current_tables(payload, workbook, xml)
        for index, card in enumerate(payload["cards"][:len(text_values)], start=2):
            _literal(workbook, xml, "概览", f"A{index}", card["label"])
            if target == "card-values":
                _literal(workbook, xml, "概览", f"B{index}", card["value"])
            else:
                _numeric(workbook["概览"][f"B{index}"], "0", "value")
        for index, card in enumerate(payload["cards"][len(text_values):], start=len(text_values) + 2):
            _literal(workbook, xml, "概览", f"A{index}", card["label"])
            _numeric(workbook["概览"][f"B{index}"], card["value"], "value")
    finally:
        workbook.close()


@pytest.mark.parametrize("alias,current_key,sheet_name", (
    ("zqtz_balance", "bond_business_types", "债券持仓"),
    ("tyw_balance", "counterparty_types", "同业持仓"),
    ("maturity_distribution", "maturity_gap", "期限分布"),
))
def test_compatibility_precedence_and_existing_coercion(alias, current_key, sheet_name, tmp_path):
    # Historical numeric-looking text dimensions are not known; preserve their coercion.
    _, payload = _core_and_json()
    current = _table(payload, current_key)
    dimension = TABLE_COLUMNS[current_key][0]
    current["rows"] = [{dimension: "current-not-selected"}]
    alias_table = deepcopy(current)
    alias_table["key"] = alias
    alias_table["columns"] = [
        {"key": dimension, "label": "=1+1"},
        {"key": "event_date", "label": "#N/A"},
    ]
    alias_table["rows"] = [
        {dimension: value, "event_date": date(2026, 10, 6)}
        for value in ("001234", "-2.125", "=1+1", "#N/A", "中文", "'literal", "  ", "", None)
    ]
    payload["tables"].insert(0, alias_table)  # Selection follows candidate-key priority, not list order.
    workbook, xml = _roundtrip(payload, tmp_path)
    try:
        sheet = workbook[sheet_name]
        assert sheet.max_row == 10
        _literal(workbook, xml, sheet_name, "A1", "=1+1")
        _literal(workbook, xml, sheet_name, "B1", "#N/A")
        _numeric(sheet["A2"], "1234", dimension)
        _numeric(sheet["A3"], "-2.125", dimension)
        for row_index, value in enumerate(("=1+1", "#N/A", "中文", "'literal"), start=4):
            _literal(workbook, xml, sheet_name, f"A{row_index}", value)
        for row_index in (8, 9, 10):
            assert sheet[f"A{row_index}"].value is None
        for row_index in range(2, 11):
            assert sheet[f"B{row_index}"].value == datetime(2026, 10, 6)
            assert sheet[f"B{row_index}"].data_type == "d"
            assert sheet[f"B{row_index}"].number_format == "yyyy-mm-dd"
    finally:
        workbook.close()


def test_empty_tables_and_missing_required_table_contract(tmp_path):
    _, payload = _core_and_json()
    for table in payload["tables"]:
        table["rows"] = []
    workbook, xml = _roundtrip(payload, tmp_path)
    try:
        _assert_current_tables(payload, workbook, xml)
        assert all(workbook[name].max_row == 1 for name in SHEETS[1:])
    finally:
        workbook.close()
    payload["tables"] = [table for table in payload["tables"] if table["key"] != "bond_business_types"]
    with pytest.raises(RuntimeError, match="Expected one of .*zqtz_balance.*bond_business_types"):
        _build_balance_analysis_workbook_xlsx_bytes(payload)
