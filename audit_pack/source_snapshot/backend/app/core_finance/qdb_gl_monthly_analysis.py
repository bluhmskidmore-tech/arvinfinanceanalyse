from __future__ import annotations

from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from backend.app.core_finance.reconciliation_checks import position_vs_ledger_diff


ZERO = Decimal("0")
ONE_HUNDRED_MILLION = Decimal("100000000")

CONFIG = {
    "DEVIATION_WARN": Decimal("5"),
    "DEVIATION_ALERT": Decimal("10"),
    "DEVIATION_CRITICAL": Decimal("20"),
    "MIN_AMOUNT_11D": Decimal("1"),
}


def _effective_analysis_config(overrides: dict[str, Any] | None) -> dict[str, Decimal]:
    merged = {key: CONFIG[key] for key in CONFIG}
    if not overrides:
        return merged
    for key, value in overrides.items():
        if key in merged:
            merged[key] = Decimal(str(value))
    return merged

CATEGORY_NAMES_3D = {
    "101": "现金",
    "110": "存放央行",
    "114": "存放同业",
    "116": "存放境外同业",
    "120": "同业借出/贴现",
    "122": "个人贷款",
    "123": "公司贷款",
    "130": "逾期贷款",
    "131": "贷款减值准备",
    "132": "福费廷/押汇",
    "133": "应收利息",
    "140": "买入返售",
    "141": "FVTPL金融资产",
    "142": "AC债券投资",
    "143": "AC其他投资",
    "144": "FVOCI金融资产",
    "145": "长期股权投资",
    "201": "单位活期存款",
    "205": "单位定期存款",
    "211": "活期储蓄存款",
    "215": "定期储蓄存款",
    "234": "同业存放",
    "235": "同业定期存款",
    "241": "拆入资金",
    "251": "保证金存款",
    "255": "卖出回购",
    "272": "应付债券/存单",
}

INDUSTRY_NAMES = {
    "01": "农林牧渔",
    "02": "采矿业",
    "03": "制造业",
    "04": "电力燃气水",
    "05": "建筑业",
    "06": "交通运输",
    "07": "信息技术",
    "08": "批发零售",
    "09": "住宿餐饮",
    "10": "金融业",
    "11": "房地产",
    "12": "租赁商务",
    "13": "科技服务",
    "14": "水利环境",
    "15": "居民服务",
    "16": "教育",
    "17": "卫生社保",
    "18": "文体娱乐",
    "19": "公共管理",
    "99": "大额存单",
}

DAILY_AVG_COLUMN_MAP = [
    (1, 2, "CNX", "3d"),
    (5, 6, "CNX", "5d"),
    (9, 10, "CNX", "7d"),
    (13, 14, "CNX", "11d"),
    (17, 18, "CNY", "3d"),
    (21, 22, "CNY", "5d"),
    (25, 26, "CNY", "7d"),
    (29, 30, "CNY", "11d"),
]


def parse_daily_avg(filepath: str | Path) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(filename=str(filepath), read_only=True, data_only=True)
    try:
        result: dict[str, list[dict[str, Any]]] = {}
        for sheet_name, period_label in (("月", "月日均"), ("年", "年日均")):
            worksheet = workbook[sheet_name]
            for code_col, value_col, currency, level in DAILY_AVG_COLUMN_MAP:
                key = f"{period_label}_{currency}_{level}"
                rows: list[dict[str, Any]] = []
                for row in worksheet.iter_rows(min_row=4, values_only=True):
                    code = _normalize_account_code(row[code_col] if len(row) > code_col else None)
                    if not code:
                        continue
                    value = _to_decimal(row[value_col] if len(row) > value_col else None)
                    if value is None:
                        continue
                    rows.append({"科目代码": code, "日均余额": value, "币种": currency, "级别": level})
                result[key] = rows
        return result
    finally:
        workbook.close()


def parse_general_ledger(filepath: str | Path) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(filename=str(filepath), read_only=True, data_only=True)
    try:
        result: dict[str, list[dict[str, Any]]] = {}
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows: list[dict[str, Any]] = []
            for row in worksheet.iter_rows(min_row=7, values_only=True):
                code = _normalize_account_code(row[0] if len(row) > 0 else None)
                if not code:
                    continue
                currency = str(row[2] or "").strip()
                if not currency:
                    continue
                opening = _to_decimal(row[3] if len(row) > 3 else None) or ZERO
                debit = _to_decimal(row[4] if len(row) > 4 else None) or ZERO
                credit = _to_decimal(row[5] if len(row) > 5 else None) or ZERO
                closing = _to_decimal(row[6] if len(row) > 6 else None) or ZERO
                rows.append(
                    {
                        "科目代码": code,
                        "科目名称": str(row[1] or "").strip(),
                        "币种": currency,
                        "期初余额": opening,
                        "本期借方": debit,
                        "本期贷方": credit,
                        "期末余额": closing,
                        "变动额": closing - opening,
                        "本期净额": debit - credit,
                        "大类1": code[:1],
                        "大类3": code[:3],
                        "大类5": code[:5],
                    }
                )
            result[sheet_name] = rows
        return result
    finally:
        workbook.close()


def merge_all(gl_data: dict[str, list[dict[str, Any]]], rj_data: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    cnx_rows = [row for row in gl_data.get("综本", []) if row["币种"] == "CNX"]
    merged: dict[str, Any] = {}

    rj_m_11 = _index_amounts(rj_data.get("月日均_CNX_11d", []))
    rj_y_11 = _index_amounts(rj_data.get("年日均_CNX_11d", []))
    merged["11位"] = [{**row, "月日均": rj_m_11.get(row["科目代码"]), "年日均": rj_y_11.get(row["科目代码"])} for row in cnx_rows]

    grouped_3 = _group_sum(cnx_rows, key="大类3")
    rj_m_3 = _index_amounts(rj_data.get("月日均_CNX_3d", []))
    rj_y_3 = _index_amounts(rj_data.get("年日均_CNX_3d", []))
    merged["3位"] = [
        {
            **row,
            "科目代码": row["group_key"],
            "名称": CATEGORY_NAMES_3D.get(row["group_key"], ""),
            "月日均": rj_m_3.get(row["group_key"]),
            "年日均": rj_y_3.get(row["group_key"]),
        }
        for row in grouped_3
    ]

    rj_m_5 = _index_amounts(rj_data.get("月日均_CNX_5d", []))
    rj_y_5 = _index_amounts(rj_data.get("年日均_CNX_5d", []))
    for prefix, label in (("123", "公司贷款"), ("201", "活期存款"), ("205", "定期存款"), ("251", "保证金")):
        subset = [row for row in cnx_rows if row["大类3"] == prefix]
        grouped_5 = _group_sum(subset, key="大类5", include_name=True)
        merged[f"5位_{label}"] = [
            {
                **row,
                "科目代码": row["group_key"],
                "行业代码": row["group_key"][3:5],
                "行业名称": INDUSTRY_NAMES.get(row["group_key"][3:5], ""),
                "月日均": rj_m_5.get(row["group_key"]),
                "年日均": rj_y_5.get(row["group_key"]),
            }
            for row in grouped_5
        ]

    cny_index = {row["科目代码"]: row for row in gl_data.get("人民币", [])}
    merged["外币分析"] = []
    for row in cnx_rows:
        cny_row = cny_index.get(row["科目代码"])
        if cny_row is None:
            continue
        cnx_value = row["期末余额"]
        cny_value = cny_row["期末余额"]
        foreign_value = cnx_value - cny_value
        foreign_share = None if cnx_value == ZERO else foreign_value / abs(cnx_value) * Decimal("100")
        merged["外币分析"].append(
            {
                "科目代码": row["科目代码"],
                "科目名称": row["科目名称"],
                "期末余额_综本": cnx_value,
                "期末余额_人民币": cny_value,
                "外币部分": foreign_value,
                "外币占比%": foreign_share,
            }
        )
    return merged


def build_qdb_gl_monthly_analysis_workbook(
    *,
    report_month: str,
    merged_data: dict[str, Any],
    threshold_overrides: dict[str, int | float] | None = None,
) -> dict[str, Any]:
    analysis_cfg = _effective_analysis_config(threshold_overrides)
    m3 = compute_deviation(merged_data.get("3位", []))
    m11 = compute_deviation(merged_data.get("11位", []))
    metrics = compute_asset_liability_structure(m3)
    gap_rows = compute_industry_gap(merged_data)
    alert_rows = [
        *generate_alerts(merged_data, analysis_config=analysis_cfg),
        *_position_ledger_reconciliation_alerts(
            _qdb_gl_position_vs_ledger_check(
                position_totals=metrics,
                ledger_rows_3d=m3,
            )
        ),
    ]
    foreign_rows = build_foreign_currency_rows(merged_data.get("外币分析", []))

    sheets = [
        _sheet("overview", "经营概览", ["指标", "值"], [
            {"指标": "总资产(亿)", "值": _display_number(_to_yi(metrics["总资产"]))},
            {"指标": "总负债(亿)", "值": _display_number(_to_yi(metrics["总负债"]))},
            {"指标": "贷款总额(亿)", "值": _display_number(_to_yi(metrics["贷款总额"]))},
            {"指标": "存款总额(亿)", "值": _display_number(_to_yi(metrics["存款总额"]))},
            {"指标": "投资总额(亿)", "值": _display_number(_to_yi(metrics["投资总额"]))},
            {"指标": "存贷比%", "值": _display_number(metrics["存贷比%"])},
            {"指标": "拨贷比%", "值": _display_number(metrics["拨贷比%"])},
            {"指标": "定期化率%", "值": _display_number(metrics["定期化率%"])},
            {"指标": "活期率%", "值": _display_number(metrics["活期率%"])},
            {"指标": "高流动性占比%", "值": _display_number(metrics["高流动性占比%"])},
        ]),
        _sheet("summary_3d", "3位科目总览", ["科目代码", "名称", "期初余额", "期末余额", "变动额", "月日均", "年日均", "偏离额", "偏离%", "趋势额", "趋势%"], [_normalize_amount_row(row, include_trend=True) for row in _sort_rows(m3, "期末余额")]),
        _sheet("asset_structure", "资产结构", ["科目代码", "名称", "期末余额", "月日均", "偏离%", "趋势%"], [
            {"科目代码": row["科目代码"], "名称": row["名称"], "期末余额": _display_number(_to_yi(row["期末余额"])), "月日均": _display_number(_to_yi(row.get("月日均"))), "偏离%": _display_number(row.get("偏离%")), "趋势%": _display_number(row.get("趋势%"))}
            for row in _sort_rows([row for row in m3 if str(row["科目代码"]).startswith("1")], "期末余额")
        ]),
        _sheet("liability_structure", "负债结构", ["科目代码", "名称", "期末余额", "月日均", "偏离%", "趋势%"], [
            {"科目代码": row["科目代码"], "名称": row["名称"], "期末余额": _display_number(_to_yi(row["期末余额"])), "月日均": _display_number(_to_yi(row.get("月日均"))), "偏离%": _display_number(row.get("偏离%")), "趋势%": _display_number(row.get("趋势%"))}
            for row in _sort_rows([row for row in m3 if str(row["科目代码"]).startswith("2")], "期末余额")
        ]),
        _industry_sheet("loan_industry", "贷款行业", compute_deviation(merged_data.get("5位_公司贷款", []))),
        _industry_sheet("deposit_demand_industry", "存款行业_活期", compute_deviation(merged_data.get("5位_活期存款", []))),
        _industry_sheet("deposit_term_industry", "存款行业_定期", compute_deviation(merged_data.get("5位_定期存款", []))),
        _sheet("industry_gap", "行业存贷差", ["行业", "贷款期末", "存款期末", "存贷差_时点", "贷款月日均", "存款月日均", "存贷差_日均"], gap_rows),
        _sheet("top_11d", "11位偏离TOP", ["科目代码", "科目名称", "期末余额", "月日均", "年日均", "偏离额", "偏离%", "趋势%"], build_11d_top_rows(m11, analysis_config=analysis_cfg)),
        _sheet("alerts", "异动预警", ["科目代码", "科目名称", "预警级别", "期末余额(亿)", "月日均(亿)", "偏离额(亿)", "偏离%", "异动类型"], alert_rows),
        _sheet("foreign_currency", "外币分析", ["科目代码", "科目名称", "期末余额_综本", "期末余额_人民币", "外币部分", "外币占比%"], foreign_rows),
    ]
    return {"report_month": report_month, "sheets": sheets}


def export_qdb_gl_monthly_analysis_workbook_xlsx_bytes(workbook_payload: dict[str, Any]) -> bytes:
    workbook = Workbook()
    active = workbook.active
    for index, sheet_payload in enumerate(workbook_payload.get("sheets", [])):
        worksheet = active if index == 0 else workbook.create_sheet()
        worksheet.title = str(sheet_payload["title"])
        columns = list(sheet_payload.get("columns", []))
        for column_index, header in enumerate(columns, start=1):
            worksheet.cell(row=1, column=column_index, value=header)
        for row_index, row in enumerate(sheet_payload.get("rows", []), start=2):
            for column_index, header in enumerate(columns, start=1):
                worksheet.cell(row=row_index, column=column_index, value=_excel_value(row.get(header)))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def compute_deviation(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    computed: list[dict[str, Any]] = []
    for row in rows:
        current = dict(row)
        month_avg = _as_decimal(row.get("月日均"))
        year_avg = _as_decimal(row.get("年日均"))
        ending = _as_decimal(row.get("期末余额"))
        current["偏离额"] = None if month_avg is None else ending - month_avg
        current["偏离%"] = _safe_pct(current["偏离额"], month_avg)
        current["趋势额"] = None if month_avg is None or year_avg is None else month_avg - year_avg
        current["趋势%"] = _safe_pct(current["趋势额"], year_avg)
        computed.append(current)
    return computed


def compute_asset_liability_structure(rows_3d: list[dict[str, Any]]) -> dict[str, Decimal]:
    rows_by_code = {row["科目代码"]: row for row in rows_3d}

    def value(code: str, field: str = "期末余额") -> Decimal:
        row = rows_by_code.get(code)
        return ZERO if row is None else (_as_decimal(row.get(field)) or ZERO)

    loan_total = sum((value(code) for code in ("122", "123", "132", "133")), ZERO)
    deposit_total = abs(sum((value(code) for code in ("201", "205", "211", "215", "251")), ZERO))
    provision = abs(value("131"))
    total_assets = sum((_as_decimal(row.get("期末余额")) or ZERO for row in rows_3d if (_as_decimal(row.get("期末余额")) or ZERO) > ZERO), ZERO)
    total_liabilities = abs(sum((_as_decimal(row.get("期末余额")) or ZERO for row in rows_3d if str(row["科目代码"]).startswith("2")), ZERO))
    investment_total = sum((value(code) for code in ("141", "142", "143", "144", "145")), ZERO)
    liquid_total = sum((value(code) for code in ("101", "110", "114", "116")), ZERO)
    term_deposit = abs(sum((value(code) for code in ("205", "215")), ZERO))
    demand_deposit = abs(sum((value(code) for code in ("201", "211")), ZERO))
    return {
        "贷款总额": loan_total,
        "存款总额": deposit_total,
        "总资产": total_assets,
        "总负债": total_liabilities,
        "投资总额": investment_total,
        "存贷比%": _pct(loan_total, deposit_total),
        "拨贷比%": _pct(provision, loan_total),
        "定期化率%": _pct(term_deposit, deposit_total),
        "活期率%": _pct(demand_deposit, deposit_total),
        "高流动性占比%": _pct(liquid_total, total_assets),
    }


def _qdb_gl_position_vs_ledger_check(
    *,
    position_totals: dict[str, Any],
    ledger_rows_3d: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    ledger_totals = _qdb_gl_ledger_totals(ledger_rows_3d)
    total_assets = _as_decimal(position_totals.get("总资产")) or ZERO
    total_liabilities = _as_decimal(position_totals.get("总负债")) or ZERO
    return position_vs_ledger_diff(
        {
            "total_assets": float(total_assets),
            "total_liabilities": float(total_liabilities),
            "net_assets": float(total_assets - total_liabilities),
        },
        ledger_totals,
        threshold_yuan=0.01,
    )


def _qdb_gl_ledger_totals(rows_3d: list[dict[str, Any]]) -> dict[str, float]:
    total_assets = sum(
        (_as_decimal(row.get("期末余额")) or ZERO)
        for row in rows_3d
        if (_as_decimal(row.get("期末余额")) or ZERO) > ZERO
    )
    total_liabilities = abs(
        sum(
            (_as_decimal(row.get("期末余额")) or ZERO)
            for row in rows_3d
            if str(row.get("科目代码") or "").startswith("2")
        )
    )
    return {
        "total_assets": float(total_assets),
        "total_liabilities": float(total_liabilities),
        "net_assets": float(total_assets - total_liabilities),
    }


def _position_ledger_reconciliation_alerts(checks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    alerts: list[dict[str, Any]] = []
    for check in checks:
        if not check["breached"]:
            continue
        alerts.append(
            {
                "科目代码": check["dimension"],
                "科目名称": "Position vs Ledger",
                "预警级别": "严重",
                "期末余额(亿)": _display_number(_to_yi(Decimal(str(check["position_value"])))),
                "月日均(亿)": _display_number(_to_yi(Decimal(str(check["ledger_value"])))),
                "偏离额(亿)": _display_number(_to_yi(Decimal(str(check["diff"])))),
                "偏离%": None,
                "异动类型": "position_vs_ledger_reconciliation",
            }
        )
    return alerts


def compute_industry_gap(merged_data: dict[str, Any]) -> list[dict[str, Any]]:
    loans = merged_data.get("5位_公司贷款", [])
    demand = {row["行业代码"]: row for row in merged_data.get("5位_活期存款", [])}
    term = {row["行业代码"]: row for row in merged_data.get("5位_定期存款", [])}
    rows: list[dict[str, Any]] = []
    for row in loans:
        industry_code = row["行业代码"]
        demand_row = demand.get(industry_code)
        term_row = term.get(industry_code)
        deposit_end = abs(_as_decimal(demand_row.get("期末余额")) or ZERO) + abs(_as_decimal(term_row.get("期末余额")) or ZERO)
        deposit_avg = abs(_as_decimal(demand_row.get("月日均")) or ZERO) + abs(_as_decimal(term_row.get("月日均")) or ZERO)
        loan_end = _as_decimal(row.get("期末余额")) or ZERO
        loan_avg = _as_decimal(row.get("月日均")) or ZERO
        rows.append({"行业": row.get("行业名称", ""), "贷款期末": _display_number(_to_yi(loan_end)), "存款期末": _display_number(_to_yi(deposit_end)), "存贷差_时点": _display_number(_to_yi(loan_end - deposit_end)), "贷款月日均": _display_number(_to_yi(loan_avg)), "存款月日均": _display_number(_to_yi(deposit_avg)), "存贷差_日均": _display_number(_to_yi(loan_avg - deposit_avg))})
    return rows


def generate_alerts(
    merged_data: dict[str, Any],
    *,
    analysis_config: dict[str, Decimal] | None = None,
) -> list[dict[str, Any]]:
    cfg = analysis_config if analysis_config is not None else CONFIG
    rows_11 = compute_deviation(merged_data.get("11位", []))
    alerts: list[dict[str, Any]] = []
    for row in rows_11:
        deviation_pct = _as_decimal(row.get("偏离%"))
        deviation_amount = _as_decimal(row.get("偏离额"))
        if deviation_pct is None or deviation_amount is None:
            continue
        abs_pct = abs(deviation_pct)
        abs_yi = abs(_to_yi(deviation_amount))
        if abs_pct > cfg["DEVIATION_CRITICAL"] and abs_yi > Decimal("5"):
            level = "严重"
        elif abs_pct > cfg["DEVIATION_ALERT"] and abs_yi > Decimal("2"):
            level = "中度"
        elif abs_pct > cfg["DEVIATION_WARN"] and abs_yi > Decimal("1"):
            level = "轻度"
        else:
            continue
        code3 = str(row["科目代码"])[:3]
        first_digit = str(row["科目代码"])[:1]
        if code3 == "130":
            anomaly = "逾期贷款异动"
        elif code3 in {"140", "255", "234", "235", "114", "120"} and abs_pct > Decimal("30"):
            anomaly = "同业业务月末操纵嫌疑"
        elif first_digit == "2" and deviation_pct < Decimal("-10"):
            anomaly = "负债月末冲量"
        elif first_digit == "1" and deviation_pct < Decimal("-20"):
            anomaly = "资产月末压降"
        else:
            anomaly = "时点vs日均显著偏离"
        alerts.append({"科目代码": row["科目代码"], "科目名称": row.get("科目名称", ""), "预警级别": level, "期末余额(亿)": _display_number(_to_yi(_as_decimal(row.get("期末余额")) or ZERO)), "月日均(亿)": _display_number(_to_yi(_as_decimal(row.get("月日均")) or ZERO)), "偏离额(亿)": _display_number(_to_yi(deviation_amount)), "偏离%": _display_number(deviation_pct), "异动类型": anomaly})
    return alerts


def build_foreign_currency_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    filtered = [row for row in rows if abs(_as_decimal(row.get("外币部分")) or ZERO) > Decimal("10000000")]
    ordered = sorted(filtered, key=lambda row: abs(_as_decimal(row.get("外币部分")) or ZERO), reverse=True)
    return [{"科目代码": row["科目代码"], "科目名称": row["科目名称"], "期末余额_综本": _display_number(_to_yi(_as_decimal(row.get("期末余额_综本")) or ZERO)), "期末余额_人民币": _display_number(_to_yi(_as_decimal(row.get("期末余额_人民币")) or ZERO)), "外币部分": _display_number(_to_yi(_as_decimal(row.get("外币部分")) or ZERO)), "外币占比%": _display_number(_as_decimal(row.get("外币占比%")))} for row in ordered[:30]]


def build_11d_top_rows(
    rows: list[dict[str, Any]],
    *,
    analysis_config: dict[str, Decimal] | None = None,
) -> list[dict[str, Any]]:
    cfg = analysis_config if analysis_config is not None else CONFIG
    filtered = [row for row in rows if abs(_as_decimal(row.get("月日均")) or ZERO) > cfg["MIN_AMOUNT_11D"] * ONE_HUNDRED_MILLION]
    ordered = sorted(filtered, key=lambda row: abs(_as_decimal(row.get("偏离%")) or ZERO), reverse=True)[:50]
    return [{"科目代码": row["科目代码"], "科目名称": row.get("科目名称", ""), "期末余额": _display_number(_to_yi(_as_decimal(row.get("期末余额")) or ZERO)), "月日均": _display_number(_to_yi(_as_decimal(row.get("月日均")) or ZERO)), "年日均": _display_number(_to_yi(_as_decimal(row.get("年日均")) or ZERO)), "偏离额": _display_number(_to_yi(_as_decimal(row.get("偏离额")) or ZERO)), "偏离%": _display_number(_as_decimal(row.get("偏离%"))), "趋势%": _display_number(_as_decimal(row.get("趋势%")))} for row in ordered]


def _industry_sheet(key: str, title: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    ordered = sorted(rows, key=lambda row: abs(_as_decimal(row.get("期末余额")) or ZERO), reverse=True)
    return _sheet(key, title, ["行业名称", "期初余额", "期末余额", "变动额", "月日均", "年日均", "偏离额", "偏离%", "趋势%"], [
        {"行业名称": row.get("行业名称", ""), "期初余额": _display_number(_to_yi(_as_decimal(row.get("期初余额")) or ZERO)), "期末余额": _display_number(_to_yi(_as_decimal(row.get("期末余额")) or ZERO)), "变动额": _display_number(_to_yi(_as_decimal(row.get("变动额")) or ZERO)), "月日均": _display_number(_to_yi(_as_decimal(row.get("月日均")) or ZERO)), "年日均": _display_number(_to_yi(_as_decimal(row.get("年日均")) or ZERO)), "偏离额": _display_number(_to_yi(_as_decimal(row.get("偏离额")) or ZERO)), "偏离%": _display_number(_as_decimal(row.get("偏离%"))), "趋势%": _display_number(_as_decimal(row.get("趋势%")))}
        for row in ordered
    ])


def _normalize_amount_row(row: dict[str, Any], *, include_trend: bool) -> dict[str, Any]:
    normalized = {
        "科目代码": row["科目代码"],
        "名称": row.get("名称", ""),
        "期初余额": _display_number(_to_yi(_as_decimal(row.get("期初余额")) or ZERO)),
        "期末余额": _display_number(_to_yi(_as_decimal(row.get("期末余额")) or ZERO)),
        "变动额": _display_number(_to_yi(_as_decimal(row.get("变动额")) or ZERO)),
        "月日均": _display_number(_to_yi(_as_decimal(row.get("月日均")) or ZERO)),
        "年日均": _display_number(_to_yi(_as_decimal(row.get("年日均")) or ZERO)),
        "偏离额": _display_number(_to_yi(_as_decimal(row.get("偏离额")) or ZERO)),
        "偏离%": _display_number(_as_decimal(row.get("偏离%"))),
    }
    if include_trend:
        normalized["趋势额"] = _display_number(_to_yi(_as_decimal(row.get("趋势额")) or ZERO))
        normalized["趋势%"] = _display_number(_as_decimal(row.get("趋势%")))
    return normalized


def _sheet(key: str, title: str, columns: list[str], rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {"key": key, "title": title, "columns": columns, "rows": rows}


def _group_sum(rows: list[dict[str, Any]], *, key: str, include_name: bool = False) -> list[dict[str, Any]]:
    buckets: dict[str, dict[str, Any]] = {}
    for row in rows:
        group_key = row[key]
        bucket = buckets.setdefault(group_key, {"group_key": group_key, "期初余额": ZERO, "期末余额": ZERO, "变动额": ZERO, "本期借方": ZERO, "本期贷方": ZERO})
        for field in ("期初余额", "期末余额", "变动额", "本期借方", "本期贷方"):
            bucket[field] += _as_decimal(row.get(field)) or ZERO
        if include_name and not bucket.get("科目名称"):
            bucket["科目名称"] = row.get("科目名称", "")
    return list(buckets.values())


def _index_amounts(rows: list[dict[str, Any]]) -> dict[str, Decimal]:
    return {row["科目代码"]: _as_decimal(row.get("日均余额")) or ZERO for row in rows}


def _sort_rows(rows: list[dict[str, Any]], field: str) -> list[dict[str, Any]]:
    return sorted(rows, key=lambda row: abs(_as_decimal(row.get(field)) or ZERO), reverse=True)


def _normalize_account_code(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        return str(int(value))
    text = str(value).strip()
    return text or None


def _to_decimal(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def _as_decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    return _to_decimal(value)


def _safe_pct(numerator: Decimal | None, denominator: Decimal | None) -> Decimal | None:
    if numerator is None or denominator is None or denominator == ZERO:
        return None
    return numerator / abs(denominator) * Decimal("100")


def _pct(numerator: Decimal, denominator: Decimal) -> Decimal:
    return ZERO if denominator == ZERO else numerator / denominator * Decimal("100")


def _to_yi(value: Decimal | None) -> Decimal:
    return (value or ZERO) / ONE_HUNDRED_MILLION


def _display_number(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    normalized = value.quantize(Decimal("0.01"))
    return int(normalized) if normalized == normalized.to_integral_value() else float(normalized)


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value
